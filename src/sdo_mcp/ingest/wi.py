"""Wisconsin Department of Administration population and household adapter.

The Demographic Services Center's December 2024 (Vintage 2024) population
workbook publishes state and all 72 county totals at 2020, 2030, 2040, and
2050. The separate household workbook has the same year grid but no stated
edition date. Municipal
and minor-civil-division products are published by DOA but are intentionally
not ingested in this pass.
"""

from __future__ import annotations

import io
import math
import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source

RAW_DIR = REPO_ROOT / "data" / "raw" / "wi"
POP_FILE = "Pop_Proj_Co_State_2020_2050.xlsx"
HH_FILE = "HH_Proj_Co_State_2020-2050.xlsx"
POP_METHOD_FILE = "Method_Pop_Proj_2020_2050.pdf"
HH_METHOD_FILE = "Method_HH_Proj_Co_2020-2050.pdf"
POP_URL = "https://doa.wi.gov/DIR/Pop_Proj_Co_State_2020_2050.xlsx"
HH_URL = "https://doa.wi.gov/DIR/HH_Proj_Co_State_2020-2050.xlsx"
POP_METHOD_URL = "https://doa.wi.gov/DIR/Method_Pop_Proj_2020_2050.pdf"
HH_METHOD_URL = "https://doa.wi.gov/DIR/Method_HH_Proj_Co_2020-2050.pdf"
STATE_FIPS = "55"
POP_SOURCE_ID = "wi-pop-2024"
HH_SOURCE_ID = "wi-households-2020-2050"
BASE_YEAR = 2020
HORIZON_YEAR = 2050
AGENCY = "Wisconsin Department of Administration, Division of Intergovernmental Relations, Demographic Services Center"
LICENSE = "Public Wisconsin state government statistics; cite the Department of Administration."
YEARS = (2020, 2030, 2040, 2050)


def acquire() -> None:
    """Download and validate both official workbooks and their methods PDFs."""
    import openpyxl

    entries = []
    for filename, url, sheets, title in (
        (POP_FILE, POP_URL, {"Notes", "Co_Proj_2020_2050_totals", "Co_Proj_2020_2050_x_age_x sex"}, "County Age-Sex Population Projections, 2020-2050"),
        (HH_FILE, HH_URL, {"Notes", "Household Projection Summary"}, "Household Projections by County, 2020-2050"),
    ):
        body = fetch.get_bytes(url)
        if not body.startswith(b"PK"):
            raise RuntimeError(f"{url}: body is not an xlsx (HTML or soft-200 error)")
        wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        try:
            if set(wb.sheetnames) != sheets:
                raise RuntimeError(f"{url}: unexpected workbook sheets {wb.sheetnames!r}")
            if wb["Notes"]["A1"].value != title:
                raise RuntimeError(f"{url}: unexpected workbook title")
        finally:
            wb.close()
        fetch.write_snapshot(RAW_DIR / filename, body)
        entries.append(dict(file=filename, source=url, retrieved_at=date.today().isoformat(),
                            sha256=fetch.sha256(body), records=73 * 4,
                            note=("Wisconsin DOA Vintage 2024 state/county population projection workbook."
                                  if filename == POP_FILE else
                                  "Wisconsin DOA state/county household projection workbook, 2020-2050 (edition date not stated in workbook).")))

    for filename, url in (
        (POP_METHOD_FILE, POP_METHOD_URL),
        (HH_METHOD_FILE, HH_METHOD_URL),
    ):
        body = fetch.get_bytes(url)
        if not body.startswith(b"%PDF"):
            raise RuntimeError(f"{url}: body is not a PDF (HTML or soft-200 error)")
        fetch.validate_pdf(body, url)
        fetch.write_snapshot(RAW_DIR / filename, body)
        entries.append(dict(file=filename, source=url, retrieved_at=date.today().isoformat(),
                            sha256=fetch.sha256(body), records=None,
                            note="Official Wisconsin DOA projection methodology."))
    fetch.write_manifest(RAW_DIR, entries)
    print(f"WI acquire: {POP_FILE}, {HH_FILE}")


def _retrieved(filename: str) -> str:
    return next(e["retrieved_at"] for e in fetch.read_manifest(RAW_DIR) if e["file"] == filename)


def _county_fips(conn) -> dict[str, str]:
    return {
        row["name"].removesuffix(" County"): row["geo_id"]
        for row in conn.execute("SELECT geo_id, name FROM geographies WHERE state='WI' AND geo_level='county'")
    }


def _insert_population(conn, fips_by_name: dict[str, str]) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(RAW_DIR / POP_FILE, read_only=True, data_only=True)
    try:
        ws = wb["Co_Proj_2020_2050_totals"]
        rows = list(ws.iter_rows(values_only=True))
        if tuple(rows[2][:5]) != ("County", "Both Sexes Population 2020", "Both Sexes Projection 2030", "Both Sexes Projection 2040", "Both Sexes Projection 2050"):
            raise ValueError("WI population: unexpected totals header")
        observations = []
        seen = set()
        for rec in rows[3:]:
            name = rec[0]
            if not isinstance(name, str) or not name.strip():
                continue
            values = rec[1:5]
            if len(values) != 4 or any(not isinstance(v, (int, float)) or isinstance(v, bool) or not math.isfinite(float(v)) or float(v) < 0 or int(v) != v for v in values):
                raise ValueError(f"WI population: invalid values for {name!r}")
            if name == "Wisconsin":
                level, gid = "state", STATE_FIPS
            else:
                gid = fips_by_name.get(name)
                if gid is None:
                    raise RuntimeError(f"WI population: county {name!r} has no Census FIPS match")
                level = "county"
            if gid in seen:
                raise ValueError(f"WI population: duplicate geography {gid}")
            seen.add(gid)
            for year, value in zip(YEARS, values):
                observations.append(dict(state="WI", geo_level=level, geo_id=gid, year=year,
                                          variable="population", value=int(value),
                                          datatype="estimate" if year == BASE_YEAR else "projection",
                                          source_id=POP_SOURCE_ID))
        expected = set(fips_by_name.values()) | {STATE_FIPS}
        if seen != expected:
            raise RuntimeError(f"WI population: geography mismatch; missing {sorted(expected - seen)}")
        return observations
    finally:
        wb.close()


def _insert_households(conn, fips_by_name: dict[str, str]) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(RAW_DIR / HH_FILE, read_only=True, data_only=True)
    try:
        ws = wb["Household Projection Summary"]
        rows = list(ws.iter_rows(values_only=True))
        if tuple(rows[4][2:6]) != ("2020 ", "2030 ", "2040 ", "2050 "):
            raise ValueError("WI households: unexpected year header")
        observations = []
        seen = set()
        for rec in rows[5:]:
            if rec[1] != "All Ages":
                continue
            name = rec[0]
            values = rec[2:6]
            if not isinstance(name, str) or len(values) != 4:
                raise ValueError(f"WI households: malformed row {rec[:6]!r}")
            if any(not isinstance(v, (int, float)) or isinstance(v, bool) or not math.isfinite(float(v)) or float(v) < 0 or int(v) != v for v in values):
                raise ValueError(f"WI households: invalid values for {name!r}")
            if name == "Wisconsin":
                level, gid = "state", STATE_FIPS
            else:
                gid = fips_by_name.get(name)
                if gid is None:
                    raise RuntimeError(f"WI households: county {name!r} has no Census FIPS match")
                level = "county"
            if gid in seen:
                raise ValueError(f"WI households: duplicate geography {gid}")
            seen.add(gid)
            for year, value in zip(YEARS, values):
                observations.append(dict(state="WI", geo_level=level, geo_id=gid, year=year,
                                          variable="households", value=int(value),
                                          datatype="estimate" if year == BASE_YEAR else "projection",
                                          source_id=HH_SOURCE_ID))
        expected = set(fips_by_name.values()) | {STATE_FIPS}
        if seen != expected:
            raise RuntimeError(f"WI households: geography mismatch; missing {sorted(expected - seen)}")
        return observations
    finally:
        wb.close()


def load(conn) -> None:
    fips_by_name = _county_fips(conn)
    if len(fips_by_name) != 72:
        raise RuntimeError(f"WI load: expected 72 counties, found {len(fips_by_name)}")
    rows = _insert_population(conn, fips_by_name)
    rows += _insert_households(conn, fips_by_name)
    upsert_source(
        conn, source_id=POP_SOURCE_ID, state="WI", agency=AGENCY,
        product_name="State and County Age-Sex Population Projections, 2020-2050",
        vintage="Vintage 2024 (December 2024)", base_year=BASE_YEAR, horizon_year=HORIZON_YEAR,
        source_url=POP_URL, retrieved_at=_retrieved(POP_FILE), license=LICENSE,
        citation="Wisconsin Department of Administration. State and County Age-Sex Population Projections, Vintage 2024.",
        method_notes=(
            "County totals are published at 2020, 2030, 2040, and 2050. The 2020 values are the Census base/estimate; later points are projections. "
            "The modified Hamilton-Perry method projects county age-sex cohorts; published cohorts are rounded to the nearest five and summed, with state control totals. "
            "This adapter loads totals only. DOA separately publishes minor civil division and municipal projections, which remain available but un-ingested."
        ),
    )
    upsert_source(
        conn, source_id=HH_SOURCE_ID, state="WI", agency=AGENCY,
        product_name="State and County Household Projections, 2020-2050",
        vintage="2020-2050 series (edition date not stated in workbook)", base_year=BASE_YEAR, horizon_year=HORIZON_YEAR,
        source_url=HH_URL, retrieved_at=_retrieved(HH_FILE), license=LICENSE,
        citation="Wisconsin Department of Administration. State and County Household Projections, 2020-2050.",
        method_notes=(
            "All-age household totals are loaded at 2020, 2030, 2040, and 2050 for the state and all 72 counties. The workbook does not state a publication edition date. "
            "2020 is the Census-based estimate; later points are projections. Household projections are anchored to DOA's county population projections and use age-specific household rates. "
            "County sums differ from the published state total by at most two households because of rounding; published state values are retained."
        ),
    )
    n = insert_observations(conn, rows)
    conn.commit()
    print(f"WI load: {n} observations")


FIXTURES = [
    ("population", "state", "55", 2020, 5893718, "estimate", POP_SOURCE_ID),
    ("population", "state", "55", 2050, 5710120, "projection", POP_SOURCE_ID),
    ("population", "county", "55009", 2040, 293080, "projection", POP_SOURCE_ID),  # Brown
    ("population", "county", "55025", 2050, 779265, "projection", POP_SOURCE_ID),  # Dane
    ("households", "state", "55", 2020, 2428361, "estimate", HH_SOURCE_ID),
    ("households", "state", "55", 2050, 2468672, "projection", HH_SOURCE_ID),
    ("households", "county", "55025", 2050, 335030, "projection", HH_SOURCE_ID),  # Dane
]


def verify(conn) -> None:
    for variable, level, gid, year, value, datatype, source_id in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype, source_id FROM observations WHERE state='WI' AND geo_level=? AND geo_id=? "
            "AND variable=? AND year=?", (level, gid, variable, year)
        ).fetchone()
        assert row is not None, f"WI verify: missing {variable} {level} {gid} {year}"
        assert (int(row["value"]), row["datatype"], row["source_id"]) == (value, datatype, source_id)
    for variable, source in (("population", POP_SOURCE_ID), ("households", HH_SOURCE_ID)):
        shape = conn.execute(
            "SELECT COUNT(*) n, COUNT(DISTINCT geo_id) ng, COUNT(DISTINCT year) ny, MIN(year) lo, MAX(year) hi "
            "FROM observations WHERE state='WI' AND variable=? AND source_id=?", (variable, source)
        ).fetchone()
        assert tuple(shape) == (292, 73, 4, 2020, 2050)
    # Population county totals are controlled to the state exactly. Household
    # workbook documentation permits small rounding differences (<= 2).
    bad = conn.execute(
        "SELECT s.year, s.value - SUM(c.value) AS diff FROM observations s JOIN observations c "
        "ON c.state='WI' AND c.geo_level='county' AND c.variable='population' AND c.year=s.year "
        "WHERE s.state='WI' AND s.geo_level='state' AND s.variable='population' GROUP BY s.year "
        "HAVING ABS(diff) > 0"
    ).fetchall()
    assert not bad, f"WI verify: population county sums differ: {[tuple(r) for r in bad]}"
    hh_bad = conn.execute(
        "SELECT s.year, s.value - SUM(c.value) AS diff FROM observations s JOIN observations c "
        "ON c.state='WI' AND c.geo_level='county' AND c.variable='households' AND c.year=s.year "
        "WHERE s.state='WI' AND s.geo_level='state' AND s.variable='households' GROUP BY s.year "
        "HAVING ABS(diff) > 2"
    ).fetchall()
    assert not hh_bad, f"WI verify: household county sums differ by >2: {[tuple(r) for r in hh_bad]}"
    print(f"WI verify: {len(FIXTURES)} fixtures + population/household shapes and sums ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.wi --acquire")
