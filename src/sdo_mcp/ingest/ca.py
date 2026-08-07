"""California adapter.

Sources (all Department of Finance; verified 2026-08-07, plan/RESEARCH.md):
  - Population: P-series, Baseline 2024, Vintage 2026 (Demographic Research
    Unit). Totals ingested from P-2A (state + 58 counties, 2020-2070); the
    committed P3_Complete.zip microdata snapshot is the cross-check gate.
  - Components: P-CC (same vintage), one row per July 1 -> July 1 period.
  - Households: P-4, Vintage 2025 (lags the population series), 2010-2040.
    Explicitly labeled "2010-2025 Estimates and 2026-2040 Projections".
  - Employment: LRCF — Long-Run County Level Forecasts (Economic Research
    Unit + Caltrans, Jan 2026): farm + total nonfarm jobs 2000-2055. In scope
    because ERU sits in the same department as the state demographer.

Estimate/projection boundaries (recorded per product, per the domain rule):
  - P-series/P-CC: 2025/2026. The Vintage 2026 release note says history was
    "updated to reflect latest published July 1 population estimates" — the
    latest published July series is 2025 — and P-4, from the same unit and
    cycle, states the 2025/2026 split explicitly.
  - P-4: 2025/2026 (stated in the file title).
  - LRCF: 2024/2025 (the file's own note: forecasts start 2025).

California quirks handled here:
  - DOF overwrites files in place on vintage refresh (stable filenames,
    mutable content) — snapshots + retrieval dates are the record
  - legacy URLs 301 to an HTML page with HTTP 200; bodies are validated
  - naked fetchers get 403 (fetch.py sends a User-Agent)
  - the P-CC state sheet carries a bogus FIPS (6117); rows are keyed by the
    Geography column instead
  - LRCF sheets are named by county with a trailing note row under the data
"""

from __future__ import annotations

import sys

from ..common import fetch
from ..store import REPO_ROOT, upsert_source, insert_observations

RAW_DIR = REPO_ROOT / "data" / "raw" / "ca"
P3_ZIP = REPO_ROOT / "data" / "raw" / "CA_P3_Complete.zip"  # committed 2026-08-07
STATE_FIPS = "06"

DEMOG = "https://dof.ca.gov/media/docs/forecasting/Demographics/projections/"
LRCF_URL = (
    "https://dof.ca.gov/media/docs/forecasting/economics/"
    "economic-forecasts-u-s-california-counties/long-run-county-level-forecasts/"
    "Master-File-Output-Public.xlsx"
)
FILES = {
    "P2A_County_Total.xlsx": DEMOG + "P2A_County_Total.xlsx",
    "P4_HHProjections_B2024.xlsx": DEMOG + "P4_HHProjections_B2024.xlsx",
    "PCC_Components_of_Change.xlsx": DEMOG + "PCC_Components_of_Change.xlsx",
    "LRCF_Master-File-Output-Public.xlsx": LRCF_URL,
}

POP_BOUNDARY = 2025   # last July 1 estimate in the Vintage 2026 P-series
HH_BOUNDARY = 2025    # stated in the P-4 title
JOBS_BOUNDARY = 2024  # stated in the LRCF note (forecasts start 2025)

DRU = "California Department of Finance, Demographic Research Unit"
ERU = (
    "California Department of Finance, Economic Research Unit (with Caltrans); "
    "same department as the state demographer"
)
LICENSE = (
    "Standard California conditions of use; public data. Cite baseline + "
    "vintage as embedded in the files."
)


def acquire() -> None:
    """Snapshot the four DOF files into data/raw/ca/. The P3 microdata zip is
    already committed at data/raw/CA_P3_Complete.zip and is not re-fetched."""
    from datetime import date

    RAW_DIR.mkdir(parents=True, exist_ok=True)
    entries = []
    for name, url in FILES.items():
        body = fetch.get_bytes(url)
        if not body.startswith(b"PK"):
            raise RuntimeError(f"{url}: body is not an xlsx (200-with-HTML trap)")
        (RAW_DIR / name).write_bytes(body)
        entries.append({
            "file": name, "source": url,
            "retrieved_at": date.today().isoformat(),
            "sha256": fetch.sha256(body), "records": None, "note": "",
        })
        print(f"CA acquire: {name} ({len(body):,} bytes)")
    fetch.write_manifest(RAW_DIR, entries)


def _wb(name: str):
    import openpyxl

    return openpyxl.load_workbook(RAW_DIR / name, read_only=True)


def _retrieved(name: str) -> str:
    entry = next(e for e in fetch.read_manifest(RAW_DIR) if e["file"] == name)
    return entry["retrieved_at"]


def _geo(fips) -> tuple[str, str]:
    f = int(fips)
    if f == 6000:
        return "state", STATE_FIPS
    return "county", f"{f:05d}"


def _melt_wide(rows: list[tuple], variable: str, boundary: int,
               source_id: str, year_col0: int) -> list[dict]:
    """P-2A / P-4 shape: header row index 1, FIPS + name + wide year columns
    starting at `year_col0`, state row FIPS 6000."""
    years = [int(c) for c in rows[1][year_col0:] if c is not None]
    out = []
    for rec in rows[2:]:
        if rec[0] is None or not str(rec[0]).isdigit():
            continue
        level, gid = _geo(rec[0])
        for year, val in zip(years, rec[year_col0:]):
            out.append(dict(state="CA", geo_level=level, geo_id=gid, year=year,
                            variable=variable, value=int(val),
                            datatype="estimate" if year <= boundary else "projection",
                            source_id=source_id))
    return out


def load(conn) -> None:
    """Snapshots -> observations + sources. No network."""
    rows: list[dict] = []

    # Population totals (P-2A: FIPS, Geography, Age, Sex, Race, 2020...2070).
    wb = _wb("P2A_County_Total.xlsx")
    rows += _melt_wide(list(wb["Data"].iter_rows(values_only=True)),
                       "population", POP_BOUNDARY, "ca-pop-2026", 5)
    wb.close()

    # Households (P-4: FIPS, County, 2010...2040).
    wb = _wb("P4_HHProjections_B2024.xlsx")
    rows += _melt_wide(list(wb["Household Projections"].iter_rows(values_only=True)),
                       "households", HH_BOUNDARY, "ca-hh-2025", 2)
    wb.close()

    # Components of change (P-CC): one row per July 1 -> July 1 period,
    # recorded under the period's END year (the year the change lands in).
    wb = _wb("PCC_Components_of_Change.xlsx")
    for sheet, is_state in (("California Components of Change", True),
                            ("Counties Components of Change", False)):
        for rec in wb[sheet].iter_rows(values_only=True):
            if rec[0] is None or not str(rec[0]).isdigit():
                continue
            if is_state:
                # The state sheet's FIPS column is bogus (6117); key on the
                # Geography column instead.
                if str(rec[1]).strip().upper() != "CALIFORNIA":
                    continue
                level, gid = "state", STATE_FIPS
            else:
                level, gid = _geo(rec[0])
            year = int(rec[7])  # End Year
            datatype = "estimate" if year <= POP_BOUNDARY else "projection"
            for var, col in (("births", 4), ("deaths", 5), ("net_migration", 6)):
                rows.append(dict(state="CA", geo_level=level, geo_id=gid,
                                 year=year, variable=var, value=int(rec[col]),
                                 datatype=datatype, source_id="ca-pop-2026"))
    wb.close()

    # Employment (LRCF): one sheet per county + "California"; total jobs =
    # the file's two published columns (farm + total nonfarm) added together —
    # an additive identity over the complete published breakdown, recorded as
    # such in the method notes.
    fips_by_name = {
        row["name"].removesuffix(" County").strip().lower(): row["geo_id"]
        for row in conn.execute(
            "SELECT geo_id, name FROM geographies WHERE state='CA' AND geo_level='county'"
        )
    }
    wb = _wb("LRCF_Master-File-Output-Public.xlsx")
    for sheet in wb.sheetnames:
        if sheet == "California":
            level, gid = "state", STATE_FIPS
        else:
            gid = fips_by_name.get(sheet.strip().lower())
            if gid is None:
                raise RuntimeError(f"CA load: LRCF sheet {sheet!r} has no Census FIPS match")
            level = "county"
        data = list(wb[sheet].iter_rows(values_only=True))
        header = [str(c) for c in data[0]]
        i_farm, i_nonfarm = header.index("Farm"), header.index("Total Nonfarm")
        for rec in data[1:]:
            if rec[0] is None or not str(rec[0]).isdigit():
                continue  # trailing note row
            year = int(rec[0])
            # Alpine's Farm column is "NA" (not reported) in every year; the
            # served value sums only the published components, and the gap is
            # recorded in the source's method notes.
            parts = [float(rec[i]) for i in (i_farm, i_nonfarm)
                     if str(rec[i]).upper() != "NA"]
            if not parts:
                continue
            rows.append(dict(state="CA", geo_level=level, geo_id=gid, year=year,
                             variable="employment", value=sum(parts),
                             datatype="estimate" if year <= JOBS_BOUNDARY else "projection",
                             source_id="ca-jobs-2026"))
    wb.close()

    upsert_source(
        conn, source_id="ca-pop-2026", state="CA", agency=DRU,
        product_name="P-series population projections (Baseline 2024, Vintage 2026) "
                     "with P-CC components of change",
        vintage="Baseline 2024, Vintage 2026 (released 2026-03-13)",
        base_year=POP_BOUNDARY, horizon_year=2070,
        source_url=DEMOG, retrieved_at=_retrieved("P2A_County_Total.xlsx"),
        license=LICENSE,
        citation=("California Department of Finance, Demographic Research Unit. "
                  "Baseline 2024 Population Projections, Vintage 2026 Release. "
                  "Sacramento: California. March 2026."),
        method_notes=(
            "Cohort-component person-level microsimulation; Lee-Carter mortality; "
            "July 1 reference dates. Estimate/projection boundary 2025/2026: the "
            "Vintage 2026 release incorporates published July 1 estimates through "
            "2025 (release note), with model years from 2026. Components are "
            "recorded under each period's end year. Ingested totals are "
            "cross-checked against the committed P3 microdata snapshot at "
            "verify time."
        ),
    )
    upsert_source(
        conn, source_id="ca-hh-2025", state="CA", agency=DRU,
        product_name="P-4 total household projections",
        vintage="Vintage 2025 (released 2025-11, lags the population series)",
        base_year=HH_BOUNDARY, horizon_year=2040,
        source_url=DEMOG + "P4_HHProjections_B2024.xlsx",
        retrieved_at=_retrieved("P4_HHProjections_B2024.xlsx"),
        license=LICENSE,
        citation=("California Department of Finance, Demographic Research Unit. "
                  "Report P-4: Total Households, California Counties, 2010-2040 "
                  "(Vintage 2025 Release). Sacramento: California. November 2025."),
        method_notes=(
            "January 1 reference dates (unlike the July 1 P-series). The file "
            "title states the boundary: 2010-2025 estimates, 2026-2040 "
            "projections. The LRCF product separately carries county household "
            "forecasts to 2055; the two disagree and are never merged — this "
            "server serves P-4 as the household series."
        ),
    )
    upsert_source(
        conn, source_id="ca-jobs-2026", state="CA", agency=ERU,
        product_name="LRCF — Long-Run County Level Forecasts (total jobs)",
        vintage="January 2026 (annual cadence)",
        base_year=JOBS_BOUNDARY, horizon_year=2055,
        source_url=LRCF_URL,
        retrieved_at=_retrieved("LRCF_Master-File-Output-Public.xlsx"),
        license=LICENSE,
        citation=("California Department of Finance, Economic Research Unit, with "
                  "the California Department of Transportation. Long-Run County "
                  "Level Forecasts. Sacramento: California. January 2026."),
        method_notes=(
            "The served value is the sum of the file's two published employment "
            "columns (Farm + Total Nonfarm) — the complete published breakdown, "
            "no industry detail. Exception: Alpine County's Farm column is NA "
            "(not reported) in every year, so its series is Total Nonfarm only. "
            "Forecast years start 2025 (the file's own note). Values are model "
            "output and non-integer. EDD occupational projections are a "
            "different product, out of scope."
        ),
    )
    n = insert_observations(conn, rows)
    conn.commit()
    print(f"CA load: {n} observations")


# Known published values read off the snapshot files (recorded 2026-08-07/
# acquisition day). LRCF floats are matched to 6 decimal places.
FIXTURES = [
    # (geo_level, geo_id, variable, year, value, datatype)
    ("state",  "06",    "population", 2025, 39528899, "estimate"),
    ("state",  "06",    "population", 2026, 39567317, "projection"),
    ("state",  "06",    "population", 2070, 39853687, "projection"),
    ("county", "06037", "population", 2025, 9853435,  "estimate"),
    ("county", "06037", "population", 2026, 9811358,  "projection"),
    ("state",  "06",    "households", 2025, 13993318, "estimate"),
    ("state",  "06",    "households", 2026, 14089261, "projection"),
    ("county", "06001", "households", 2040, 681258,   "projection"),
    ("state",  "06",    "employment", 2024, 415600.488238705 + 17934705.3830389, "estimate"),
    ("state",  "06",    "employment", 2025, 420919.372466538 + 17956935.3964827, "projection"),
    ("county", "06001", "employment", 2055, 730.964983064773 + 864885.192079286, "projection"),
]

COMPONENT_FIXTURES = [
    # (geo_id, end_year, births, deaths, net_migration, datatype)
    ("06001", 2021, 16738, 11595, -27689, "estimate"),
    ("06001", 2070, 11205, 21276, 8840, "projection"),
]


def _p3_county_sums(years: set[int]) -> dict[tuple[str, int], int]:
    """Sum the committed P3 microdata (fips,year,sex,race7,agerc,perwt) to
    county totals for the given years."""
    import csv
    import io
    import zipfile

    sums: dict[tuple[str, int], int] = {}
    with zipfile.ZipFile(P3_ZIP) as z, z.open("P3_Complete.csv") as f:
        reader = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig"))
        next(reader)
        for fips, year, _sex, _race, _age, perwt in reader:
            y = int(year)
            if y in years:
                key = (f"{int(fips):05d}", y)
                sums[key] = sums.get(key, 0) + int(perwt)
    return sums


def verify(conn) -> None:
    for level, gid, var, year, value, datatype in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype FROM observations WHERE state='CA' AND"
            " geo_level=? AND geo_id=? AND variable=? AND year=? AND scenario=''"
            " AND age_group='all' AND sex='all' AND race='all'",
            (level, gid, var, year),
        ).fetchone()
        assert row is not None, f"CA verify: missing {level} {gid} {var} {year}"
        assert abs(row["value"] - value) < 1e-6 and row["datatype"] == datatype, (
            f"CA verify: {level} {gid} {var} {year} -> "
            f"{row['value']}/{row['datatype']}, expected {value}/{datatype}"
        )

    for gid, year, births, deaths, netmig, datatype in COMPONENT_FIXTURES:
        got = {
            r["variable"]: (int(r["value"]), r["datatype"])
            for r in conn.execute(
                "SELECT variable, value, datatype FROM observations WHERE"
                " state='CA' AND geo_id=? AND year=? AND variable IN"
                " ('births','deaths','net_migration')", (gid, year))
        }
        expected = {"births": (births, datatype), "deaths": (deaths, datatype),
                    "net_migration": (netmig, datatype)}
        assert got == expected, f"CA verify: components {gid} {year}: {got}"

    spans = {
        ("population", "county", 58): (2020, 2070),
        ("households", "county", 58): (2010, 2040),
        ("employment", "county", 58): (2000, 2055),
    }
    for (var, level, n_geos), (y0, y1) in spans.items():
        row = conn.execute(
            "SELECT MIN(year) lo, MAX(year) hi, COUNT(DISTINCT geo_id) n"
            " FROM observations WHERE state='CA' AND geo_level=? AND variable=?",
            (level, var),
        ).fetchone()
        assert (row["lo"], row["hi"]) == (y0, y1), f"CA verify: {var} span {row['lo']}-{row['hi']}"
        assert row["n"] == n_geos, f"CA verify: {var} has {row['n']} counties"

    # Reproducibility gate: the served totals must be reproducible from the
    # committed P3 microdata snapshot — every county, two years (one estimate
    # year, one far projection year), exact match.
    p3 = _p3_county_sums({2025, 2070})
    mismatches = []
    for (gid, year), total in p3.items():
        row = conn.execute(
            "SELECT value FROM observations WHERE state='CA' AND geo_level='county'"
            " AND geo_id=? AND year=? AND variable='population'", (gid, year),
        ).fetchone()
        if row is None or int(row["value"]) != total:
            mismatches.append((gid, year, total, row and int(row["value"])))
    assert len(p3) == 116, f"CA verify: P3 sum produced {len(p3)} county-years"
    assert not mismatches, f"CA verify: P3 microdata mismatch: {mismatches[:5]}"

    print(f"CA verify: {len(FIXTURES)} fixtures + components + spans + "
          f"P3 microdata cross-check (116 county-years) ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.ca --acquire")
