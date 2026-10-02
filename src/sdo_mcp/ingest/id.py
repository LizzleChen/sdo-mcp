"""Idaho's revised August 2025 population totals, state and all 44 counties."""
from __future__ import annotations

import io
import math
import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "id"
CATALOG_URL = "https://lmi.idaho.gov/data-tools/population-projections/"
PREFIX = "https://lmi.idaho.gov/wp-content/uploads/2025/09/"
CATALOG_FILE = "population_catalog_2026-10-02.html"
METHOD_FILE = "population_forecast_methodology_2017.pdf"
JOBS_FILE = "industry_occupation_catalog_2026-10-02.html"
STATE_FILE = "Population_Statewide_revised2025-08-07.xlsx"
REGIONS = {
    "Northern": ("1. North", ("Benewah", "Bonner", "Boundary", "Kootenai", "Shoshone")),
    "North-central": ("2. North Central", ("Clearwater", "Idaho", "Latah", "Lewis", "Nez Perce")),
    "Southwestern": ("3. Southwest", ("Ada", "Adams", "Boise", "Canyon", "Elmore", "Gem", "Owyhee", "Payette", "Valley", "Washington")),
    "South-central": ("4. South Central", ("Blaine", "Camas", "Cassia", "Gooding", "Jerome", "Lincoln", "Minidoka", "Twin Falls")),
    "Southeastern": ("5. Southeast", ("Bannock", "Bear Lake", "Bingham", "Caribou", "Franklin", "Oneida", "Power")),
    "Eastern": ("6. East", ("Bonneville", "Butte", "Clark", "Custer", "Fremont", "Jefferson", "Lemhi", "Madison", "Teton")),
}
FILES = {STATE_FILE: PREFIX + "Population_Statewide.xlsx"}
FILES.update({f"Population_{region}_revised2025-08-07.xlsx": PREFIX + f"Population_{region}.xlsx" for region in REGIONS})
FILES.update({CATALOG_FILE: CATALOG_URL,
    METHOD_FILE: "https://lmi.idaho.gov/wp-content/uploads/publications/2017/PopulationProjections/Population-Forecast-Methodology.pdf?ver=2017-05-31-171419-813",
    JOBS_FILE: "https://lmi.idaho.gov/data-tools/occupational-industry-projections/"})
SOURCE_ID = "id-pop-revised2025-state"
COUNTY_SOURCE_ID = "id-pop-revised2025-counties"
AGENCY = "Idaho Department of Labor, Labor Market Information"
VINTAGE = "August 7, 2025 (revised)"
YEARS = tuple(range(2024, 2035))
STATE_VALUES = (2047050, 2080283, 2113206, 2145803, 2178059, 2209994, 2241600, 2272856, 2303768, 2334322, 2364544)
AGE_LABELS = ("Under 5 years", *(f"{age} to {age + 4} years" for age in range(5, 85, 5)), "85 years and over")


def _geographies() -> dict[str, str]:
    result = {p[4].removesuffix(" County"): p[1] + p[2]
              for line in COUNTY_FILE.read_text().splitlines()[1:] if (p := line.split("|"))[0] == "ID"}
    if len(result) != 44 or set(result) != {n for _, names in REGIONS.values() for n in names}:
        raise ValueError("ID: unexpected Census county roster")
    return result


def _integer(value) -> int:
    if (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)
            or value < 0 or not float(value).is_integer()):
        raise ValueError("ID: invalid or suppressed integer population")
    return int(value)


def _sheet_totals(sheet) -> tuple[int, ...]:
    header = ("Age Group", *YEARS, None, "Net Change", "Growth %", "Annualized")
    if (tuple(sheet.cell(1, c).value for c in range(1, 17)) != header
            or any(sheet.cell(1, c).value is not None for c in range(17, sheet.max_column + 1))):
        raise ValueError("ID: unexpected workbook header or year grid")
    labels = tuple(sheet.cell(r, 1).value for r in range(2, 30))
    expected = (*AGE_LABELS, None, "Total", "Under 15", "15 to 64", "65 and Older", None, "Men ", "Women", None, "Labor Force")
    if labels != expected:
        raise ValueError("ID: missing, duplicate or changed demographic row labels")
    if any(sheet.cell(r, 1).value is not None for r in range(30, sheet.max_row + 1)):
        raise ValueError("ID: unexpected trailing table rows")
    totals = tuple(_integer(sheet.cell(21, c).value) for c in range(2, 13))
    for c, total in enumerate(totals, 2):
        ages = [_integer(sheet.cell(r, c).value) for r in range(2, 20)]
        sex = [_integer(sheet.cell(r, c).value) for r in (26, 27)]
        broad = [_integer(sheet.cell(r, c).value) for r in (22, 23, 24)]
        if sum(ages) != total or sum(sex) != total or sum(broad) != total:
            raise ValueError("ID: published age/sex totals disagree")
    return totals


def parse_workbook(body: bytes, expected_sheets: tuple[str, ...]) -> dict[str, tuple[int, ...]]:
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(body), data_only=True)
    except Exception as exc:
        raise ValueError("ID: response is not a readable XLSX workbook") from exc
    if tuple(book.sheetnames) != expected_sheets:
        raise ValueError("ID: unexpected, missing or duplicate workbook geography sheets")
    return {sheet.title: _sheet_totals(sheet) for sheet in book}


def parse(bodies: dict[str, bytes]) -> list[dict]:
    state = parse_workbook(bodies[STATE_FILE], ("Total", *(v[0] for v in REGIONS.values())))
    if state["Total"] != STATE_VALUES:
        raise ValueError("ID: independently published state totals changed")
    counties = {}
    for region, (state_sheet, names) in REGIONS.items():
        filename = f"Population_{region}_revised2025-08-07.xlsx"
        regional = parse_workbook(bodies[filename], ("Total", *names))
        if regional["Total"] != state[state_sheet]:
            raise ValueError("ID: region totals differ across workbooks")
        for i in range(len(YEARS)):
            if sum(regional[n][i] for n in names) != regional["Total"][i]:
                raise ValueError("ID: county/region totals disagree")
        counties.update({n: regional[n] for n in names})
    if set(counties) != set(_geographies()):
        raise ValueError("ID: incomplete county grid")
    if any(sum(v[i] for v in counties.values()) != state["Total"][i] for i in range(len(YEARS))):
        raise ValueError("ID: county/state totals disagree")
    geographies = _geographies() | {"Idaho State": "16"}
    rows = []
    for name, values in {"Idaho State": state["Total"], **counties}.items():
        level = "state" if name == "Idaho State" else "county"
        rows.extend(dict(state="ID", geo_level=level, geo_id=geographies[name], year=year,
                         variable="population", value=value, datatype="projection",
                         source_id=SOURCE_ID if level == "state" else COUNTY_SOURCE_ID)
                    for year, value in zip(YEARS, values, strict=True))
    return rows


def _validate(bodies: dict[str, bytes]) -> list[dict]:
    rows = parse(bodies)
    catalog = " ".join(bodies[CATALOG_FILE].decode("utf-8", "replace").split())
    if any(s not in catalog for s in ("<html", "2,364,544", "2024-2034", "Revised Aug. 7, 2025", "American Community Survey")):
        raise ValueError("ID: invalid publisher catalog, edition or ACS evidence")
    if any(FILES[name] not in catalog for name in FILES if name.endswith(".xlsx")):
        raise ValueError("ID: publisher catalog omits selected workbook")
    fetch.validate_pdf(bodies[METHOD_FILE], "ID methods")
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(bodies[METHOD_FILE]))
    text = " ".join((reader.pages[0].extract_text() or "").split())
    if len(reader.pages) != 1 or any(s not in text for s in ("Idaho Department of Labor", "each of Idaho’s counties", "smaller the population", "Migration")):
        raise ValueError("ID: invalid producer, county or method evidence")
    jobs = bodies[JOBS_FILE].decode("utf-8", "replace").lower()
    if "<html" not in jobs or "industries and occupations" not in jobs or "long-term employment projections" not in jobs:
        raise ValueError("ID: invalid occupational/industry product evidence")
    return rows


def acquire() -> None:
    if (RAW_DIR / "manifest.json").exists():
        entries = fetch.read_manifest(RAW_DIR)
        if {e["file"] for e in entries} != set(FILES):
            raise ValueError("ID: incomplete snapshot manifest")
        _validate({name: (RAW_DIR / name).read_bytes() for name in FILES})
        print("ID acquire: reused validated August 2025 snapshots")
        return
    bodies = {name: fetch.get_bytes(url) for name, url in FILES.items()}
    _validate(bodies)
    selected_counts = {STATE_FILE: 11} | {f"Population_{region}_revised2025-08-07.xlsx": 11 * len(names) for region, (_, names) in REGIONS.items()}
    entries = []
    for name, body in bodies.items():
        fetch.write_snapshot(RAW_DIR / name, body)
        entries.append(dict(file=name, source=FILES[name], retrieved_at=date.today().isoformat(),
                            sha256=fetch.sha256(body), records=selected_counts.get(name),
                            note="Original publisher bytes. August 7, 2025 catalog revision; workbook upload path is not vintage. Population Total rows selected; labor force is not jobs."))
    fetch.write_manifest(RAW_DIR, entries)


def load(conn) -> None:
    entries = {e["file"]: e for e in fetch.read_manifest(RAW_DIR)}
    rows = _validate({name: (RAW_DIR / name).read_bytes() for name in FILES})
    if {(r["geo_level"], r["geo_id"]) for r in conn.execute("SELECT * FROM geographies WHERE state='ID'")} != {(r["geo_level"], r["geo_id"]) for r in rows}:
        raise ValueError("ID: seeded geographies differ from publication")
    for level, source_id in (("state", SOURCE_ID), ("county", COUNTY_SOURCE_ID)):
        upsert_source(conn, source_id=source_id, state="ID", agency=AGENCY,
            product_name=f"Ten-Year Population Forecast 2024-2034: {level} Total rows",
            vintage=VINTAGE, base_year=None, horizon_year=2034,
            source_url=FILES[STATE_FILE] if level == "state" else CATALOG_URL,
            retrieved_at=(entries[STATE_FILE]["retrieved_at"] if level == "state" else
                max(entries[f"Population_{region}_revised2025-08-07.xlsx"]["retrieved_at"] for region in REGIONS)),
            license="No explicit data license identified; cite the Idaho Department of Labor.",
            citation="Idaho Department of Labor. Ten-Year Population Forecast 2024-2034, revised August 7, 2025. Original Statewide or six regional workbooks, Total row 21, B:L.",
            method_notes=("Explicit statewide and all 44 county Total rows publish integer persons at annual 2024-2034 points. "
                "The catalog labels the entire grid projections and identifies American Community Survey estimates as inputs; "
                "the last underlying estimate year and reference day are unstated, so base_year is null and every loaded point remains projection. "
                "The September 2025 URL directory and workbook creation dates are not used as edition dates. "
                "The currently linked 2017 general methods describe county/age cohort birth, death, aging and migration modeling; county is the smallest modeled geography. "
                "County, region and state Total rows reconcile exactly across all published years. No interpolation or balancing is applied. "
                "Age blocks, sex and six region series are archived but not served. Labor Force is a potential labor pool, not jobs. "
                "The separate industry/occupational projections are out of scope. Household-count, published component tables and eligible demographic jobs coverage remain unverified."))
    insert_observations(conn, rows)
    conn.commit()
    print(f"ID load: {len(rows)} observations")


# Independent direct original XML cell reads, plus the catalog's state horizon.
FIXTURES = (("16", 2024, 2047050), ("16", 2025, 2080283), ("16", 2034, 2364544),
            ("16001", 2024, 533654), ("16001", 2025, 542752), ("16001", 2034, 618118),
            ("16027", 2034, 320984), ("16033", 2024, 825), ("16033", 2025, 817), ("16033", 2034, 725),
            ("16025", 2024, 1289), ("16025", 2034, 1425))


def verify(conn) -> None:
    rows = conn.execute("SELECT * FROM observations WHERE state='ID'").fetchall()
    ids = set(_geographies().values()) | {"16"}
    assert len(rows) == 495 and {(r["geo_id"], r["year"]) for r in rows} == {(gid, y) for gid in ids for y in YEARS}
    assert all(r["geo_level"] == ("state" if r["geo_id"] == "16" else "county") and r["datatype"] == "projection"
               and r["source_id"] == (SOURCE_ID if r["geo_id"] == "16" else COUNTY_SOURCE_ID)
               and r["variable"] == "population" and r["scenario"] == "" and r["age_group"] == r["sex"] == r["race"] == "all"
               and r["value"] >= 0 and float(r["value"]).is_integer() for r in rows)
    values = {(r["geo_id"], r["year"]): r["value"] for r in rows}
    for gid, year, value in FIXTURES:
        assert values[gid, year] == value
    for year in YEARS:
        assert sum(values[gid, year] for gid in ids - {"16"}) == values["16", year]
    print("ID verify: state/all 44 counties, exact grids and original XML fixtures ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.id --acquire")
