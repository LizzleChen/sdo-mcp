"""Hawaii DBEDT April 2024 resident population and civilian total jobs."""

from __future__ import annotations

import io
import math
import re
import sys
from datetime import date, datetime

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source

RAW_DIR = REPO_ROOT / "data" / "raw" / "hi"
CATALOG_URL = "https://dbedt.hawaii.gov/economic/economic-forecast/long-range-projections/"
WORKBOOK_URL = "https://files.hawaii.gov/dbedt/economic/data_reports/LRF/2050-long-range-projections-appendices.xlsx"
REPORT_URL = "https://files.hawaii.gov/dbedt/economic/data_reports/LRF/2050-long-range-projections.pdf"
GEOGRAPHY_URL = "https://files.hawaii.gov/dbedt/economic/databook/2024-individual/01/012324.pdf"
BEA_URL = "https://www.bea.gov/sites/default/files/methodologies/BEA-Local-Area-Personal-Income-and-Employment-Concepts-and-Methods.pdf"
CATALOG_FILE = "projection_catalog_2026-09-30.html"
WORKBOOK_FILE = "population_economic_2050_appendices_apr2024.xlsx"
REPORT_FILE = "population_economic_2050_report_apr2024.pdf"
GEOGRAPHY_FILE = "dbedt_population_geography_table_1-23_2024.pdf"
BEA_FILE = "bea_local_income_employment_methods_2026-09-30.pdf"
FILES = {CATALOG_FILE: CATALOG_URL, WORKBOOK_FILE: WORKBOOK_URL, REPORT_FILE: REPORT_URL,
         GEOGRAPHY_FILE: GEOGRAPHY_URL, BEA_FILE: BEA_URL}
POP_SOURCE_ID = "hi-pop-apr2024"
JOBS_SOURCE_ID = "hi-jobs-apr2024"
AGENCY = "Hawaii Department of Business, Economic Development and Tourism, Research and Economic Analysis Division"
COMBINED_ID = "hi-maui-kalawao"
COMBINED_NAME = "Maui and Kalawao counties (combined)"
GEOGRAPHIES = (("15", "Hawaii State"), ("15001", "Hawaii County"), ("15003", "Honolulu County"),
               ("15007", "Kauai County"), (COMBINED_ID, "Maui County"))
POP_YEARS = (2020, 2030, 2040, 2050)
JOBS_YEARS = (2022, 2030, 2040, 2050)
POP_FOOTNOTE = "1) As of July 1st\n2) For 2020 figures, total population is from the 2022 vintage population estimates by the U.S. Census Bureau while the population by age group are DBEDT estimates \n3) Visitor arrivals to Hawaii were limited in 2020 due to pandemic related travel restrictions"
JOBS_FOOTNOTE = "1)  Actual figures from the U.S. Bureau of Economic Analysis (BEA) with missing values estimated by DBEDT\n *  Projected values were rounded to the nearest ten."


def _number(value, context: str, *, integer: bool = False):
    if (not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value)
            or value < 0 or (integer and int(value) != value)):
        raise ValueError(f"HI: invalid numeric total at {context}")
    return value


def parse_workbook(body: bytes) -> list[dict]:
    """Keep published model decimals; never substitute visitors or labor force."""
    import openpyxl
    try:
        book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("HI: response is not a readable XLSX workbook") from exc
    try:
        if (book.sheetnames != ["List of Appendix Tables"] + [f"Table A-{i}" for i in range(1, 61)]
                or book.properties.modified != datetime(2024, 4, 18, 1, 15, 52)
                or book["List of Appendix Tables"]["A1"].value != "DBEDT Long Range Projections 2050 Series: List of Appendix Tables"):
            raise ValueError("HI: unexpected workbook edition or sheets")
        rows, population, jobs = [], {}, {}
        for i, (gid, label) in enumerate(GEOGRAPHIES):
            sn = f"Table A-{i + 2}"
            grid = list(book[sn].values)
            if (len(grid) != 36 or grid[0][0] != f"{sn}.  {label} Population Projection, Selected Components, 2020-2050 1)"
                    or grid[1][1:5] != ("2020 2)", 2030, 2040, 2050)
                    or grid[3][0] != "Total resident population" or grid[16][0] != "De facto population 3)"
                    or grid[-1][0] != POP_FOOTNOTE):
                raise ValueError(f"HI: unexpected population years, labels or estimate boundary in {sn}")
            population[gid] = tuple(_number(v, f"{sn}/{year}") for year, v in zip(POP_YEARS, grid[3][1:5]))
            age_sheet = book[f"Table A-{7 + i * 3}"]
            if (age_sheet["A21"].value != "Total"
                    or tuple(age_sheet.cell(21, c).value for c in range(2, 6)) != population[gid]
                    or not age_sheet["A1"].value.startswith(f"Table A-{7 + i * 3}.  {label} Total Resident Population")):
                raise ValueError(f"HI: population summary and age-table totals disagree for {gid}")
            # Six exhaustive age groups independently reconcile to the publisher's total.
            for column, value in enumerate(population[gid], 1):
                age_sum = sum(_number(row[column], f"{sn}/age") for row in grid[9:15])
                if not math.isclose(age_sum, value, rel_tol=0, abs_tol=1e-7):
                    raise ValueError(f"HI: population age groups do not reconcile for {gid}")
            rows.extend({"state": "HI", "geo_level": "state" if gid == "15" else "county", "geo_id": gid,
                         "year": year, "variable": "population", "value": value,
                         "datatype": "estimate" if year == 2020 else "projection", "source_id": POP_SOURCE_ID}
                        for year, value in zip(POP_YEARS, population[gid]))
            sn = f"Table A-{38 + i}"
            grid = list(book[sn].values)
            punctuation = "" if i == 3 else "."
            if (len(grid) != 43 or grid[0][0] != f"{sn}{punctuation} {label} Civilian Total Jobs by Sector, 2022-2050"
                    or grid[1][1:5] != ("20221)", 2030, 2040, 2050) or grid[2][0] != "Total"
                    or grid[-1][0] != JOBS_FOOTNOTE):
                raise ValueError(f"HI: unexpected jobs years, units or estimate boundary in {sn}")
            jobs[gid] = tuple(_number(v, f"{sn}/{year}", integer=True) for year, v in zip(JOBS_YEARS, grid[2][1:5]))
            rows.extend({"state": "HI", "geo_level": "state" if gid == "15" else "county", "geo_id": gid,
                         "year": year, "variable": "employment", "value": value,
                         "datatype": "estimate" if year == 2022 else "projection", "source_id": JOBS_SOURCE_ID}
                        for year, value in zip(JOBS_YEARS, jobs[gid]))
        for i, year in enumerate(POP_YEARS):
            if not math.isclose(population["15"][i], sum(v[i] for g, v in population.items() if g != "15"), rel_tol=0, abs_tol=1e-7):
                raise ValueError(f"HI: county-group population totals do not reconcile in {year}")
        for i, year in enumerate(JOBS_YEARS):
            # Source county totals are independently rounded to tens; preserve its discrepancy.
            if abs(jobs["15"][i] - sum(v[i] for g, v in jobs.items() if g != "15")) > 20:
                raise ValueError(f"HI: county-group jobs exceed publication rounding in {year}")
        return rows
    finally:
        book.close()


def parse_report(body: bytes) -> dict[tuple[str, int, str], int]:
    """Independent PDF population/jobs totals; report rounds model totals to tens."""
    from pypdf import PdfReader
    try:
        report = PdfReader(io.BytesIO(body))
        if len(report.pages) != 36 or "April 2024" not in report.pages[0].extract_text():
            raise ValueError("unexpected report edition")
        methodology = " ".join(report.pages[21].extract_text().split())
        if "cohort-component method with 2020 as the base year" not in methodology:
            raise ValueError("missing population base-year statement")
        job_methods = " ".join(report.pages[30].extract_text().split())
        if "military jobs were subtracted" not in job_methods or "Projections of jobs" not in job_methods:
            raise ValueError("missing civilian jobs definition")
        values = {}
        text = report.pages[5].extract_text()
        if "Table 1. Resident population by county: history and projections" not in text:
            raise ValueError("missing population PDF table")
        for line in text.splitlines():
            match = re.match(r"^(2020|2030|2040|2050)[12]?\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s*$", line)
            if match:
                year = int(match[1])
                for (gid, _), token in zip(GEOGRAPHIES, match.groups()[1:]):
                    key = (gid, year, "population")
                    if key in values:
                        raise ValueError("duplicate PDF population row")
                    values[key] = int(token.replace(",", ""))
        text = report.pages[17].extract_text()
        if "Table 8. Projection of civilian jobs, 2022-2050" not in text:
            raise ValueError("missing jobs PDF table")
        gid = None
        label_map = {"State Total": "15", "Hawaii County": "15001", "Honolulu County": "15003",
                     "Kauai County": "15007", "Maui County": COMBINED_ID}
        for line in text.splitlines():
            normalized = " ".join(line.split())
            if normalized in label_map:
                gid = label_map[normalized]
            match = re.match(r"^Civilian total jobs\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s*$", normalized)
            if match:
                if gid is None:
                    raise ValueError("jobs PDF geography missing")
                for year, token in zip(JOBS_YEARS, match.groups()):
                    key = (gid, year, "employment")
                    if key in values:
                        raise ValueError("duplicate jobs PDF row")
                    values[key] = int(token.replace(",", ""))
        if len(values) != 40:
            raise ValueError("incomplete PDF population/jobs grid")
        return values
    except Exception as exc:
        raise ValueError("HI: response is not the expected population/economic report PDF") from exc


def _validate_geographies(population_body: bytes, jobs_body: bytes) -> None:
    from pypdf import PdfReader
    try:
        proof = PdfReader(io.BytesIO(population_body))
        if len(proof.pages) != 1 or "Maui County including Kalawao County" not in " ".join(proof.pages[0].extract_text().split()):
            raise ValueError("missing DBEDT combined population geography")
        bea = PdfReader(io.BytesIO(jobs_body))
        # Target the published geography page, avoiding optional font decoding elsewhere.
        if len(bea.pages) != 122 or "Kalawao County, Hawaii is combined with Maui County" not in " ".join(bea.pages[22].extract_text().split()):
            raise ValueError("missing BEA combined jobs geography")
    except Exception as exc:
        raise ValueError("HI: invalid combined-county geography evidence PDF") from exc


def _validate_all(bodies: dict[str, bytes]) -> list[dict]:
    catalog = " ".join(bodies[CATALOG_FILE].decode("utf-8", "replace").lower().split())
    if any(marker not in catalog for marker in ("<html", "2050 series", "four counties", "2050-long-range-projections-appendices.xlsx")):
        raise ValueError("HI: invalid publisher catalog")
    _validate_geographies(bodies[GEOGRAPHY_FILE], bodies[BEA_FILE])
    rows = parse_workbook(bodies[WORKBOOK_FILE])
    checks = parse_report(bodies[REPORT_FILE])
    for row in rows:
        key = (row["geo_id"], row["year"], row["variable"])
        reference = checks[key]
        if key == (COMBINED_ID, 2030, "population"):
            # Publisher conflict: A-6 = 173620.3066, PDF Table 1 = 173520.
            # It is recorded and preserved; no silent repair or arbitrary wide tolerance.
            valid = math.isclose(row["value"], 173620.30660481827, rel_tol=0, abs_tol=1e-7) and reference == 173520
        elif row["variable"] == "population":
            valid = abs(row["value"] - reference) <= 10  # PDF independently rounds to tens.
        elif row["year"] == 2022:
            valid = abs(row["value"] - reference) <= 5  # Appendix also rounds historical county jobs.
        else:
            valid = row["value"] == reference
        if not valid:
            raise ValueError(f"HI: workbook and independent PDF disagree at {key}")
    return rows


def acquire() -> None:
    if (RAW_DIR / "manifest.json").exists():
        if {e["file"] for e in fetch.read_manifest(RAW_DIR)} != set(FILES):
            raise ValueError("HI: incomplete snapshot manifest")
        _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
        print("HI acquire: reused validated snapshots")
        return
    bodies = {name: fetch.get_bytes(url) for name, url in FILES.items()}
    rows = _validate_all(bodies)
    entries = []
    for name, url in FILES.items():
        fetch.write_snapshot(RAW_DIR / name, bodies[name])
        entries.append({"file": name, "source": url, "retrieved_at": date.today().isoformat(),
                        "sha256": fetch.sha256(bodies[name]), "records": len(rows) if name == WORKBOOK_FILE else None,
                        "note": "Original DBEDT April 2024 workbook and report; DBEDT/BEA evidence confirms combined Maui/Kalawao geography. Workbook model precision retained; Maui 2030 differs from report by about 100 people."})
    fetch.write_manifest(RAW_DIR, entries)


def load(conn) -> None:
    manifest = {e["file"]: e for e in fetch.read_manifest(RAW_DIR)}
    rows = _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
    conn.execute("INSERT OR REPLACE INTO geographies (geo_id,geo_level,state,name) VALUES (?,?,?,?)",
                 (COMBINED_ID, "county", "HI", COMBINED_NAME))
    for source_id, variable, base_year, description in (
        (POP_SOURCE_ID, "population", 2020, "resident population, including military and dependents but excluding visitors"),
        (JOBS_SOURCE_ID, "employment", 2022, "civilian total jobs: wage/salary plus proprietors' jobs, excluding military jobs"),
    ):
        upsert_source(conn, source_id=source_id, state="HI", agency=AGENCY,
                      product_name=f"DBEDT 2050 Series {description}", vintage="April 2024", base_year=base_year,
                      horizon_year=2050, source_url=CATALOG_URL, retrieved_at=manifest[WORKBOOK_FILE]["retrieved_at"],
                      license="No explicit data license identified; cite Hawaii DBEDT Research and Economic Analysis.",
                      citation=f"{AGENCY}. Population and Economic Projections for the State of Hawaii to 2050. April 2024.",
                      method_notes=(
                          "Original appendix workbook is the numeric authority; its cached model decimals are preserved. "
                          "Resident population includes military personnel and dependents and excludes visitors. "
                          "Population points are July 1 2020 estimates and 2030/2040/2050 projections. Methodology explicitly "
                          "launches the cohort-component population model in 2020; Census Vintage 2022 identifies the estimate "
                          "edition, not a 2022 population base. Jobs launch from 2022 BEA/DBEDT estimates and are projected for "
                          "2030/2040/2050, rounded to tens in the appendix; they count civilian jobs, including proprietors, "
                          "not employed residents or labor force. Military jobs are excluded. DBEDT publishes the state and "
                          "four county groups; Maui includes Kalawao. The combined series uses hi-maui-kalawao, never Maui "
                          "15009 or Kalawao 15005 individually. DBEDT Data Book Table 1.23 confirms population aggregation; "
                          "BEA's local employment methodology confirms the same county combination. Workbook county "
                          "population totals reconcile to its published state total. Jobs county sums can differ by ten "
                          "from the state because of independent rounding. The workbook's combined Maui/Kalawao 2030 "
                          "population is 173620.30660481827; PDF Table 1 and Data Book Table 1.23 print 173520, a material "
                          "publisher discrepancy retained rather than corrected. Other independent PDF totals agree within "
                          "their documented rounding. De facto population, age/sex, labor force and economic sector detail "
                          "are archived but not served as population/jobs totals."
                      ))
    insert_observations(conn, rows)
    conn.commit()
    print(f"HI load: {len(rows)} observations")


def verify(conn) -> None:
    shape = conn.execute("SELECT COUNT(*),COUNT(DISTINCT geo_id),MIN(year),MAX(year) FROM observations WHERE state='HI'").fetchone()
    assert tuple(shape) == (40, 5, 2020, 2050)
    for gid, year, variable, expected in (("15", 2050, "population", 1560885.7897387256),
                                         (COMBINED_ID, 2030, "population", 173620.30660481827),
                                         ("15", 2022, "employment", 853100), ("15", 2050, "employment", 979350),
                                         ("15001", 2050, "employment", 133520), (COMBINED_ID, 2050, "employment", 129980)):
        row = conn.execute("SELECT value,datatype FROM observations WHERE state='HI' AND geo_id=? AND year=? AND variable=?",
                           (gid, year, variable)).fetchone()
        assert tuple(row) == (expected, "estimate" if year in (2020, 2022) else "projection")
    assert conn.execute("SELECT COUNT(*) FROM observations WHERE state='HI' AND geo_id IN ('15005','15009')").fetchone()[0] == 0
    print("HI verify: resident population/civilian jobs and explicit combined county geography ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.hi --acquire")
