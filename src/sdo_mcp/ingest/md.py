"""Maryland Planning state and county population, household and jobs totals.

Only the first (level) table is served. Regional rows, historical years before
2020, and subsequent change/rate tables remain in the immutable originals.
"""

from __future__ import annotations

import io
import math
import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "md"
CATALOG_FILE = "projection_catalog_2026-09-29.html"
CATALOG_URL = "https://planning.maryland.gov/MSDC/Pages/s3_projection.aspx"
METHOD_FILE = "population_methodology_2026-09-29.pdf"
METHOD_URL = "https://planning.maryland.gov/MSDC/Documents/popproj/PopProjMethod-text.pdf"
AGENCY = "Maryland Department of Planning, State Data and Analysis Center"
LICENSE = "No explicit data license identified; cite the Maryland Department of Planning."
BASE_YEAR = 2020
PRODUCTS = {
    "population": {
        "file": "population_2025-11.xlsx", "sheet": "DataPop",
        "url": "https://planning.maryland.gov/MSDC/Documents/popproj/TotalPopProj.xlsx",
        "source_id": "md-pop-2025-11", "vintage": "November 2025",
        "title": "Historical and Projected Total Population for Maryland's Jurisdictions",
        "horizon": 2055,
    },
    "households": {
        "file": "households_2025-11.xlsx", "sheet": "Households",
        "url": "https://planning.maryland.gov/MSDC/Documents/popproj/HouseholdProj.xlsx",
        "source_id": "md-households-2025-11", "vintage": "November 2025",
        "title": "Historical and Projected Households for Maryland's Jurisdictions",
        "horizon": 2055,
    },
    "employment": {
        "file": "jobs_2022-10.xlsx", "sheet": "Summary",
        "url": "https://planning.maryland.gov/MSDC/Documents/projection/Jobs/Summary_TotalJobs.xlsx",
        "source_id": "md-jobs-2022-10", "vintage": "October 2022 (preliminary)",
        "title": "Historical and Projected Total Jobs by Place of Work for Maryland's Jurisdictions",
        "horizon": 2050,
    },
}
REGIONS = {
    "BALTIMORE REGION", "WASHINGTON SUBURBAN REGION", "SOUTHERN MARYLAND REGION",
    "WESTERN MARYLAND REGION", "UPPER EASTERN SHORE REGION", "LOWER EASTERN SHORE REGION",
}


def _county_names(jobs: bool) -> dict[str, str]:
    names = {}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        parts = line.split("|")
        if len(parts) >= 5 and parts[0] == "MD":
            name = parts[4]
            # The Census roster uses lowercase 'city'; this is county 24510,
            # never Baltimore County (24005) or a place-level observation.
            name = "Baltimore City" if name == "Baltimore city" else name
            if jobs:
                name = name.removesuffix(" County")
            names[name] = parts[1] + parts[2]
    if len(names) != 24:
        raise ValueError("MD: expected 24 Census county-equivalent jurisdictions")
    return names


def parse_workbook(body: bytes, variable: str) -> list[dict]:
    """Validate source semantics and the complete published totals grid."""
    import openpyxl

    spec = PRODUCTS[variable]
    jobs = variable == "employment"
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"MD: response is not a readable {variable} workbook") from exc
    try:
        if workbook.sheetnames != [spec["sheet"]]:
            raise ValueError(f"MD: unexpected {variable} workbook sheets")
        sheet = workbook.active
        title_cell, date_cell = ("C1", "C2") if jobs else ("A1", "A2")
        edition = "(Revisions October 2022)" if jobs else "(November 2025)"
        if (str(sheet[title_cell].value).strip() != spec["title"]
                or sheet[date_cell].value != edition):
            raise ValueError(f"MD: unexpected {variable} title or edition")
        years = ((1970, 1980, 1990, 2000, 2010, *range(2020, 2051, 5))
                 if jobs else (2000, 2010, *range(2020, 2056, 5)))
        header_row, name_col = (4, 1) if jobs else (3, 0)
        first_col = name_col + 1
        header = tuple(next(sheet.iter_rows(min_row=header_row, max_row=header_row,
                                            values_only=True)))
        expected = years if jobs else ("Census 2000", "Census 2010", "Census 2020",
                                       *map(str, range(2025, 2056, 5)))
        if (header[first_col:first_col + len(years)] != expected
                or any(v is not None for v in header[first_col + len(years):])
                or (not jobs and header[0] != "Geography")):
            raise ValueError(f"MD: unexpected {variable} year header")
        if jobs:
            markers = {
                "B43": "Historical data from the U.S. Bureau of Economic Analysis, Tables CA25 and CA25N",
                "B45": "Projections from 2025 to 2050 prepared by the Maryland Department of Planning, October 2022.",
            }
        else:
            markers = {
                "A41": "Projected numbers are rounded to the nearest 10.",
                "A44": "Prepared by the Maryland Department of Planning, State Data and Analysis Center, November 2025.",
            }
        for cell, expected_text in markers.items():
            if sheet[cell].value != expected_text:
                raise ValueError(f"MD: unexpected {variable} source footnote at {cell}")

        counties = _county_names(jobs)
        regions = REGIONS.copy()
        if jobs:
            regions.remove("WASHINGTON SUBURBAN REGION")
            regions.add("SUBURBAN WASHINGTON REGION")
        expected_names = set(counties) | regions | {"MARYLAND"}
        seen = set()
        totals = {}
        for row in sheet.iter_rows(min_row=6 if jobs else 4, max_row=42 if jobs else 40,
                                   values_only=True):
            if all(v is None for v in row):
                continue
            name = row[name_col]
            if name not in expected_names:
                raise ValueError(f"MD: unknown {variable} geography {name!r}")
            if name in seen:
                raise ValueError(f"MD: duplicate {variable} geography {name!r}")
            seen.add(name)
            if any(v is not None for v in (*row[:name_col], *row[first_col + len(years):])):
                raise ValueError(f"MD: unexpected extra {variable} columns for {name}")
            values = row[first_col:first_col + len(years)]
            for value in values:
                if (not isinstance(value, (int, float)) or isinstance(value, bool)
                        or not math.isfinite(value) or value < 0 or int(value) != value):
                    raise ValueError(f"MD: invalid {variable} value at {name}: {value!r}")
            if name not in regions:
                totals[name] = tuple(map(int, values))
        if seen != expected_names:
            raise ValueError(f"MD: incomplete {variable} geographies: {sorted(expected_names - seen)}")
        for index, year in enumerate(years):
            if sum(totals[name][index] for name in counties) != totals["MARYLAND"][index]:
                raise ValueError(f"MD: {variable} county sum differs from state at {year}")
        return [
            {"state": "MD", "geo_level": "state" if name == "MARYLAND" else "county",
             "geo_id": "24" if name == "MARYLAND" else counties[name],
             "year": year, "variable": variable, "value": value,
             "datatype": "estimate" if year == 2020 else "projection",
             "source_id": spec["source_id"]}
            for name, values in totals.items()
            for year, value in zip(years, values) if year >= 2020
        ]
    finally:
        workbook.close()


def _validate_catalog(body: bytes) -> None:
    text = " ".join(body.decode("utf-8", "replace").split()).lower()
    for marker in ("<html", "24 jurisdictions", "preliminary", "10/22",
                   "totalpopproj.xlsx", "householdproj.xlsx", "summary_totaljobs.xlsx"):
        if marker not in text:
            raise ValueError(f"MD: unexpected projection catalog; missing {marker!r}")


def _validate_method(body: bytes) -> None:
    from pypdf import PdfReader

    fetch.validate_pdf(body, "MD methodology")
    text = " ".join(" ".join(p.extract_text() for p in PdfReader(io.BytesIO(body)).pages).split())
    for marker in ("Maryland Department of Planning", "cohort-component", "2020 to 2055"):
        if marker not in text:
            raise ValueError(f"MD: unexpected population methodology; missing {marker!r}")


def _snapshots():
    for variable, spec in PRODUCTS.items():
        yield (spec["file"], spec["url"],
               lambda body, variable=variable: parse_workbook(body, variable))
    yield CATALOG_FILE, CATALOG_URL, _validate_catalog
    yield METHOD_FILE, METHOD_URL, _validate_method


def acquire() -> None:
    """Reuse validated originals; changed editions require new filenames."""
    if (RAW_DIR / "manifest.json").exists():
        entries = fetch.read_manifest(RAW_DIR)
        if not {f for f, _, _ in _snapshots()} <= {e["file"] for e in entries}:
            raise ValueError("MD: incomplete saved manifest")
        for filename, _, validate in _snapshots():
            validate((RAW_DIR / filename).read_bytes())
        print("MD acquire: reused validated snapshots")
        return
    snapshots = []
    for filename, url, validate in _snapshots():
        body = fetch.get_bytes(url)
        rows = validate(body)
        snapshots.append((filename, url, body, len(rows) if rows is not None else None))
    entries = []
    for filename, url, body, count in snapshots:
        fetch.write_snapshot(RAW_DIR / filename, body)
        entries.append(dict(file=filename, source=url, retrieved_at=date.today().isoformat(),
                            sha256=fetch.sha256(body), records=count,
                            note="Original Maryland Planning source; edition is recorded in the workbook, not inferred from the catalog or retrieval date."))
    fetch.write_manifest(RAW_DIR, entries)


def load(conn) -> None:
    entries = {e["file"]: e for e in fetch.read_manifest(RAW_DIR)}
    _validate_catalog((RAW_DIR / CATALOG_FILE).read_bytes())
    _validate_method((RAW_DIR / METHOD_FILE).read_bytes())
    count = 0
    for variable, spec in PRODUCTS.items():
        rows = parse_workbook((RAW_DIR / spec["file"]).read_bytes(), variable)
        jobs = variable == "employment"
        method = (
            "Preliminary total jobs by place of work, not employed residents or labor force. "
            "The historical 2020 value is from BEA Tables CA25/CA25N; MDP explicitly labels "
            "2025-2050 as projections. This is a separate October 2022 product."
            if jobs else
            "2020 Census counts are the historical anchor and use the MCP estimate label; "
            "2025-2055 remain projections, rounded by the publisher to the nearest 10. "
            "The methodology describes cohort-component population projections with a 2020-2055 "
            "span; the latest year of auxiliary rate inputs is not specified. Population and "
            "household totals incorporate BMC and MWCOG cooperative forecasts. Household counts "
            "come from the household workbook, not household population or household size. "
            "The workbook's November 2025 edition supersedes the catalog's October 2025 label."
        )
        upsert_source(
            conn, source_id=spec["source_id"], state="MD", agency=AGENCY,
            product_name=spec["title"], vintage=spec["vintage"], base_year=BASE_YEAR,
            horizon_year=spec["horizon"], source_url=spec["url"],
            retrieved_at=entries[spec["file"]]["retrieved_at"], license=LICENSE,
            citation=f"{AGENCY}. {spec['title']}. {spec['vintage']}.",
            method_notes=method + " Base year 2020 is the last historical observation in the "
            "published table. State and all 24 county equivalents are loaded from 2020 in "
            "five-year steps. Baltimore City (24510) is separate from Baltimore County (24005). "
            "Regional rows and earlier history are archived only; no annual interpolation is performed.",
        )
        count += insert_observations(conn, rows)
    conn.commit()
    print(f"MD load: {count} observations")


# Cross-checked with the corresponding publisher PDF tables, page 1.
FIXTURES = [
    ("population", "24", 2020, 6177224),
    ("population", "24", 2055, 7341680),
    ("population", "24510", 2025, 569530),
    ("population", "24005", 2055, 933690),
    ("households", "24", 2055, 2865380),
    ("households", "24510", 2055, 265760),
    ("employment", "24", 2020, 3536398),
    ("employment", "24", 2050, 4179100),
    ("employment", "24510", 2050, 454800),
]


def verify(conn) -> None:
    for variable, gid, year, value in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype, source_id FROM observations WHERE state='MD' "
            "AND variable=? AND geo_id=? AND year=?", (variable, gid, year),
        ).fetchone()
        assert row is not None, f"MD: missing fixture {variable}/{gid}/{year}"
        assert tuple(row) == (value, "estimate" if year == 2020 else "projection",
                              PRODUCTS[variable]["source_id"])
    for variable, spec in PRODUCTS.items():
        n_years = (spec["horizon"] - 2020) // 5 + 1
        shape = conn.execute(
            "SELECT COUNT(*), COUNT(DISTINCT geo_id), COUNT(DISTINCT year), MIN(year), MAX(year) "
            "FROM observations WHERE state='MD' AND variable=?", (variable,),
        ).fetchone()
        assert tuple(shape) == (25 * n_years, 25, n_years, 2020, spec["horizon"])
    bad = conn.execute(
        "SELECT variable, year FROM observations WHERE state='MD' GROUP BY variable, year "
        "HAVING SUM(CASE WHEN geo_level='state' THEN value ELSE -value END) != 0"
    ).fetchall()
    assert not bad, f"MD: county/state sums differ: {bad}"
    print("MD verify: fixtures, grids and county/state sums ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.md --acquire")
