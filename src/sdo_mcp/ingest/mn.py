"""Minnesota adapter.

Source: Minnesota State Demographic Center (SDC), May 2024 Long-Term
Population Projections. State totals and components run 2024-2075; county
age-by-sex detail runs 2024-2055. Verified 2026-08-10
(plan/EXPANSION.md).

Minnesota's key trap is the estimate/projection boundary. The workbook starts
in 2024, but 2024 is itself a projection. Its last underlying estimate is the
2022 Census vintage. Every observation loaded here is therefore labeled
``projection`` and the source records 2022 as its base year.
"""

from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source

RAW_DIR = REPO_ROOT / "data" / "raw" / "mn"
FILE = "2024-Minnesota-Long-Term-Population-Projections v3_tcm36-626199_tcm36-626199.xlsx"
URL = (
    "https://mn.gov/admin/assets/"
    "2024-Minnesota-Long-Term-Population-Projections%20v3_"
    "tcm36-626199_tcm36-626199.xlsx"
)

STATE_FIPS = "27"
SOURCE_ID = "mn-pop-2024"
VINTAGE = "May 2024 (corrected July 1, 2024)"
BASE_YEAR = 2022
STATE_YEARS = (2024, 2075)
COUNTY_YEARS = (2024, 2055)

AGENCY = "Minnesota State Demographic Center, Minnesota Department of Administration"
CITATION = (
    "Minnesota State Demographic Center. Long-Term Population Projections for "
    "Minnesota, May 2024 (corrected July 1, 2024). Retrieved from mn.gov."
)
LICENSE = (
    "No separate license statement; public government data under the Minnesota "
    "Government Data Practices Act."
)


def acquire() -> None:
    """Download and validate the official workbook snapshot."""
    body = fetch.get_bytes(URL)
    if not body.startswith(b"PK"):
        raise RuntimeError(f"{URL}: body is not an xlsx (soft-200 wall or captcha)")

    import io
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        about = wb["ABOUT"]["A1"].value
        required = {"statewide total & components", "county age sex"}
        if about != "Long-Term Population Projections for Minnesota" or not required <= set(wb.sheetnames):
            raise RuntimeError(f"{URL}: unexpected workbook layout")
    finally:
        wb.close()

    RAW_DIR.mkdir(parents=True, exist_ok=True)
    fetch.write_snapshot(RAW_DIR / FILE, body)
    fetch.write_manifest(RAW_DIR, [{
        "file": FILE,
        "source": URL,
        "retrieved_at": date.today().isoformat(),
        "sha256": fetch.sha256(body),
        "records": None,
        "note": (
            "May 2024 vintage. State series 2024-2075; county series "
            "2024-2055. Every point is a projection; the underlying estimate "
            "base is Census Vintage 2022."
        ),
    }])
    print(f"MN acquire: {FILE} ({len(body):,} bytes)")


def _retrieved() -> str:
    return next(
        e["retrieved_at"] for e in fetch.read_manifest(RAW_DIR) if e["file"] == FILE
    )


def load(conn) -> None:
    """Workbook snapshot -> state/county totals and statewide components."""
    import openpyxl

    fips_by_name = {
        row["name"].removesuffix(" County"): row["geo_id"]
        for row in conn.execute(
            "SELECT geo_id, name FROM geographies WHERE state='MN' AND geo_level='county'"
        )
    }

    wb = openpyxl.load_workbook(RAW_DIR / FILE, read_only=True, data_only=True)
    try:
        state_rows = list(wb["statewide total & components"].iter_rows(values_only=True))
        state_years = [int(y) for y in state_rows[0][1:] if y is not None]
        assert (state_years[0], state_years[-1], len(state_years)) == (2024, 2075, 52), (
            f"MN state year span {state_years[0]}-{state_years[-1]} "
            f"({len(state_years)} years)"
        )
        state_values = {str(r[0]): r[1:1 + len(state_years)] for r in state_rows[1:]}

        rows: list[dict] = []
        for year, value in zip(state_years, state_values["TotPop"]):
            rows.append(dict(
                state="MN", geo_level="state", geo_id=STATE_FIPS, year=year,
                variable="population", value=int(value), datatype="projection",
                source_id=SOURCE_ID,
            ))
        for source_name, variable in (
            ("Births", "births"),
            ("Deaths", "deaths"),
            ("NetMig", "net_migration"),
        ):
            for year, value in zip(state_years, state_values[source_name]):
                rows.append(dict(
                    state="MN", geo_level="state", geo_id=STATE_FIPS, year=year,
                    variable=variable, value=int(value), datatype="projection",
                    source_id=SOURCE_ID,
                ))

        county_sheet = wb["county age sex"]
        county_header = next(county_sheet.iter_rows(values_only=True))
        county_years = [int(y) for y in county_header[4:] if y is not None]
        assert (county_years[0], county_years[-1], len(county_years)) == (2024, 2055, 32), (
            f"MN county year span {county_years[0]}-{county_years[-1]} "
            f"({len(county_years)} years)"
        )

        totals: defaultdict[tuple[str, int], int] = defaultdict(int)
        cells: defaultdict[tuple[str, int], int] = defaultdict(int)
        source_names: set[str] = set()
        for rec in county_sheet.iter_rows(min_row=2, values_only=True):
            name = str(rec[0]).strip()
            source_names.add(name)
            for year, value in zip(county_years, rec[4:4 + len(county_years)]):
                totals[(name, year)] += int(value)
                cells[(name, year)] += 1

        unmatched = sorted(source_names - set(fips_by_name))
        if unmatched:
            raise RuntimeError(f"MN load: county names with no Census FIPS match: {unmatched}")
        if len(source_names) != 87:
            raise RuntimeError(f"MN load: found {len(source_names)} counties, expected 87")

        for (name, year), value in totals.items():
            assert cells[(name, year)] == 36, (
                f"MN load: {name} {year} has {cells[(name, year)]} age/sex cells, expected 36"
            )
            rows.append(dict(
                state="MN", geo_level="county", geo_id=fips_by_name[name], year=year,
                variable="population", value=value, datatype="projection",
                source_id=SOURCE_ID,
            ))
    finally:
        wb.close()

    upsert_source(
        conn,
        source_id=SOURCE_ID,
        state="MN",
        agency=AGENCY,
        product_name="Long-Term Population Projections for Minnesota",
        vintage=VINTAGE,
        base_year=BASE_YEAR,
        horizon_year=STATE_YEARS[1],
        source_url=URL,
        retrieved_at=_retrieved(),
        license=LICENSE,
        citation=CITATION,
        method_notes=(
            "State total and components are published directly. County totals are "
            "the exact sums of 18 age groups by two sexes in the county age-sex "
            "sheet. The file contains no estimate rows: its 2024 launch value is "
            "itself a projection based on Census Vintage 2022 estimates, so every "
            "served point is labeled projection and base_year is 2022. County "
            "detail ends in 2055 while the state series continues through 2075."
        ),
    )
    n = insert_observations(conn, rows)
    conn.commit()
    print(f"MN load: {n} observations")


FIXTURES = [
    ("state", "27", 2024, 5777995),
    ("state", "27", 2050, 6139681),
    ("state", "27", 2075, 6112793),
    ("county", "27053", 2024, 1270261),  # Hennepin
    ("county", "27053", 2050, 1227870),
    ("county", "27037", 2024, 454162),   # Dakota
    ("county", "27037", 2050, 552431),
]


def verify(conn) -> None:
    for level, gid, year, value in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype FROM observations WHERE state='MN' AND "
            "geo_level=? AND geo_id=? AND variable='population' AND year=?",
            (level, gid, year),
        ).fetchone()
        assert row is not None, f"MN verify: missing {level} {gid} {year}"
        assert int(row["value"]) == value and row["datatype"] == "projection", (
            f"MN verify: {level} {gid} {year} -> {row['value']}/{row['datatype']}, "
            f"expected {value}/projection"
        )

    bad = conn.execute(
        "SELECT s.year, s.value state_total, c.value county_sum FROM observations s "
        "JOIN (SELECT year, SUM(value) value FROM observations WHERE state='MN' "
        "AND geo_level='county' AND variable='population' GROUP BY year) c "
        "ON c.year=s.year WHERE s.state='MN' AND s.geo_level='state' "
        "AND s.variable='population' AND s.year<=2055 AND s.value!=c.value"
    ).fetchall()
    assert not bad, f"MN verify: county sums differ from state: {[tuple(r) for r in bad[:5]]}"

    county = conn.execute(
        "SELECT MIN(year) lo, MAX(year) hi, COUNT(DISTINCT year) ny, "
        "COUNT(DISTINCT geo_id) ng FROM observations WHERE state='MN' "
        "AND geo_level='county' AND variable='population'"
    ).fetchone()
    assert (county["lo"], county["hi"], county["ny"], county["ng"]) == (2024, 2055, 32, 87)

    component = conn.execute(
        "SELECT variable, MIN(year) lo, MAX(year) hi, COUNT(*) n FROM observations "
        "WHERE state='MN' AND variable IN ('births','deaths','net_migration') "
        "GROUP BY variable"
    ).fetchall()
    assert {r["variable"]: (r["lo"], r["hi"], r["n"]) for r in component} == {
        "births": (2024, 2075, 52),
        "deaths": (2024, 2075, 52),
        "net_migration": (2024, 2075, 52),
    }

    estimates = conn.execute(
        "SELECT COUNT(*) n FROM observations WHERE state='MN' AND datatype='estimate'"
    ).fetchone()["n"]
    assert estimates == 0, f"MN verify: found {estimates} estimate rows; expected none"
    print(f"MN verify: {len(FIXTURES)} fixtures + 32-year county partition + components ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.mn --acquire")
