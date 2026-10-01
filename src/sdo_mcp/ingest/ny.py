"""Cornell PAD's separate 2024 state and 2018 county population products.

State projections retain all three scenarios and the publisher's fractional
counts. County totals come directly from each official county response, never
from the older county product's statewide aggregate. Detail and components
remain archived in the original responses but are not served by this adapter.
"""

from __future__ import annotations

import io
import json
import math
import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "ny"
AGENCY = "Cornell University Program on Applied Demographics (New York's Federal-State population projections representative)"
LICENSE = "No explicit data license identified in the reviewed products; cite Cornell Program on Applied Demographics."
STATE_SOURCE_ID = "ny-pop-state-2024"
COUNTY_SOURCE_ID = "ny-pop-county-2018"
STATE_BASE_YEAR = 2022
COUNTY_BASE_YEAR = 2015
STATE_YEARS = tuple(range(2022, 2051))
COUNTY_YEARS = tuple(range(2015, 2041))
SCENARIOS = {"Middle": "middle", "High": "high", "Low": "low"}
STATE_URL = "https://pad.human.cornell.edu/state_projections/datatools.cfm"
COUNTY_URL = "https://pad.human.cornell.edu/counties/projections.cfm"
STATE_FILES = {
    label: (f"state_2024_{label.lower()}.xlsx",
            f"https://pad.human.cornell.edu/state_projections/downloads/results_Assumptions_{label}.xlsx")
    for label in SCENARIOS
}
COUNTY_API = "https://pad.human.cornell.edu/counties/getprojdata.cfm?county="
COUNTY_HEADER = ("COUNTY", "COUNTY_DESCR", "SEXCODE", "SEX_DESCR", "AGEGRPCODE", "AGEGRP_DESCR", "RACECODE", "RACE_DESCR") + tuple(f"YR_{y}" for y in COUNTY_YEARS)
COMPONENT_HEADER = ("COUNTY", "YEAR", "PROJPERIOD", "STARTPOP", "ENDTOTAL", "NETMIGRANTS", "BIRTHS", "DEATHS", "NATURAL_INCREASE", "SP_CHANGE")
REFERENCES = (
    ("state_catalog_2026-10-01.html", "https://pad.human.cornell.edu/state_projections/", "Fall of 2024"),
    ("state_datatools_2026-10-01.html", STATE_URL, "results_Assumptions_Middle.xlsx"),
    ("state_base_2026-10-01.html", "https://pad.human.cornell.edu/state_projections/meth_base.cfm", "July 1, 2022"),
    ("state_methodology_2026-10-01.html", "https://pad.human.cornell.edu/state_projections/meth_overview.cfm", "cohort-component"),
    ("state_scenarios_2026-10-01.html", "https://pad.human.cornell.edu/state_projections/meth_scenarios.cfm", "Results presented in this set"),
    ("state_trends_2026-10-01.html", "https://pad.human.cornell.edu/state_projections/poptrends.cfm", "18.3 million"),
    ("state_chart_data_2026-10-01.js", "https://pad.human.cornell.edu/state_projections/loadDataTables.js", "middlePopProj.addRows"),
    ("county_catalog_2026-10-01.html", "https://pad.human.cornell.edu/counties/index.cfm", "2018 County Projections"),
    ("county_explorer_2026-10-01.html", COUNTY_URL, "var baseyr = 2015"),
    ("county_methodology_2018.pdf", "https://pad.human.cornell.edu/counties/downloads/2018Methodology.pdf", "September 21, 2018"),
    ("county_user_guide.pdf", "https://pad.human.cornell.edu/counties/downloads/ProjectionsUserManual.pdf", "APPENDIX A"),
    ("producer_2026-10-01.html", "https://pad.human.cornell.edu/people.cfm", "Federal-State Cooperatives"),
)
# Independent Excel exports for spot checks against the official chart API.
CHECK_EXPORTS = (1, 47, 89)


def _counties() -> dict[int, str]:
    result = {}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        fields = line.split("|")
        if fields[0] == "NY":
            result[int(fields[2])] = fields[4].removesuffix(" County")
    if set(result) != set(range(1, 124, 2)):
        raise ValueError("NY: expected 62 Census county identifiers")
    return result


def _number(value, context: str, *, whole=False, negative=False):
    if (type(value) not in (int, float) or not math.isfinite(value)
            or (not negative and value < 0) or (whole and int(value) != value)):
        raise ValueError(f"NY: invalid numeric value at {context}: {value!r}")
    return int(value) if whole else value


def _workbook(body: bytes):
    import openpyxl

    try:
        return openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("NY: response is not a readable population workbook") from exc


def parse_state(body: bytes, scenario_label: str) -> list[dict]:
    """Validate the published edition, complete age/year grid and components."""
    if scenario_label not in SCENARIOS:
        raise ValueError("NY state: unknown scenario")
    workbook = _workbook(body)
    try:
        if workbook.sheetnames != ["Readme", "Components", "PopByAge"]:
            raise ValueError("NY state: unexpected sheets")
        readme = list(workbook["Readme"].values)
        if readme != [
            ("Results population projections for New York State", None), (None, None),
            ("Produced by:", "Cornell Program on Applied Demographics"),
            ("Date:", "October 2024"), ("Scenario name:", scenario_label),
        ]:
            raise ValueError("NY state: unexpected edition, producer or scenario")
        records = iter(workbook["PopByAge"].values)
        if next(records, None) != ("year", "AgeCode", "total", "males", "females"):
            raise ValueError("NY state: unexpected population header")
        expected = {(y, age) for y in STATE_YEARS for age in (-1, *range(101), 999)}
        table = {}
        for row in records:
            if len(row) != 5:
                raise ValueError("NY state: malformed population row")
            year, age, total, males, females = row
            if type(year) is not int or type(age) is not int or (year, age) not in expected:
                raise ValueError("NY state: unexpected year/age identifier")
            if (year, age) in table:
                raise ValueError("NY state: duplicate year/age identifier")
            for value in (total, males, females):
                _number(value, f"{scenario_label}/{year}/{age}", whole=year == STATE_BASE_YEAR and age != 999)
            if age == 999:
                if max(total, males, females) > 120:
                    raise ValueError("NY state: invalid median age")
            elif not math.isclose(total, males + females, rel_tol=0, abs_tol=1e-6):
                raise ValueError("NY state: population sex totals differ")
            table[year, age] = (total, males, females)
        if set(table) != expected:
            raise ValueError("NY state: incomplete population year/age grid")
        for year in STATE_YEARS:
            for i in range(3):
                if not math.isclose(table[year, -1][i], sum(table[year, age][i] for age in range(101)), rel_tol=0, abs_tol=1e-6):
                    raise ValueError("NY state: population age totals differ")
        components = iter(workbook["Components"].values)
        if next(components, None) != ("year", "startpop", "births", "deaths", "dom_in", "dom_out", "int_in", "int_out", "endpop"):
            raise ValueError("NY state: unexpected components header")
        seen = set()
        for row in components:
            if len(row) != 9 or type(row[0]) is not int or row[0] not in STATE_YEARS[1:]:
                raise ValueError("NY state: unexpected component year grid")
            year, start, births, deaths, dom_in, dom_out, int_in, int_out, end = row
            if year in seen:
                raise ValueError("NY state: duplicate component year")
            seen.add(year)
            for value in row[1:]:
                _number(value, f"{scenario_label}/components/{year}")
            if (not math.isclose(start, table[year-1, -1][0], rel_tol=0, abs_tol=1e-6)
                    or not math.isclose(end, table[year, -1][0], rel_tol=0, abs_tol=1e-6)
                    or not math.isclose(end, start + births - deaths + dom_in - dom_out + int_in - int_out, rel_tol=0, abs_tol=1e-6)):
                raise ValueError("NY state: component balance or population differs")
        if seen != set(STATE_YEARS[1:]):
            raise ValueError("NY state: incomplete component year grid")
        return [dict(state="NY", geo_level="state", geo_id="36", year=y,
                     variable="population", value=table[y, -1][0],
                     datatype="estimate" if y == STATE_BASE_YEAR else "projection",
                     scenario=SCENARIOS[scenario_label], source_id=STATE_SOURCE_ID)
                for y in STATE_YEARS]
    finally:
        workbook.close()


def _county_table(header, records, county: int) -> list[dict]:
    names = _counties()
    if county not in names:
        raise ValueError("NY county: unexpected requested county identifier")
    if tuple(header) != COUNTY_HEADER:
        raise ValueError("NY county: unexpected population header")
    expected = {(sex, age) for sex in range(3) for age in (-999, *range(86), 999)}
    table = {}
    for record in records:
        if not isinstance(record, (tuple, list)) or len(record) != len(COUNTY_HEADER):
            raise ValueError("NY county: malformed population row")
        code, name, sex, sex_name, age, age_name, race, race_name, *values = record
        for value in (code, sex, age, race):
            _number(value, "county identifiers", whole=True, negative=True)
        # Source omits the period in St. Lawrence; retain Census spelling in
        # the geography table and require this specific documented alias.
        published_name = "St Lawrence" if county == 89 else names[county]
        if code != county or name != published_name:
            raise ValueError("NY county: mismatched county identifier/name")
        expected_age_name = {-999: "Total", 0: "Newborn", 85: "85+", 999: "Median"}.get(age, str(int(age)))
        # The JSON labels ages as strings; the independent Excel export uses
        # whole numeric cells for labels 1-84. Both identify the same age.
        if type(age_name) in (int, float) and 1 <= age <= 84:
            age_name = str(_number(age_name, "county age label", whole=True))
        if ((sex, age) not in expected or sex_name != {0: "All", 1: "Male", 2: "Female"}.get(sex)
                or age_name != expected_age_name or race != 0 or race_name != "All"):
            raise ValueError("NY county: unexpected sex/age/race identifier or label")
        if (sex, age) in table:
            raise ValueError("NY county: duplicate sex/age row")
        for year, value in zip(COUNTY_YEARS, values):
            _number(value, f"{county}/{sex}/{age}/{year}", whole=age != 999)
        if age == 999 and max(values) > 120:
            raise ValueError("NY county: invalid median age")
        table[sex, age] = values
    if set(table) != expected:
        raise ValueError("NY county: incomplete sex/age grid")
    for i, year in enumerate(COUNTY_YEARS):
        for sex in range(3):
            if table[sex, -999][i] != sum(table[sex, age][i] for age in range(86)):
                raise ValueError("NY county: population age totals differ")
        for age in (-999, *range(86)):
            if table[0, age][i] != table[1, age][i] + table[2, age][i]:
                raise ValueError("NY county: population sex totals differ")
    return [dict(state="NY", geo_level="county", geo_id=f"36{county:03d}", year=year,
                 variable="population", value=int(value),
                 datatype="estimate" if year == COUNTY_BASE_YEAR else "projection",
                 source_id=COUNTY_SOURCE_ID)
            for year, value in zip(COUNTY_YEARS, table[0, -999])]


def parse_county(body: bytes, county: int) -> list[dict]:
    """Parse only complete official county JSON; HTTP 200 is not validation."""
    try:
        data = json.loads(body)
    except Exception as exc:
        raise ValueError("NY county: response is not readable JSON") from exc
    if (not isinstance(data, dict) or set(data) != {"char", "comp"}
            or any(not isinstance(data[key], dict) or set(data[key]) != {"COLUMNS", "DATA"}
                   or not isinstance(data[key]["COLUMNS"], list) or not isinstance(data[key]["DATA"], list)
                   for key in ("char", "comp"))):
        raise ValueError("NY county: unexpected JSON tables")
    observations = _county_table(data["char"]["COLUMNS"], data["char"]["DATA"], county)
    components = data["comp"]
    if tuple(components["COLUMNS"]) != COMPONENT_HEADER:
        raise ValueError("NY county: unexpected components header")
    seen = set()
    for row in components["DATA"]:
        if not isinstance(row, (tuple, list)) or len(row) != 10:
            raise ValueError("NY county: malformed components row")
        code, year, period, start, end, migrants, births, deaths, natural, special = row
        _number(code, "component county identifier", whole=True)
        _number(year, "component year identifier", whole=True)
        if code != county or year not in COUNTY_YEARS[1:] or period != f"yr{int(year) % 100:02d}":
            raise ValueError("NY county: unexpected component identifier/year")
        if year in seen:
            raise ValueError("NY county: duplicate component year")
        seen.add(year)
        for i, value in enumerate(row[3:], 3):
            _number(value, f"{county}/components/{year}", whole=True, negative=i in (5, 8, 9))
        totals = {point["year"]: point["value"] for point in observations}
        if (start != totals[year-1] or end != totals[year] or natural != births - deaths
                or end != start + natural + migrants + special):
            raise ValueError("NY county: component balance or population differs")
    if seen != set(COUNTY_YEARS[1:]):
        raise ValueError("NY county: incomplete component year grid")
    return observations


def parse_county_export(body: bytes, county: int) -> list[dict]:
    workbook = _workbook(body)
    try:
        if workbook.sheetnames != ["Projections"]:
            raise ValueError("NY county export: unexpected sheets")
        sheet = workbook["Projections"]
        # Cornell's original xlsx downloads have incorrect worksheet dimensions.
        sheet.reset_dimensions()
        records = iter(sheet.values)
        return _county_table(next(records, ()), records, county)
    finally:
        workbook.close()


def _reference(body: bytes, filename: str, marker: str) -> None:
    if filename.endswith(".pdf"):
        from pypdf import PdfReader

        fetch.validate_pdf(body, filename)
        text = " ".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(body)).pages)
    else:
        text = body.decode("utf-8-sig")
        if filename.endswith(".html") and "<html" not in text.lower():
            raise ValueError(f"NY: unexpected reference HTML {filename}")
    if marker not in " ".join(text.split()):
        raise ValueError(f"NY: unexpected reference content {filename}")


def _downloads():
    for label, (filename, url) in STATE_FILES.items():
        yield filename, url, lambda body, label=label: parse_state(body, label)
    for county in _counties():
        yield f"county_2018_{county:03d}.json", COUNTY_API + str(county), lambda body, county=county: parse_county(body, county)
    for county in CHECK_EXPORTS:
        yield (f"county_2018_{county:03d}_check.xlsx",
               f"https://pad.human.cornell.edu/counties/expprojdata.cfm?county={county}",
               lambda body, county=county: parse_county_export(body, county))
    for filename, url, marker in REFERENCES:
        yield filename, url, lambda body, filename=filename, marker=marker: _reference(body, filename, marker)


def acquire() -> None:
    """Archive valid original responses; preserve saved bytes and retrieval dates."""
    downloads = list(_downloads())
    if (RAW_DIR / "manifest.json").exists():
        entries = fetch.read_manifest(RAW_DIR)
        if not {name for name, _, _ in downloads} <= {e["file"] for e in entries}:
            raise ValueError("NY: incomplete saved manifest")
        for filename, _, validate in downloads:
            validate((RAW_DIR / filename).read_bytes())
        print("NY acquire: reused validated 2024 state and 2018 county originals")
        return
    validated = []
    for filename, url, validate in downloads:
        body = fetch.get_bytes(url)
        rows = validate(body)
        validated.append((filename, url, body, len(rows) if rows is not None else None))
    entries = []
    for filename, url, body, count in validated:
        fetch.write_snapshot(RAW_DIR / filename, body)
        entries.append(dict(file=filename, source=url, retrieved_at=date.today().isoformat(),
                            sha256=fetch.sha256(body), records=count,
                            note="Original Cornell PAD response. October 2024 state scenarios: base 2022, horizon 2050. September 2018 county edition: base 2015, horizon 2040. Population totals only served; demographic detail and components remain archived."))
    fetch.write_manifest(RAW_DIR, entries)
    print(f"NY acquire: {len(entries)} original data/evidence files")


def load(conn) -> None:
    entries = fetch.read_manifest(RAW_DIR)
    retrieved = {e["file"]: e["retrieved_at"] for e in entries}
    rows = []
    for label, (filename, _) in STATE_FILES.items():
        rows.extend(parse_state((RAW_DIR / filename).read_bytes(), label))
    for county in _counties():
        rows.extend(parse_county((RAW_DIR / f"county_2018_{county:03d}.json").read_bytes(), county))
    expected = {(r["geo_level"], r["geo_id"]) for r in rows}
    actual = {tuple(r) for r in conn.execute("SELECT geo_level, geo_id FROM geographies WHERE state='NY'")}
    if actual != expected:
        raise ValueError("NY: seeded geography set differs from published county grid")
    upsert_source(
        conn, source_id=STATE_SOURCE_ID, state="NY", agency=AGENCY,
        product_name="New York State Population Projections 2022-2050, Middle/High/Low",
        vintage="October 2024", base_year=STATE_BASE_YEAR, horizon_year=2050,
        source_url=STATE_URL, retrieved_at=retrieved[STATE_FILES["Middle"][0]], license=LICENSE,
        citation="Cornell Program on Applied Demographics. New York State Projections 2022-2050. October 2024. Middle, High and Low scenarios.",
        method_notes=(
            "Provisional October 2024 state-only cohort-component edition; catalog reviewed October 1, 2026 still lists the 2024 product. "
            "July 1, 2022 Census Vintage 2023 estimates form the base; base_year=2022 is distinct from the estimate vintage and publication year. "
            "2022 is an estimate; every published annual point 2023-2050 remains a projection even after that year. "
            "Middle, High and Low scenarios are retained separately. Middle is this server's default central scenario; PAD describes it as pre-Covid demographic change but does not identify a recommended scenario. "
            "Published fractional population totals from PopByAge/AgeCode=-1 are preserved without rounding or recomputation. "
            "Age/sex detail and annual components are archived but not served. This state product is not a sum of the older 2018 county series. "
            f"Base/methods: https://pad.human.cornell.edu/state_projections/meth_base.cfm; scenario downloads: {STATE_URL}."),
    )
    upsert_source(
        conn, source_id=COUNTY_SOURCE_ID, state="NY", agency=AGENCY,
        product_name="2018 New York County Population Projections 2015-2040",
        vintage="September 21, 2018", base_year=COUNTY_BASE_YEAR, horizon_year=2040,
        source_url=COUNTY_URL, retrieved_at=retrieved["county_2018_001.json"], license=LICENSE,
        citation="Cornell Program on Applied Demographics. 2018 County Population Projections 2015-2040. Released September 21, 2018.",
        method_notes=(
            "Separate 2018 county cohort-component product, still listed in the October 1, 2026 catalog pending county updates to the newer state controls. "
            "July 2015 starting population comes from Vintage 2017 Bridged-Race Postcensal Population Estimates downloaded through CDC-Wonder. "
            "2015 is the estimate base; all annual 2016-2040 points are projections, regardless of retrieval date. One published scenario; no Low/Middle/High options apply to counties. "
            "All 62 county all-sex/all-age totals are read directly from the official chart JSON responses. County IDs are the Census county suffix; prefix 36 forms full FIPS. "
            "Published statewide and ten economic-region aggregates of this county product are not loaded; no county sums are substituted for the 2024 statewide forecast. "
            "Demographic detail, median ages and components remain archived only. Methodology: https://pad.human.cornell.edu/counties/downloads/2018Methodology.pdf."),
    )
    count = insert_observations(conn, rows)
    conn.commit()
    print(f"NY load: {count} population observations")


def verify(conn) -> None:
    state = conn.execute("SELECT COUNT(*), COUNT(DISTINCT scenario), MIN(year), MAX(year), SUM(datatype='estimate') FROM observations WHERE source_id=?", (STATE_SOURCE_ID,)).fetchone()
    county = conn.execute("SELECT COUNT(*), COUNT(DISTINCT geo_id), MIN(year), MAX(year), SUM(datatype='estimate') FROM observations WHERE source_id=?", (COUNTY_SOURCE_ID,)).fetchone()
    if tuple(state) != (87, 3, 2022, 2050, 3) or tuple(county) != (1612, 62, 2015, 2040, 62):
        raise ValueError("NY: incomplete state/county population grid")
    bases = conn.execute("SELECT value, datatype FROM observations WHERE source_id=? AND year=2022", (STATE_SOURCE_ID,)).fetchall()
    if len(bases) != 3 or any(tuple(row) != (19673203, "estimate") for row in bases):
        raise ValueError("NY state: independent chart base fixture differs")
    # Separately published chart data limits totals to ten significant digits.
    for scenario, expected in (("middle", 18288497.49), ("high", 19503775.2), ("low", 16927438.21)):
        row = conn.execute("SELECT value, datatype FROM observations WHERE source_id=? AND scenario=? AND year=2050", (STATE_SOURCE_ID, scenario)).fetchone()
        if row is None or not math.isclose(row[0], expected, rel_tol=0, abs_tol=0.005) or row[1] != "projection":
            raise ValueError(f"NY state: independent chart fixture differs: {scenario}")
    for code in CHECK_EXPORTS:
        filename = f"county_2018_{code:03d}_check.xlsx"
        for point in parse_county_export((RAW_DIR / filename).read_bytes(), code):
            row = conn.execute("SELECT value, datatype FROM observations WHERE source_id=? AND geo_id=? AND year=?", (COUNTY_SOURCE_ID, point["geo_id"], point["year"])).fetchone()
            if row is None or tuple(row) != (point["value"], point["datatype"]):
                raise ValueError(f"NY county: independent Excel fixture differs: {code}/{point['year']}")
    print("NY verify: separate grids, three state scenarios and independent chart/Excel fixtures ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        from .. import store

        connection = store.connect(create=True)
        try:
            load(connection)
            verify(connection)
        finally:
            connection.close()
