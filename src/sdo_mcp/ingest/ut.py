"""Utah Kem C. Gardner Policy Institute adapter (2025-2065 vintage).

The official workbook contains annual July 1 baseline projections for Utah and
all 29 counties: population, components of change, households, and total jobs.
The 2025 launch row is explicitly part of the projection product, so it is
labeled ``projection``.  The workbook does not state a last underlying estimate
year; ``base_year`` therefore remains NULL rather than being inferred.
"""

from __future__ import annotations

import io
import math
import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source

RAW_DIR = REPO_ROOT / "data" / "raw" / "ut"
FILE = "Gardner-Policy-Institute-State-and-County-Projections-2025-2065-Data.xlsx"
REPORT_FILE = "ProjectionsSummary-Mar2026.pdf"
URL = "https://d36oiwf74r1rap.cloudfront.net/wp-content/uploads/2026/08/Gardner-Policy-Institute-State-and-County-Projections-2025-2065-Data.xlsx"
REPORT_URL = "https://d36oiwf74r1rap.cloudfront.net/wp-content/uploads/2026/03/ProjectionsSummary-Mar2026.pdf"
STATE_FIPS = "49"
POP_SOURCE_ID = "ut-pop-2025"
HH_SOURCE_ID = "ut-households-2025"
JOBS_SOURCE_ID = "ut-employment-2025"
BASE_YEAR = None
HORIZON_YEAR = 2065
VINTAGE = "2025-2065 baseline (November 2025; workbook at August 2026 asset URL)"
YEARS = tuple(range(2025, 2066))
AGENCY = "Kem C. Gardner Policy Institute, University of Utah"
LICENSE = "Public University of Utah data; cite the Kem C. Gardner Policy Institute."


def acquire() -> None:
    """Download and validate the official workbook and accompanying report."""
    import openpyxl

    body = fetch.get_bytes(URL)
    if not body.startswith(b"PK"):
        raise RuntimeError(f"{URL}: body is not an xlsx (HTML or soft-200 error)")
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        expected = {"File Layout", "Demographic Detail", "Total Employment by County", "State Industry Employment"}
        if set(wb.sheetnames) != expected:
            raise RuntimeError(f"{URL}: unexpected workbook sheets {wb.sheetnames!r}")
        if wb["File Layout"]["A1"].value != "Source: Kem C. Gardner Policy Institute State and County Projections 2025-2065":
            raise RuntimeError(f"{URL}: unexpected workbook title")
        if wb["Demographic Detail"]["A4"].value != "Geography" or wb["Total Employment by County"]["A4"].value != "Geography":
            raise RuntimeError(f"{URL}: unexpected table headers")
    finally:
        wb.close()
    report = fetch.get_bytes(REPORT_URL)
    if not report.startswith(b"%PDF"):
        raise RuntimeError(f"{REPORT_URL}: body is not a PDF (HTML or soft-200 error)")
    fetch.validate_pdf(report, REPORT_URL)
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    fetch.write_snapshot(RAW_DIR / FILE, body)
    fetch.write_snapshot(RAW_DIR / REPORT_FILE, report)
    today = date.today().isoformat()
    fetch.write_manifest(RAW_DIR, [
        dict(file=FILE, source=URL, retrieved_at=today, sha256=fetch.sha256(body), records=30 * 41,
             note="November 2025 Gardner Policy Institute state/county baseline workbook. Annual July 1 population, households, and total employment 2025-2065; state/county components included in demographic detail."),
        dict(file=REPORT_FILE, source=REPORT_URL, retrieved_at=today, sha256=fetch.sha256(report), records=None,
             note="March 2026 official Utah 2065 projection report and methodology summary."),
    ])
    print(f"UT acquire: {FILE} ({len(body):,} bytes)")


def _retrieved(filename: str = FILE) -> str:
    return next(e["retrieved_at"] for e in fetch.read_manifest(RAW_DIR) if e["file"] == filename)


def _geographies(conn) -> dict[str, tuple[str, str]]:
    by_name = {
        row["name"]: ("county", row["geo_id"])
        for row in conn.execute("SELECT geo_id, name FROM geographies WHERE state='UT' AND geo_level='county'")
    }
    if len(by_name) != 29:
        raise RuntimeError(f"UT load: expected 29 counties, found {len(by_name)}")
    by_name["Utah State"] = ("state", STATE_FIPS)
    return by_name


def load(conn) -> None:
    """Load baseline population, components, households, and jobs totals."""
    import openpyxl

    geographies = _geographies(conn)
    wb = openpyxl.load_workbook(RAW_DIR / FILE, read_only=True, data_only=True)
    try:
        detail = wb["Demographic Detail"]
        header = tuple(next(detail.iter_rows(min_row=4, max_row=4, values_only=True))[:18])
        expected_header = ("Geography", "Year", "Population", "Births", "Deaths", "Natural Change", "Net Migration",
                           "Population Ages 0-4", "Population Ages 5-17", "Population Ages 18-24", "Population Ages 25-64",
                           "Population Ages 65-84", "Population Ages 85+", "Households", "Persons Per Household",
                           "Household Population", "Group Quarters Population", "Median Age")
        if header != expected_header:
            raise ValueError("UT load: unexpected demographic-detail header")
        rows: list[dict] = []
        seen = set()
        for rec in detail.iter_rows(min_row=5, values_only=True):
            name, year = rec[:2]
            if name is None and year is None:
                continue
            if name not in geographies or year not in YEARS:
                raise ValueError(f"UT load: unexpected geography/year {name!r}/{year!r}")
            if (name, year) in seen:
                raise ValueError(f"UT load: duplicate demographic row {name}/{year}")
            seen.add((name, year))
            level, gid = geographies[name]
            for index, variable in ((2, "population"), (3, "births"), (4, "deaths"), (6, "net_migration"), (13, "households")):
                value = rec[index]
                if (not isinstance(value, (int, float)) or isinstance(value, bool)
                        or not math.isfinite(float(value)) or int(value) != value
                        or (variable != "net_migration" and float(value) < 0)):
                    raise ValueError(f"UT load: invalid {variable} at {name}/{year}: {value!r}")
                rows.append(dict(state="UT", geo_level=level, geo_id=gid, year=int(year), variable=variable,
                                 value=int(value), datatype="projection", scenario="baseline",
                                 source_id=(HH_SOURCE_ID if variable == "households" else POP_SOURCE_ID)))
        expected = {(name, year) for name in geographies for year in YEARS}
        if seen != expected:
            raise RuntimeError(f"UT load: demographic rows differ; missing {len(expected - seen)}")

        jobs = wb["Total Employment by County"]
        if tuple(next(jobs.iter_rows(min_row=4, max_row=4, values_only=True))[:3]) != ("Geography", "Year", "Total Jobs"):
            raise ValueError("UT load: unexpected county-employment header")
        county_seen = set()
        county_state_jobs = {}
        for rec in jobs.iter_rows(min_row=5, values_only=True):
            name, year, value = rec[:3]
            if name is None and year is None:
                continue
            # State jobs are taken from the dedicated State Industry sheet below;
            # the county table's state row is cross-checked and matches it. The
            # sum of the 29 county rows can differ by a few rounded jobs.
            if name == "Utah State":
                if year in county_state_jobs:
                    raise ValueError(f"UT load: duplicate county-table state jobs row {year}")
                if year not in YEARS or not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(float(value)) or int(value) != value or float(value) < 0:
                    raise ValueError(f"UT load: invalid county-table state jobs row {year!r}/{value!r}")
                county_state_jobs[year] = int(value)
                continue
            if name not in geographies or geographies[name][0] != "county" or year not in YEARS:
                raise ValueError(f"UT load: unexpected county job row {name!r}/{year!r}")
            if (name, year) in county_seen:
                raise ValueError(f"UT load: duplicate county jobs row {name}/{year}")
            if (not isinstance(value, (int, float)) or isinstance(value, bool)
                    or not math.isfinite(float(value)) or int(value) != value or value < 0):
                raise ValueError(f"UT load: invalid county jobs at {name}/{year}: {value!r}")
            county_seen.add((name, year))
            level, gid = geographies[name]
            rows.append(dict(state="UT", geo_level=level, geo_id=gid, year=int(year), variable="employment",
                             value=int(value), datatype="projection", scenario="baseline", source_id=JOBS_SOURCE_ID))
        expected_county = {(name, year) for name, (level, _gid) in geographies.items() if level == "county" for year in YEARS}
        if county_seen != expected_county:
            raise RuntimeError(f"UT load: county employment rows differ; missing {len(expected_county - county_seen)}")

        industry = wb["State Industry Employment"]
        if tuple(next(industry.iter_rows(min_row=4, max_row=4, values_only=True))[:2]) != ("Year", "Total Jobs"):
            raise ValueError("UT load: unexpected state-employment header")
        state_jobs = {}
        for rec in industry.iter_rows(min_row=5, values_only=True):
            year, value = rec[:2]
            if year is None and value is None:
                continue
            if year in state_jobs:
                raise ValueError(f"UT load: duplicate state jobs row {year}")
            if (year not in YEARS or not isinstance(value, (int, float)) or isinstance(value, bool)
                    or not math.isfinite(float(value)) or int(value) != value or value < 0):
                raise ValueError(f"UT load: invalid state jobs row {year!r}/{value!r}")
            state_jobs[year] = int(value)
            rows.append(dict(state="UT", geo_level="state", geo_id=STATE_FIPS, year=int(year), variable="employment",
                             value=int(value), datatype="projection", scenario="baseline", source_id=JOBS_SOURCE_ID))
        if set(state_jobs) != set(YEARS):
            raise RuntimeError(f"UT load: state employment years differ; missing {sorted(set(YEARS) - state_jobs.keys())}")
        if set(county_state_jobs) != set(YEARS):
            raise RuntimeError("UT load: county-table state employment row is incomplete")
        if state_jobs != county_state_jobs:
            raise RuntimeError("UT load: state employment totals differ between published tables")
    finally:
        wb.close()

    upsert_source(conn, source_id=POP_SOURCE_ID, state="UT", agency=AGENCY,
                   product_name="Utah 2065 State and County Population Projections",
                   vintage=VINTAGE,
                   base_year=BASE_YEAR, horizon_year=HORIZON_YEAR, source_url=URL,
                   retrieved_at=_retrieved(), license=LICENSE,
                   citation="Kem C. Gardner Policy Institute. State and County Projections 2025-2065.",
                   method_notes=(
                       "Annual July 1 baseline projections for the state and 29 counties. Population, births, deaths, and net migration are published directly in the demographic-detail table. "
                       "The report describes the 2025-2065 product as projections; the workbook does not identify a last underlying estimate year, so base_year is intentionally NULL. "
                       "County totals differ from the state row by at most four people from independent rounding; no county values are adjusted."
                   ))
    upsert_source(conn, source_id=HH_SOURCE_ID, state="UT", agency=AGENCY,
                   product_name="Utah 2065 State and County Household Projections",
                   vintage=VINTAGE, base_year=BASE_YEAR,
                   horizon_year=HORIZON_YEAR, source_url=URL, retrieved_at=_retrieved(), license=LICENSE,
                   citation="Kem C. Gardner Policy Institute. State and County Projections 2025-2065.",
                   method_notes="Annual July 1 household totals and average household size for Utah and all 29 counties. All rows are projections; the source does not state a last estimate year.")
    upsert_source(conn, source_id=JOBS_SOURCE_ID, state="UT", agency=AGENCY,
                   product_name="Utah 2065 State and County Total Employment Projections",
                   vintage=VINTAGE, base_year=BASE_YEAR,
                   horizon_year=HORIZON_YEAR, source_url=URL, retrieved_at=_retrieved(), license=LICENSE,
                   citation="Kem C. Gardner Policy Institute. State and County Projections 2025-2065.",
                   method_notes=(
                       "Annual July 1 total jobs for Utah and 29 counties. The report defines employment as wage-and-salary plus self-employment and describes the series as projections. "
                       "State totals come from the State Industry Employment table; county totals come from Total Employment by County. Their published sums differ by a few jobs because of rounding, so neither table is reconciled to the other."
                   ))
    n = insert_observations(conn, rows)
    conn.commit()
    print(f"UT load: {n} observations")


FIXTURES = [
    ("population", "state", "49", 2025, 3556394, POP_SOURCE_ID),
    ("population", "state", "49", 2065, 5550525, POP_SOURCE_ID),
    ("population", "county", "49035", 2050, 1473675, POP_SOURCE_ID),  # Salt Lake
    ("households", "state", "49", 2065, 2332663, HH_SOURCE_ID),
    ("households", "county", "49049", 2050, 452295, HH_SOURCE_ID),    # Utah County
    ("employment", "state", "49", 2025, 2503944, JOBS_SOURCE_ID),
    ("employment", "state", "49", 2065, 3729317, JOBS_SOURCE_ID),
    ("employment", "county", "49035", 2065, 1683385, JOBS_SOURCE_ID),
]


def verify(conn) -> None:
    for variable, level, gid, year, value, source_id in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype, source_id FROM observations WHERE state='UT' AND geo_level=? AND geo_id=? "
            "AND variable=? AND year=?", (level, gid, variable, year)
        ).fetchone()
        assert row is not None, f"UT verify: missing {variable} {level} {gid} {year}"
        assert (int(row["value"]), row["datatype"], row["source_id"]) == (value, "projection", source_id)
    for variable, source in (("population", POP_SOURCE_ID), ("households", HH_SOURCE_ID), ("employment", JOBS_SOURCE_ID)):
        shape = conn.execute(
            "SELECT COUNT(*) n, COUNT(DISTINCT geo_id) ng, COUNT(DISTINCT year) ny, MIN(year) lo, MAX(year) hi "
            "FROM observations WHERE state='UT' AND variable=? AND source_id=?", (variable, source)
        ).fetchone()
        assert tuple(shape) == (1230, 30, 41, 2025, 2065), f"UT verify: {variable} shape {tuple(shape)}"
    for variable in ("births", "deaths", "net_migration"):
        shape = conn.execute(
            "SELECT COUNT(*) n, COUNT(DISTINCT geo_id) ng, COUNT(DISTINCT year) ny FROM observations "
            "WHERE state='UT' AND variable=?", (variable,)
        ).fetchone()
        assert tuple(shape) == (1230, 30, 41)
    assert conn.execute("SELECT COUNT(*) FROM observations WHERE state='UT' AND datatype='estimate'").fetchone()[0] == 0
    print(f"UT verify: {len(FIXTURES)} fixtures + annual state/county pop/HH/jobs/components shapes ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.ut --acquire")
