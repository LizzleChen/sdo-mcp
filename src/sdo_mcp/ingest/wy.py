"""Wyoming EAD February 2025 state and county population forecasts."""

from __future__ import annotations

import io
import math
import sys
from datetime import date, datetime

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "wy"
CATALOG_URL = "https://ai.wyo.gov/divisions/economic-analysis/demographic-information/projections"
SHEET_URL = "https://docs.google.com/spreadsheets/d/19ooDLwh4bIXnGUcaVr95CQ_cCXvypcFa/edit"
WORKBOOK_URL = "https://docs.google.com/spreadsheets/d/19ooDLwh4bIXnGUcaVr95CQ_cCXvypcFa/export?format=xlsx"
PDF_URL = "https://drive.google.com/file/d/1w7y6YRfgZ3k9voPjqw0dlJkUiwwiYlEd/view"
PDF_DOWNLOAD_URL = "https://drive.google.com/uc?export=download&id=1w7y6YRfgZ3k9voPjqw0dlJkUiwwiYlEd"
CATALOG_FILE = "projection_catalog_2026-09-30.html"
WORKBOOK_FILE = "population_2020-2050_feb2025.xlsx"
PDF_FILE = "population_2020-2050_feb2025.pdf"
FILES = {CATALOG_FILE: CATALOG_URL, WORKBOOK_FILE: WORKBOOK_URL, PDF_FILE: PDF_DOWNLOAD_URL}
SOURCE_ID = "wy-pop-feb2025"
AGENCY = "Wyoming Department of Administration & Information, Economic Analysis Division"
TITLE = "Population for Wyoming, Counties, Cities, and Towns: 2020 to 2050"
FOOTER = f"Prepared by {AGENCY} (http://eadiv.state.wy.us), February 2025"
YEARS = tuple(range(2020, 2051))


def _counties() -> dict[str, str]:
    counties = {}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        fields = line.split("|")
        if fields[0] == "WY":
            counties[fields[4]] = fields[1] + fields[2]
    if len(counties) != 23:
        raise ValueError("WY: expected 23 Census counties")
    return counties


def _name(source: str) -> str:
    return source.removesuffix(" Cnty") + " County" if source.endswith(" Cnty") else source


def parse_workbook(body: bytes) -> list[dict]:
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("WY: response is not a readable XLSX workbook") from exc
    try:
        if book.sheetnames != ["Sheet1", "Print"]:
            raise ValueError("WY: unexpected workbook sheets")
        grid = list(book["Sheet1"].iter_rows(values_only=True))
        if (len(grid) != 265 or any(len(row) != 34 for row in grid)
                or grid[0][0] != TITLE or tuple(grid[2][1:32]) != YEARS
                or grid[3][0] != "AREA" or grid[3][1] != "CENSUS"
                or any(label != "Estimate" for label in grid[3][2:5])
                or any(str(label).strip() != "Forecast" for label in grid[3][5:32])
                or grid[137][0] != FOOTER
                or book.properties.modified != datetime(2025, 2, 3, 15, 6, 56)):
            raise ValueError("WY: unexpected workbook edition, years or estimate boundary")
        counties = _counties()
        rows = []
        totals = {}
        seen = set()
        for source in grid[5:130]:
            name = source[0]
            if name == "WYOMING":
                gid = "56"
            elif isinstance(name, str) and name.endswith((" County", " Cnty")):
                gid = counties.get(_name(name))
                if gid is None:
                    raise ValueError(f"WY: unknown county name {name!r}")
            else:
                continue  # Published municipal rows are archived for later place support.
            if gid in seen or source[33] != name:
                raise ValueError(f"WY: duplicate or mismatched geography {name!r}")
            seen.add(gid)
            values = source[1:32]
            if any(not isinstance(v, (int, float)) or isinstance(v, bool)
                   or not math.isfinite(v) or v < 0 or int(v) != v for v in values):
                raise ValueError(f"WY: invalid population value for {name}")
            totals[gid] = tuple(map(int, values))
            rows.extend({
                "state": "WY", "geo_level": "state" if gid == "56" else "county",
                "geo_id": gid, "year": year, "variable": "population", "value": int(value),
                "datatype": "estimate" if year <= 2023 else "projection",
                "source_id": SOURCE_ID,
            } for year, value in zip(YEARS, values))
        if seen != set(counties.values()) | {"56"}:
            raise ValueError("WY: incomplete state/county population grid")
        for i, year in enumerate(YEARS):
            if totals["56"][i] != sum(values[i] for gid, values in totals.items() if gid != "56"):
                raise ValueError(f"WY: county totals do not reconcile in {year}")
        return rows
    finally:
        book.close()


def parse_pdf(body: bytes) -> dict[tuple[str, int], int]:
    """Read the separately published PDF table for an independent value check."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(body))
        if len(reader.pages) != 12:
            raise ValueError("unexpected PDF page count")
        names = {"WYOMING": "56"} | {name: gid for name, gid in _counties().items()}
        names |= {name.replace(" County", " Cnty"): gid
                  for name, gid in _counties().items() if name in ("Johnson County", "Laramie County", "Lincoln County")}
        totals = {}
        for page_number, page in enumerate(reader.pages):
            lines = page.extract_text().splitlines()
            years = tuple(range(2020, 2031)) if page_number < 4 else (
                tuple(range(2031, 2041)) if page_number < 8 else tuple(range(2041, 2051)))
            if lines[0] != TITLE or tuple(map(int, lines[1].split())) != years:
                raise ValueError("unexpected PDF title or year header")
            for line in lines:
                for name, gid in names.items():
                    if not line.startswith(name + " "):
                        continue
                    values = [int(token.replace(",", "")) for token in line[len(name) + 1:].split()]
                    if len(values) != len(years):
                        raise ValueError(f"incomplete PDF row for {name}")
                    for year, value in zip(years, values):
                        key = (gid, year)
                        if key in totals:
                            raise ValueError(f"duplicate PDF cell {key}")
                        totals[key] = value
                    break
        if len(totals) != 24 * len(YEARS) or FOOTER not in reader.pages[-1].extract_text():
            raise ValueError("incomplete PDF geography grid or edition note")
        return totals
    except Exception as exc:
        raise ValueError("WY: response is not the expected population PDF") from exc


def _validate_catalog(body: bytes) -> None:
    page = " ".join(body.decode("utf-8", "replace").lower().split())
    for marker in ("<html", "population forecasts for wyoming", "19oodlwh4bixngucavr95cq_ccxvypcfa",
                   "1w7y6yrfgz3k9vopjqw0dljkuiwwiyled"):
        if marker not in page:
            raise ValueError(f"WY: invalid publisher catalog; missing {marker!r}")


def _validate_all(bodies: dict[str, bytes]) -> list[dict]:
    _validate_catalog(bodies[CATALOG_FILE])
    rows = parse_workbook(bodies[WORKBOOK_FILE])
    pdf = parse_pdf(bodies[PDF_FILE])
    for row in rows:
        key = (row["geo_id"], row["year"])
        if pdf[key] != row["value"]:
            raise ValueError(f"WY: workbook and PDF disagree at {key}")
    return rows


def acquire() -> None:
    if (RAW_DIR / "manifest.json").exists():
        entries = {e["file"] for e in fetch.read_manifest(RAW_DIR)}
        if entries != set(FILES):
            raise ValueError("WY: incomplete saved snapshot manifest")
        _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
        print("WY acquire: reused validated snapshots")
        return
    bodies = {name: fetch.get_bytes(url) for name, url in FILES.items()}
    rows = _validate_all(bodies)
    entries = []
    for name, url in FILES.items():
        body = bodies[name]
        fetch.write_snapshot(RAW_DIR / name, body)
        entries.append({
            "file": name, "source": url, "retrieved_at": date.today().isoformat(),
            "sha256": fetch.sha256(body), "records": len(rows) if name == WORKBOOK_FILE else None,
            "note": "Original Wyoming EAD February 2025 population publication; workbook and PDF include municipal rows not served by this adapter.",
        })
    fetch.write_manifest(RAW_DIR, entries)


def load(conn) -> None:
    entries = {e["file"]: e for e in fetch.read_manifest(RAW_DIR)}
    rows = _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
    upsert_source(
        conn, source_id=SOURCE_ID, state="WY", agency=AGENCY,
        product_name="Wyoming state and county population forecasts", vintage="February 2025",
        base_year=2023, horizon_year=2050, source_url=SHEET_URL,
        retrieved_at=entries[WORKBOOK_FILE]["retrieved_at"],
        license="No explicit data license identified; cite the Wyoming Economic Analysis Division.",
        citation=f"{AGENCY}. Population for Wyoming, Counties, Cities, and Towns: 2020 to 2050. February 2025.",
        method_notes=(
            "The publisher's February 2025 workbook and companion PDF agree on every loaded state "
            "and county value. The 2020 rows are Census counts, 2021-2023 are Census Bureau "
            "estimates, and 2024-2050 are EAD forecasts; base_year=2023 records the final "
            "estimate year. All 23 county totals reconcile exactly to the separately published "
            "state row for every year. Johnson, Laramie, and Lincoln are abbreviated 'Cnty' "
            "in the source and mapped to Census county FIPS. The publisher also calculates "
            "municipal forecasts from place/county ratios; those published place rows are "
            "archived but not loaded. No household or eligible EAD jobs forecast was established "
            "during this population-source check."
        ),
    )
    insert_observations(conn, rows)
    conn.commit()
    print(f"WY load: {len(rows)} observations")


def verify(conn) -> None:
    fixtures = {("56", 2020): (576851, "estimate"),
                ("56", 2023): (584057, "estimate"),
                ("56", 2024): (586980, "projection"),
                ("56", 2050): (625230, "projection"),
                ("56001", 2024): (38590, "projection"),
                ("56003", 2050): (12830, "projection"),
                ("56021", 2050): (111060, "projection")}
    for (gid, year), expected in fixtures.items():
        row = conn.execute("SELECT value, datatype FROM observations WHERE "
                           "state='WY' AND variable='population' AND geo_id=? AND year=?",
                           (gid, year)).fetchone()
        assert tuple(row) == expected
    shape = conn.execute("SELECT COUNT(*),COUNT(DISTINCT geo_id),COUNT(DISTINCT year),"
                         "MIN(year),MAX(year) FROM observations WHERE state='WY'").fetchone()
    assert tuple(shape) == (744, 24, 31, 2020, 2050)
    print("WY verify: independent PDF values and complete annual grid ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.wy --acquire")
