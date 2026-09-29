"""Michigan MCDA Vintage 2024 population totals from the original workbook.

The download contains state, county and prosperity-region rows in five-year
steps. Serve only state/county totals, retaining the publisher's decimals.
The workbook contains no estimates, components or alternative scenarios.
"""

from __future__ import annotations

import io
import math
import sys
from datetime import date
from urllib.request import Request, urlopen

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "mi"
FILE = "population_v2024.xlsx"
URL = "https://milmi.org/_docs/publications/projectionsdata/County_Population_Projections_2050.xlsx"
DATA_URL = "https://milmi.org/DataSearch/POPPROJ"
REPORT_URL = "https://www.michigan.gov/-/media/Project/Websites/mcda/reports/2024/Michigan-Statewide-Population-Projections-through-2050.pdf?rev=bf065eb658824929b65255e0e8b9a93c"
COUNTY_METHOD_URL = "https://www.michigan.gov/-/media/Project/Websites/mcda/assets/Insight-Media-Files/Mi-County-Proj/Methodology---MiCtyPopProj.pdf?hash=396E088A6297591C3644528592AA07AA&rev=d93f987a0828444eade6640080b280e4"
COUNTY_REPORT_URL = "https://www.michigan.gov/mcda/insights/2025/03/06/mich-county-popproj-2050"
REFERENCES = (
    ("state_report_v2024.pdf", REPORT_URL, None),
    ("county_methodology_v2024.pdf", COUNTY_METHOD_URL, None),
    ("county_release_2025-03-06.html", COUNTY_REPORT_URL, "Vintage 2024 Population Projections"),
    ("data_catalog_2026-09-29.html", DATA_URL, "County_Population_Projections_2050.xlsx"),
)
STATE_SOURCE_ID = "mi-pop-state-2024"
COUNTY_SOURCE_ID = "mi-pop-county-2024"
AGENCY = "Michigan Center for Data and Analytics (MCDA), State Demographer, Department of Technology, Management and Budget"
LICENSE = "No explicit data license identified in the reviewed products; cite Michigan Center for Data and Analytics."
YEARS = tuple(range(2025, 2051, 5))
SHEET = "raked_projections_with_geo_name"
HEADER = ["stfips", "areatype", "area", "geo_name", "periodyear", "periodtype",
          "period", "popsource", "population", "female", "male"] + [
    prefix + age for age in ("under5", *(f"{a}to{a+4}" for a in range(5, 85, 5)), "85xx")
    for prefix in ("tot", "fem", "male")
]
# These are source regions, not counties or Census metropolitan areas. Validate
# them before skipping so a changed geography coding cannot silently pass.
REGIONS = {
    560100: "Upper Peninsula Prosperity Alliance",
    560200: "Northwest Prosperity Region",
    560300: "Northeast Prosperity Region",
    560400: "West Michigan Prosperity Alliance",
    560500: "East Central Michigan Prosperity Region",
    560600: "East Michigan Prosperity Region",
    560700: "South Central Prosperity Region",
    560800: "Southwest Prosperity Region",
    560900: "Southeast Michigan Prosperity Region",
    561000: "Detroit Metro Prosperity Region",
}


def _geographies() -> dict[tuple[int, int], str]:
    geos = {(1, 0): "Michigan"}
    aliases = {"St. Clair": "Saint Clair", "St. Joseph": "Saint Joseph"}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        parts = line.split("|")
        if parts[0] == "MI":
            name = parts[4].removesuffix(" County")
            geos[4, int(parts[2])] = aliases.get(name, name)
    if len(geos) != 84:
        raise ValueError("MI: expected Michigan and 83 Census counties")
    geos.update({(56, code): name for code, name in REGIONS.items()})
    return geos


def parse(body: bytes) -> list[dict]:
    """Reject soft-200 errors, changed structure and incomplete source grids."""
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("MI: response is not a readable population workbook") from exc
    try:
        if workbook.sheetnames != [SHEET]:
            raise ValueError("MI: unexpected workbook sheets")
        rows = iter(workbook[SHEET].iter_rows(values_only=True))
        if tuple(next(rows, ())) != tuple(HEADER):
            raise ValueError("MI: unexpected population header")
        geos = _geographies()
        expected = {(kind, code, year) for kind, code in geos for year in YEARS}
        seen = set()
        observations = []
        for row in rows:
            if all(v is None for v in row):
                continue
            if len(row) != len(HEADER):
                raise ValueError("MI: malformed population row")
            state, kind, code, name, year, periodtype, period, popsource, *values = row
            codes = (state, kind, code, year, periodtype, period, popsource)
            if (any(type(v) is not int for v in codes) or
                    (state, periodtype, period, popsource) != (26, 52, 0, 4)):
                raise ValueError("MI: unexpected geography, year or series codes")
            if geos.get((kind, code)) != name or (kind, code) not in geos:
                raise ValueError(f"MI: unknown geography or mismatched name: {kind}/{code}/{name!r}")
            key = kind, code, year
            if key in seen or key not in expected:
                raise ValueError(f"MI: duplicate or unexpected geography/year: {key}")
            seen.add(key)
            if any(type(v) not in (int, float) or not math.isfinite(v) or v < 0 for v in values):
                raise ValueError(f"MI: invalid population value in {key}")
            if kind == 56:
                continue
            observations.append(dict(
                state="MI", geo_level="state" if kind == 1 else "county",
                geo_id="26" if kind == 1 else f"26{code:03d}", year=year,
                variable="population", value=values[0], datatype="projection",
                scenario="net_positive_migration" if kind == 1 else "",
                source_id=STATE_SOURCE_ID if kind == 1 else COUNTY_SOURCE_ID,
            ))
        if seen != expected:
            raise ValueError(f"MI: incomplete geography/year grid; missing {len(expected - seen)} rows")
        return observations
    finally:
        workbook.close()


def _validate_reference(body: bytes, name: str, marker: str | None) -> None:
    if marker is None:
        fetch.validate_pdf(body, name)
    else:
        text = body.decode("utf-8-sig")
        if "<html" not in text.lower() or marker not in text:
            raise ValueError(f"MI: unexpected reference page {name}")


def _get_reference(url: str) -> bytes:
    # Michigan's public media endpoint rejects the shared ingest User-Agent.
    # A browser-compatible header works; format validation still runs below.
    request = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urlopen(request, timeout=60) as response:
        return response.read()


def acquire() -> None:
    """Archive originals only after validation; reuse this saved edition offline."""
    if (RAW_DIR / "manifest.json").exists():
        entries = fetch.read_manifest(RAW_DIR)
        required = {FILE} | {name for name, _, _ in REFERENCES}
        if not required <= {e["file"] for e in entries}:
            raise ValueError("MI: incomplete saved manifest")
        parse((RAW_DIR / FILE).read_bytes())
        for name, _, marker in REFERENCES:
            _validate_reference((RAW_DIR / name).read_bytes(), name, marker)
        print("MI acquire: reused validated Vintage 2024 snapshots")
        return
    body = fetch.get_bytes(URL)
    rows = parse(body)
    downloads = [(FILE, URL, body, len(rows))]
    for name, url, marker in REFERENCES:
        reference = _get_reference(url)
        _validate_reference(reference, name, marker)
        downloads.append((name, url, reference, None))
    entries = []
    for name, url, body, count in downloads:
        fetch.write_snapshot(RAW_DIR / name, body)
        entries.append(dict(
            file=name, source=url, retrieved_at=date.today().isoformat(),
            sha256=fetch.sha256(body), records=count,
            note="Original MCDA Vintage 2024 workbook or supporting reference. Workbook: 564 source rows; 504 state/county totals served, 60 prosperity-region rows retained only in the original.",
        ))
    fetch.write_manifest(RAW_DIR, entries)
    print(f"MI acquire: {len(rows)} state/county totals and four reference documents")


def load(conn) -> None:
    entries = fetch.read_manifest(RAW_DIR)
    retrieved = next(e["retrieved_at"] for e in entries if e["file"] == FILE)
    rows = parse((RAW_DIR / FILE).read_bytes())
    actual = {(r["geo_level"], r["geo_id"]) for r in conn.execute(
        "SELECT geo_level, geo_id FROM geographies WHERE state='MI'"
    )}
    if actual != {(r["geo_level"], r["geo_id"]) for r in rows}:
        raise ValueError("MI: seeded geographies do not match the source geography set")
    shared_notes = (
        "Vintage 2024 population projections, published state series April 2024 and county release March 2025. "
        "Only the workbook's five-year points (2025-2050) are served; all are projections, including 2025. "
        "No estimates or intermediate years are generated. Fractional model totals retain source precision. "
        "County sums differ from published state totals by at most 0.03 persons; original values are retained. "
        "The workbook's age/sex detail and ten prosperity regions are archived but not served. "
    )
    for source_id, level, base, method in (
        (STATE_SOURCE_ID, "state", 2022,
         "The state report (page 4) explicitly identifies 2022 as the last historical estimate year, "
         "with projections beginning in 2023. The download supplies only 2025 onward. "
         "The main net-positive-migration scenario is loaded; the report's illustrative zero-migration "
         f"scenario is not loaded. State report: {REPORT_URL}"),
        (COUNTY_SOURCE_ID, "county", None,
         "The county cohort-component methodology applies state trends to county fertility, survival and migration. "
         "It models change from 2023 but does not explicitly identify the county population base year. "
         "County base_year is therefore NULL; the explicitly documented statewide 2022 base is not silently borrowed. "
         f"County methodology: {COUNTY_METHOD_URL}"),
    ):
        upsert_source(
            conn, source_id=source_id, state="MI", agency=AGENCY,
            product_name=f"Michigan {level.title()} Population Projections through 2050",
            vintage="2024", base_year=base, horizon_year=2050,
            source_url=URL, retrieved_at=retrieved, license=LICENSE,
            citation="Michigan Center for Data and Analytics. Vintage 2024 Population Projections through 2050.",
            method_notes=shared_notes + method,
        )
    count = insert_observations(conn, rows)
    conn.commit()
    print(f"MI load: {count} observations")


FIXTURES = [
    ("26", 2025, 10112318),
    ("26", 2050, 9905991.62),
    ("26081", 2050, 763022.75),
    ("26083", 2050, 1562.25),
    ("26147", 2050, 144502.45),
    ("26149", 2050, 59891.7),
    ("26163", 2050, 1565187.79),
]


def verify(conn) -> None:
    for gid, year, value in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype FROM observations WHERE state='MI' AND geo_id=? "
            "AND year=? AND variable='population'", (gid, year),
        ).fetchone()
        assert row is not None and tuple(row) == (value, "projection"), f"MI fixture: {gid}/{year}"
    shape = conn.execute(
        "SELECT COUNT(*), COUNT(DISTINCT geo_id), MIN(year), MAX(year), "
        "SUM(datatype='estimate') FROM observations WHERE state='MI'"
    ).fetchone()
    assert tuple(shape) == (504, 84, 2025, 2050, 0)
    print(f"MI verify: {len(FIXTURES)} fixtures and complete state/county grid ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.mi --acquire")
