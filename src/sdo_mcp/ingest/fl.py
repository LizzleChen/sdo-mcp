"""Florida adapter.

Source: BEBR (Bureau of Economic and Business Research, University of
Florida), producer of the official county projections under continuous state
contract (s. 186.901 F.S.); calibrated to the state forecast adopted by the
FDEC conference convened by EDR. 2026 edition (Feb 2026): base April 1 2025
estimate, projections 2030-2050 in FIVE-YEAR steps, Low/Medium/High
scenarios (BEBR recommends Medium). Verified 2026-08-07 (plan/RESEARCH.md).

Florida quirks handled here:
  - the sheet layout is parser-hostile: each county block is a county row
    carrying only the April 1 2025 estimate, followed by three indented
    scenario rows carrying the projections; names are uppercase with
    trailing whitespace
  - county scenario values are rounded to the nearest 100 (county sums do
    not equal the state row); the state estimate is unrounded
  - no FIPS anywhere — names joined against the seeded Census county list
  - the 2025 estimate precedes the scenario split and is stored under all
    three scenarios so every scenario series carries its base point
  - BEBR URLs are vintage-stamped (/wp-content/uploads/YYYY/MM/) — a new
    edition means a new URL, so a stale one fails loudly rather than
    silently serving an old vintage

State-level households and employment forecasts exist only as conference
PDFs (FDEC packet, EDR long-run tables) — coverage marks them pdf_only and
nothing is ingested from them.
"""

from __future__ import annotations

import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, upsert_source, insert_observations

RAW_DIR = REPO_ROOT / "data" / "raw" / "fl"
STATE_FIPS = "12"

# Vintage-stamped URL for the 2026 edition (verified 2026-08-07).
URL = "https://bebr.ufl.edu/wp-content/uploads/2026/02/projections_2026.xlsx"
FILE = "projections_2026.xlsx"

BASE_YEAR = 2025
PROJ_YEARS = (2030, 2035, 2040, 2045, 2050)
SCENARIOS = ("low", "medium", "high")
SOURCE_ID = "fl-pop-2026"

AGENCY = (
    "Bureau of Economic and Business Research (BEBR), University of Florida — "
    "official county projections under s. 186.901 F.S.; state forecast adopted "
    "by the Demographic Estimating Conference convened by EDR"
)
CITATION = (
    "Bureau of Economic and Business Research, University of Florida. "
    "Projections of Florida Population by County, 2030-2050 (2026 edition, "
    "February 2026). Retrieved from bebr.ufl.edu."
)
LICENSE = (
    "BEBR free tier; no license text published (explicit-permission email "
    "open, non-blocking). EDR materials are Florida public records."
)


def acquire() -> None:
    body = fetch.get_bytes(URL)
    if not body.startswith(b"PK"):
        raise RuntimeError(
            f"{URL}: body is not an xlsx — likely a new edition at a new "
            "vintage-stamped URL; update URL from bebr.ufl.edu/population/"
        )
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    (RAW_DIR / FILE).write_bytes(body)
    fetch.write_manifest(RAW_DIR, [{
        "file": FILE, "source": URL,
        "retrieved_at": date.today().isoformat(),
        "sha256": fetch.sha256(body), "records": None,
        "note": "county L-M-H projections; county scenario values rounded to "
                "nearest 100, state estimate unrounded",
    }])
    print(f"FL acquire: {FILE} ({len(body):,} bytes)")


def _name_key(name: str) -> str:
    return name.strip().lower().replace(" ", "")


def load(conn) -> None:
    """Snapshot -> observations + sources. No network."""
    import openpyxl

    fips_by_key = {
        _name_key(row["name"].removesuffix(" County")): row["geo_id"]
        for row in conn.execute(
            "SELECT geo_id, name FROM geographies WHERE state='FL' AND geo_level='county'"
        )
    }

    wb = openpyxl.load_workbook(RAW_DIR / FILE, read_only=True)
    sheet_rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()

    rows: list[dict] = []
    level = gid = None
    blocks = 0
    for rec in sheet_rows[2:]:
        label = str(rec[0] or "").strip()
        if not label:
            continue
        if label.lower() in SCENARIOS:
            assert gid is not None, f"FL load: scenario row before any county row"
            for year, val in zip(PROJ_YEARS, rec[3:8]):
                rows.append(dict(state="FL", geo_level=level, geo_id=gid,
                                 year=year, variable="population", value=int(val),
                                 datatype="projection", scenario=label.lower(),
                                 source_id=SOURCE_ID))
        else:
            # County (or FLORIDA) row: carries only the April 1 2025 estimate.
            if label.upper() == "FLORIDA":
                level, gid = "state", STATE_FIPS
            else:
                level = "county"
                gid = fips_by_key.get(_name_key(label))
                if gid is None:
                    raise RuntimeError(f"FL load: county {label!r} has no Census FIPS match")
            blocks += 1
            estimate = int(rec[1])
            # The base estimate precedes the scenario split; store it under
            # every scenario so each series carries its base point.
            for scen in SCENARIOS:
                rows.append(dict(state="FL", geo_level=level, geo_id=gid,
                                 year=BASE_YEAR, variable="population",
                                 value=estimate, datatype="estimate",
                                 scenario=scen, source_id=SOURCE_ID))
    if blocks != 68:
        raise RuntimeError(f"FL load: {blocks} blocks parsed, expected 67 counties + state")

    upsert_source(
        conn, source_id=SOURCE_ID, state="FL", agency=AGENCY,
        product_name="Projections of Florida Population by County (Low/Medium/High)",
        vintage="2026 edition (February 2026)",
        base_year=BASE_YEAR, horizon_year=PROJ_YEARS[-1],
        source_url=URL,
        retrieved_at=next(e for e in fetch.read_manifest(RAW_DIR)
                          if e["file"] == FILE)["retrieved_at"],
        license=LICENSE, citation=CITATION,
        method_notes=(
            "Base April 1 2025 estimate; projections in five-year steps "
            "2030-2050 (no annual series). BEBR recommends Medium. County "
            "scenario values are rounded to the nearest 100, so county sums do "
            "not equal the state series; the state estimate is unrounded. "
            "Vintage skew: this edition calibrates to the December 2025 FDEC "
            "conference; the June 2026 conference later revised the state "
            "forecast down — figures trace to the December 2025 adoption."
        ),
    )
    n = insert_observations(conn, rows)
    conn.commit()
    print(f"FL load: {n} observations")


# Known published values read off the snapshot (recorded at acquisition).
FIXTURES = [
    # (geo_level, geo_id, scenario, year, value, datatype)
    ("state",  "12",    "medium", 2025, 23379261, "estimate"),
    ("state",  "12",    "medium", 2050, 28250100, "projection"),
    ("state",  "12",    "low",    2050, 24860100, "projection"),
    ("state",  "12",    "high",   2050, 31640100, "projection"),
    ("county", "12001", "low",    2030, 295400,   "projection"),  # Alachua
    ("county", "12001", "high",   2050, 410700,   "projection"),
    ("county", "12003", "medium", 2035, 31500,    "projection"),  # Baker
]


def verify(conn) -> None:
    for level, gid, scen, year, value, datatype in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype FROM observations WHERE state='FL' AND"
            " geo_level=? AND geo_id=? AND variable='population' AND year=?"
            " AND scenario=? AND age_group='all' AND sex='all' AND race='all'",
            (level, gid, year, scen),
        ).fetchone()
        assert row is not None, f"FL verify: missing {level} {gid} {scen} {year}"
        assert int(row["value"]) == value and row["datatype"] == datatype, (
            f"FL verify: {level} {gid} {scen} {year} -> "
            f"{row['value']}/{row['datatype']}, expected {value}/{datatype}"
        )

    for scen in SCENARIOS:
        span = conn.execute(
            "SELECT COUNT(DISTINCT geo_id) n, COUNT(DISTINCT year) y"
            " FROM observations WHERE state='FL' AND geo_level='county'"
            " AND scenario=?", (scen,),
        ).fetchone()
        assert span["n"] == 67, f"FL {scen} has {span['n']} counties, expected 67"
        assert span["y"] == 6, f"FL {scen} has {span['y']} years, expected 6"

    # Scenario ordering must hold everywhere: low <= medium <= high.
    bad = conn.execute(
        "SELECT lo.geo_id, lo.year FROM observations lo"
        " JOIN observations md ON md.state=lo.state AND md.geo_id=lo.geo_id"
        "  AND md.year=lo.year AND md.variable=lo.variable AND md.scenario='medium'"
        " JOIN observations hi ON hi.state=lo.state AND hi.geo_id=lo.geo_id"
        "  AND hi.year=lo.year AND hi.variable=lo.variable AND hi.scenario='high'"
        " WHERE lo.state='FL' AND lo.scenario='low' AND lo.datatype='projection'"
        "  AND NOT (lo.value <= md.value AND md.value <= hi.value)"
    ).fetchall()
    assert not bad, f"FL verify: scenario ordering violated at {[tuple(r) for r in bad[:5]]}"

    print(f"FL verify: {len(FIXTURES)} round-trip values + spans + scenario ordering ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.fl --acquire")
