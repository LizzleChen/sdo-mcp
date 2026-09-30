"""South Carolina RFA Vintage 2023 state and county population totals."""

from __future__ import annotations

import io
import math
import sys
from datetime import date
from decimal import Decimal, ROUND_HALF_UP

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "sc"
CATALOG_URL = "https://www.rfa.sc.gov/programs/census-state-liaison-and-data-center/search-and-download-population-data/rfa-population-projections"
OVERVIEW_URL = "https://www.rfa.sc.gov/programs/census-state-liaison-and-data-center/population-projections"
TOTALS_URL = ("https://www.rfa.sc.gov/sites/rfa/files/2026-08/Vintage%202023%20RFA%20Population%20"
              "Projections%20-%20Total%20Population%2C%20Population%20Change%2C%20and%20"
              "Components%20of%20Change%20by%20Sex%20and%20County%202024-2043.xlsx")
AGE5_URL = ("https://www.rfa.sc.gov/sites/rfa/files/2026-08/Vintage%202023%20RFA%20Population%20"
            "Projections%20by%205-Year%20Age%20Group%20and%20County%202024-2043.xlsx")
FILES = {
    "projection_catalog_2026-09-30.html": CATALOG_URL,
    "projection_overview_2026-09-30.html": OVERVIEW_URL,
    "population_components_v2023.xlsx": TOTALS_URL,
    "age5_v2023.xlsx": AGE5_URL,
}
SOURCE_ID = "sc-pop-v2023"
AGENCY = "South Carolina Revenue and Fiscal Affairs Office"
YEARS = tuple(range(2023, 2044))


def _counties() -> dict[str, str]:
    counties = {}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        fields = line.split("|")
        if fields[0] == "SC":
            counties[fields[1] + fields[2]] = fields[4].removesuffix(" County")
    if len(counties) != 46:
        raise ValueError("SC: expected 46 Census counties")
    return counties


def _workbook(body: bytes, expected_sheets: list[str]):
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("SC: response is not a readable XLSX workbook") from exc
    if book.sheetnames != expected_sheets:
        book.close()
        raise ValueError("SC: unexpected workbook sheets")
    return book


def parse_totals(body: bytes) -> list[dict]:
    book = _workbook(body, ["TOTAL POPULATION", "FEMALES", "MALES"])
    try:
        sheet = book.worksheets[0]
        grid = list(sheet.iter_rows(values_only=True))
        if len(grid) != 93 or any(len(row) != 103 for row in grid):
            raise ValueError("SC: unexpected total-population table layout")
        if (grid[0][0] != "VINTAGE 2023 SOUTH CAROLINA POPULATION PROJECTIONS"
                or grid[1][0] != "TOTAL POPULATION, POPULATION CHANGE, AND COMPONENTS OF CHANGE BY SEX AND COUNTY"
                or grid[2][0] != "2024-2043"
                or grid[4][0] != "Population Base Period: 2019-2023"
                or "Vintage 2023 Postcensal Estimates for 2020-2023" not in grid[5][0]
                or grid[10][:2] != ("COUNTY", "FIPS")
                or grid[11][2] != "2023 POPULATION"):
            raise ValueError("SC: unexpected edition, estimate base or year header")
        for year in range(2024, 2044):
            group = 2 if year == 2024 else 8 + 5 * (year - 2025)
            column = 7 + 5 * (year - 2024)
            if (grid[10][group] != f"{year-1} - {year}"
                    or not str(grid[11][column]).startswith(f"{year} POPULATION")):
                raise ValueError(f"SC: unexpected {year} period or population header")
        counties = _counties()
        seen = set()
        rows = []
        totals = {}
        for source in grid[12:59]:
            name, code = source[:2]
            gid = str(code)
            if gid == "45":
                if name != "SOUTH CAROLINA":
                    raise ValueError("SC: invalid state row")
            elif gid not in counties or name != counties[gid].upper()[:9]:
                raise ValueError(f"SC: unexpected county name or FIPS: {name}/{code}")
            if gid in seen:
                raise ValueError(f"SC: duplicate geography {gid}")
            seen.add(gid)
            values = (source[2], *(source[7 + 5 * i] for i in range(20)))
            if any(not isinstance(v, (int, float)) or isinstance(v, bool)
                   or not math.isfinite(v) or v < 0 or int(v) != v for v in values):
                raise ValueError(f"SC: invalid population value for {gid}")
            totals[gid] = tuple(map(int, values))
            rows.extend({
                "state": "SC", "geo_level": "state" if gid == "45" else "county",
                "geo_id": gid, "year": year, "variable": "population", "value": int(value),
                "datatype": "estimate" if year == 2023 else "projection",
                "source_id": SOURCE_ID,
            } for year, value in zip(YEARS, values))
        if seen != set(counties) | {"45"}:
            raise ValueError("SC: incomplete state/county population grid")
        for i, year in enumerate(YEARS):
            delta = totals["45"][i] - sum(values[i] for gid, values in totals.items() if gid != "45")
            if abs(delta) > 5:
                raise ValueError(f"SC: county/state difference exceeds published rounding in {year}")
        return rows
    finally:
        book.close()


def parse_age5_totals(body: bytes) -> dict[tuple[str, int], int]:
    """Companion all-ages cells independently validate the displayed totals."""
    book = _workbook(body, ["Vintage 2023 RFA Population Pro"])
    try:
        sheet = book.active
        grid = list(sheet.iter_rows(values_only=True))
        if (len(grid) != 937
                or grid[0][0] != "VINTAGE 2023 SOUTH CAROLINA POPULATION PROJECTIONS"
                or grid[1][0] != "BY COUNTY, 5-YEAR AGE GROUPS, AND SEX"
                or grid[2][0] != "2024-2043"
                or tuple(grid[9][3:23]) != tuple(f"T{year}" for year in range(2024, 2044))):
            raise ValueError("SC: unexpected companion age-group edition or years")
        counties = _counties()
        result = {}
        for row in grid[10:]:
            if row[2] != "ALL AGES":
                continue
            gid = str(row[0])
            if gid == "45":
                expected_name = "SOUTH CAROLINA"
            else:
                expected_name = counties.get(gid, "").upper()[:9]
            if not expected_name or row[1] != expected_name:
                raise ValueError(f"SC: unexpected companion geography {gid}")
            for year, value in zip(range(2024, 2044), row[3:23]):
                if not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
                    raise ValueError(f"SC: invalid companion value for {gid}/{year}")
                key = (gid, year)
                if key in result:
                    raise ValueError(f"SC: duplicate companion all-ages row {gid}/{year}")
                result[key] = int(Decimal(str(value)).quantize(1, rounding=ROUND_HALF_UP))
        if len(result) != 47 * 20:
            raise ValueError("SC: incomplete companion all-ages grid")
        return result
    finally:
        book.close()


def _validate_page(name: str, body: bytes) -> None:
    page = " ".join(body.decode("utf-8", "replace").lower().split())
    markers = (("<html", "vintage 2023", "2024-2043.xlsx") if name.startswith("projection_catalog")
               else ("<html", "revenue and fiscal affairs office", "annual state and county population projections"))
    for marker in markers:
        if marker not in page:
            raise ValueError(f"SC: invalid publisher page; missing {marker!r}")


def _validate_all(bodies: dict[str, bytes]) -> list[dict]:
    for name in FILES:
        if name.endswith(".html"):
            _validate_page(name, bodies[name])
    rows = parse_totals(bodies["population_components_v2023.xlsx"])
    companion = parse_age5_totals(bodies["age5_v2023.xlsx"])
    for row in rows:
        if row["year"] >= 2024 and companion[(row["geo_id"], row["year"])] != row["value"]:
            raise ValueError(f"SC: publisher workbooks disagree at {row['geo_id']}/{row['year']}")
    return rows


def acquire() -> None:
    if (RAW_DIR / "manifest.json").exists():
        entries = {e["file"] for e in fetch.read_manifest(RAW_DIR)}
        if entries != set(FILES):
            raise ValueError("SC: incomplete saved snapshot manifest")
        _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
        print("SC acquire: reused validated snapshots")
        return
    bodies = {name: fetch.get_bytes(url) for name, url in FILES.items()}
    _validate_all(bodies)
    entries = []
    for name, url in FILES.items():
        body = bodies[name]
        fetch.write_snapshot(RAW_DIR / name, body)
        entries.append({"file": name, "source": url, "retrieved_at": date.today().isoformat(),
                        "sha256": fetch.sha256(body), "records": 987 if name == "population_components_v2023.xlsx" else None,
                        "note": "Original RFA Vintage 2023 publication; the 2026-08 URL folder is not the model vintage."})
    fetch.write_manifest(RAW_DIR, entries)


def load(conn) -> None:
    entries = {e["file"]: e for e in fetch.read_manifest(RAW_DIR)}
    rows = _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
    upsert_source(
        conn, source_id=SOURCE_ID, state="SC", agency=AGENCY,
        product_name="RFA state and county population projections", vintage="Vintage 2023",
        base_year=2023, horizon_year=2043, source_url=TOTALS_URL,
        retrieved_at=entries["population_components_v2023.xlsx"]["retrieved_at"],
        license="No explicit data license identified; cite the South Carolina RFA.",
        citation=f"{AGENCY}. Vintage 2023 Population Projections, 2024-2043.",
        method_notes=(
            "The workbook identifies a 2019-2023 population base period and Vintage 2023 "
            "postcensal estimates through 2023. The 2023 point is an estimate; 2024-2043 are "
            "annual projections. State and all 46 county totals are taken from the TOTAL POPULATION "
            "sheet and independently checked against all-ages values rounded as displayed in "
            "the separate age-group workbook. County sums can differ from the state by up to five "
            "people due to rounding; the published state row is preserved. The totals workbook "
            "also publishes natural change and net migration, but separate births/deaths totals "
            "and interval semantics have not been established for the MCP components tool. "
            "The 2026-08 URL directory is a storage path, not a revised model vintage."
        ),
    )
    insert_observations(conn, rows)
    conn.commit()
    print(f"SC load: {len(rows)} observations")


def verify(conn) -> None:
    fixtures = {("45", 2023): (5373555, "estimate"),
                ("45", 2024): (5445291, "projection"),
                ("45", 2043): (6870048, "projection"),
                ("45001", 2024): (24442, "projection"),
                ("45005", 2024): (7166, "projection")}
    for (gid, year), expected in fixtures.items():
        row = conn.execute("SELECT value, datatype FROM observations WHERE "
                           "state='SC' AND variable='population' AND geo_id=? AND year=?",
                           (gid, year)).fetchone()
        assert tuple(row) == expected
    shape = conn.execute("SELECT COUNT(*),COUNT(DISTINCT geo_id),COUNT(DISTINCT year),"
                         "MIN(year),MAX(year) FROM observations WHERE state='SC'").fetchone()
    assert tuple(shape) == (987, 47, 21, 2023, 2043)
    print("SC verify: cross-workbook fixtures and complete annual grid ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.sc --acquire")
