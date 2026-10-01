"""Alaska September 2026 resident population: state scenarios and county totals."""

from __future__ import annotations

import io
import math
import re
import sys
from datetime import date, datetime

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "ak"
CATALOG_URL = "https://live.laborstats.alaska.gov/article/alaska-population-projections"
STATE_URL = "https://live.laborstats.alaska.gov/sites/default/files/2026-09/StateProjections.xlsx"
COUNTY_URL = "https://live.laborstats.alaska.gov/sites/default/files/2026-09/Boroughs%20and%20Census%20Areas.xlsx"
REPORT_URL = "https://live.laborstats.alaska.gov/pop/projections/pub/popproj.pdf"
GEOGRAPHY_URL = "https://tigerweb.geo.census.gov/tigerwebmain/Files/acs26/tigerweb_acs26_county_2020_tab20_ak.html"
CATALOG_FILE = "projection_catalog_2026-09-30.html"
STATE_FILE = "state_projections_2025-2055_sep2026.xlsx"
COUNTY_FILE_NAME = "borough_census_area_projections_2025-2055_sep2026.xlsx"
REPORT_FILE = "population_projections_2025-2055_sep2026.pdf"
GEOGRAPHY_FILE = "census_2020_county_geographies_2026-09-30.html"
FILES = {CATALOG_FILE: CATALOG_URL, STATE_FILE: STATE_URL, COUNTY_FILE_NAME: COUNTY_URL,
         REPORT_FILE: REPORT_URL, GEOGRAPHY_FILE: GEOGRAPHY_URL}
SOURCE_ID = "ak-pop-sep2026"
AGENCY = "Alaska Department of Labor and Workforce Development, Research and Analysis Section"
YEARS = tuple(range(2025, 2056))
COUNTY_YEARS = tuple(range(2025, 2056, 5))
SCENARIOS = ("low", "mid", "high")
FOOTER = f"Source: {AGENCY}"
ALIASES = {
    "Anchorage, Municipality of": "Anchorage Municipality",
    "Anchorage, Municipality": "Anchorage Municipality",
    "Skagway Borough, Municipality of": "Skagway Municipality",
    "Skagway, Municipality": "Skagway Municipality",
    **{f"{name}, City and Borough{suffix}": f"{name} City and Borough"
       for name in ("Juneau", "Sitka", "Wrangell", "Yakutat") for suffix in ("", " of")},
}


def _counties() -> dict[str, str]:
    counties = {}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        parts = line.split("|")
        if parts[0] == "AK":
            counties[parts[4]] = parts[1] + parts[2]
    if len(counties) != 30 or counties.get("Chugach Census Area") != "02063" or counties.get("Copper River Census Area") != "02066":
        raise ValueError("AK: expected all 30 current Census county equivalents")
    return counties


def _number(value, context: str) -> int:
    if (not isinstance(value, (int, float)) or isinstance(value, bool)
            or not math.isfinite(value) or value < 0 or int(value) != value):
        raise ValueError(f"AK: invalid population value at {context}")
    return int(value)


def _book(body: bytes):
    import openpyxl
    try:
        return openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("AK: response is not a readable XLSX workbook") from exc


def _age_totals(sheet, expected_title: str) -> dict[int, int]:
    """Cross-check published five-year totals within the age/sex panels."""
    if sheet["A1"].value != expected_title:
        raise ValueError("AK: unexpected age/sex table title")
    totals = {}
    for year, (r, c) in zip(COUNTY_YEARS, ((3, 1), (3, 8), (29, 1), (29, 8), (55, 1), (55, 8), (81, 1))):
        label = f"July 1, {year} {'Estimate' if year == 2025 else 'Projected'}"
        if sheet.cell(r, c).value != label or sheet.cell(r + 23, c).value != "Total":
            raise ValueError("AK: unexpected age/sex year or estimate boundary")
        total = _number(sheet.cell(r + 23, c + 1).value, f"{sheet.title}/{year}")
        male = _number(sheet.cell(r + 23, c + 2).value, f"{sheet.title}/{year}/male")
        female = _number(sheet.cell(r + 23, c + 3).value, f"{sheet.title}/{year}/female")
        if total != male + female:
            raise ValueError("AK: age/sex totals do not reconcile")
        totals[year] = total
    return totals


def parse_state_workbook(body: bytes) -> list[dict]:
    book = _book(body)
    try:
        if (book.sheetnames != ["TableOfContents", "ProjComponentsOfChange", "AlaskaMiddle", "AlaskaLow", "AlaskaHigh"]
                or book.properties.modified != datetime(2026, 9, 15, 22, 51, 31)):
            raise ValueError("AK: unexpected statewide workbook edition or sheets")
        grid = list(book["ProjComponentsOfChange"].values)
        if (len(grid) != 44 or len(grid[0]) != 29
                or grid[0][0] != "Annual Components of Population Change for Alaska, 2025 to 2055"
                or grid[3][2] != "End of Period" or grid[4][2] != "Population"
                or grid[5][2:5] != ("Low", "Middle", "High") or grid[-1][0] != FOOTER):
            raise ValueError("AK: unexpected statewide population grid")
        rows = []
        totals = {}
        for year, row in zip(YEARS, grid[7:38]):
            if row[0] != f"{year - 1}-{str(year)[-2:]}":
                raise ValueError("AK: unexpected annual population year grid")
            for scenario, value in zip(SCENARIOS, row[2:5]):
                value = _number(value, f"{scenario}/{year}")
                totals[scenario, year] = value
                rows.append({"state": "AK", "geo_level": "state", "geo_id": "02", "year": year,
                             "variable": "population", "value": value,
                             "datatype": "estimate" if year == 2025 else "projection",
                             "scenario": scenario, "source_id": SOURCE_ID})
        for scenario, display in (("mid", "Middle"), ("low", "Low"), ("high", "High")):
            checks = _age_totals(book[f"Alaska{display}"],
                f"Alaska Population by Age and Sex, and Components of Change: 2025 to 2055, {display} Scenario")
            if any(totals[scenario, year] != value for year, value in checks.items()):
                raise ValueError("AK: annual and five-year state totals disagree")
        if any(totals["low", y] > totals["mid", y] or totals["mid", y] > totals["high", y] for y in YEARS):
            raise ValueError("AK: statewide scenario order is invalid")
        return rows
    finally:
        book.close()


def parse_county_workbook(body: bytes) -> list[dict]:
    book = _book(body)
    try:
        counties = _counties()
        if (len(book.sheetnames) != 39 or book.sheetnames[:3] != ["TableOfContents", "ProjectionsByArea", "Alaska"]
                or book.properties.modified != datetime(2026, 9, 15, 23, 26, 17)):
            raise ValueError("AK: unexpected county workbook edition or sheets")
        grid = list(book["ProjectionsByArea"].values)
        if (len(grid) != 51 or len(grid[0]) != 211 or grid[0][0] != "Alaska Population by Area: 2025 to 2055"
                or grid[2][1:14:2] != tuple(f"July 1, {year}" for year in COUNTY_YEARS)
                or grid[3][1:14:2] != ("Estimate",) + ("Projection",) * 6 or grid[-1][0] != FOOTER):
            raise ValueError("AK: unexpected county years or estimate boundary")
        rows, totals = [], {}
        for row in grid[5:48]:
            label = row[0]
            if not isinstance(label, str) or (label.strip() != "Alaska" and not label.startswith("   ")):
                continue  # Economic regions are separate published geography, not counties.
            name = ALIASES.get(label.strip(), label.strip())
            gid = "02" if name == "Alaska" else counties.get(name)
            if gid is None or gid in totals:
                raise ValueError(f"AK: unknown or duplicate county geography {label!r}")
            totals[gid] = tuple(_number(v, f"{gid}/{y}") for y, v in zip(COUNTY_YEARS, row[1:14:2]))
            if gid != "02":
                rows.extend({"state": "AK", "geo_level": "county", "geo_id": gid, "year": year,
                             "variable": "population", "value": value,
                             "datatype": "estimate" if year == 2025 else "projection", "source_id": SOURCE_ID}
                            for year, value in zip(COUNTY_YEARS, totals[gid]))
        if set(totals) != set(counties.values()) | {"02"}:
            raise ValueError("AK: incomplete county geography grid")
        for i, year in enumerate(COUNTY_YEARS):
            if totals["02"][i] != sum(values[i] for gid, values in totals.items() if gid != "02"):
                raise ValueError(f"AK: county totals do not reconcile in {year}")
        return rows
    finally:
        book.close()


def parse_report(body: bytes) -> dict[tuple[str, int, str], int]:
    """Parse independent PDF Tables 1.4 and 3.1, including every loaded point."""
    from pypdf import PdfReader
    try:
        report = PdfReader(io.BytesIO(body))
        if len(report.pages) != 103 or "Published September 2026" not in report.pages[1].extract_text():
            raise ValueError("unexpected report edition")
        values = {}
        state_text = report.pages[11].extract_text()
        if "Population Projections for Alaska, 2025 to 2055" not in state_text:
            raise ValueError("missing statewide PDF table")
        for line in state_text.splitlines():
            match = re.match(r"^(\d{4})-(\d{2})\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)(?:\s|$)", line)
            if match:
                year = int(match[1]) + 1
                if match[2] != str(year)[-2:] or year not in YEARS:
                    raise ValueError("unexpected statewide PDF year")
                for scenario, token in zip(SCENARIOS, match.groups()[2:]):
                    key = ("02", year, scenario)
                    if key in values:
                        raise ValueError("duplicate statewide PDF year")
                    values[key] = int(token.replace(",", ""))
        counties = _counties()
        for page_number, years in ((26, COUNTY_YEARS[:4]), (27, COUNTY_YEARS[4:])):
            text = report.pages[page_number].extract_text()
            if "Alaska Population by Region, Borough, and Census Area, 2025 to 2055" not in text:
                raise ValueError("missing county PDF table")
            seen = set()
            for line in text.splitlines():
                match = re.match(r"^(.+?)\s+([\d,]+(?:\s+[\d,]+)+)\s*$", line)
                if not match:
                    continue
                label = match[1].strip()
                name = ALIASES.get(label, label)
                gid = counties.get(name)
                if gid is None:
                    continue
                numbers = tuple(int(n.replace(",", "")) for n in match[2].split())
                if gid in seen or len(numbers) != len(years):
                    raise ValueError("duplicate or incomplete county PDF row")
                seen.add(gid)
                values.update({(gid, year, ""): value for year, value in zip(years, numbers)})
            if seen != set(counties.values()):
                raise ValueError("incomplete PDF county grid")
        if len(values) != 303:
            raise ValueError("incomplete PDF state/county grid")
        return values
    except Exception as exc:
        raise ValueError("AK: response is not the expected population report PDF") from exc


def _validate_all(bodies: dict[str, bytes]) -> list[dict]:
    catalog = bodies[CATALOG_FILE].decode("utf-8", "replace").lower()
    if any(marker not in catalog for marker in ("<html", "2025 to 2055", "stateprojections.xlsx", "boroughs%20and%20census%20areas.xlsx")):
        raise ValueError("AK: invalid publisher catalog")
    geography = bodies[GEOGRAPHY_FILE].decode("utf-8", "replace")
    if any(marker not in geography for marker in ("<html", "02063", "02066", "Chugach Census Area", "Copper River Census Area")):
        raise ValueError("AK: invalid Census county geography evidence")
    rows = parse_state_workbook(bodies[STATE_FILE]) + parse_county_workbook(bodies[COUNTY_FILE_NAME])
    checks = parse_report(bodies[REPORT_FILE])
    if any(checks[r["geo_id"], r["year"], r.get("scenario", "")] != r["value"] for r in rows):
        raise ValueError("AK: workbook and independent PDF disagree")
    return rows


def acquire() -> None:
    if (RAW_DIR / "manifest.json").exists():
        if {e["file"] for e in fetch.read_manifest(RAW_DIR)} != set(FILES):
            raise ValueError("AK: incomplete snapshot manifest")
        _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
        print("AK acquire: reused validated snapshots")
        return
    bodies = {name: fetch.get_bytes(url) for name, url in FILES.items()}
    _validate_all(bodies)
    entries = []
    for name, url in FILES.items():
        fetch.write_snapshot(RAW_DIR / name, bodies[name])
        entries.append({"file": name, "source": url, "retrieved_at": date.today().isoformat(),
                        "sha256": fetch.sha256(bodies[name]),
                        "records": {STATE_FILE: 93, COUNTY_FILE_NAME: 210}.get(name),
                        "note": "Original Alaska September 2026 edition; state annual scenarios and county five-year totals; Census evidence verifies current county equivalents."})
    fetch.write_manifest(RAW_DIR, entries)


def load(conn) -> None:
    manifest = {e["file"]: e for e in fetch.read_manifest(RAW_DIR)}
    rows = _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
    upsert_source(conn, source_id=SOURCE_ID, state="AK", agency=AGENCY,
                  product_name="Alaska state and borough/census-area resident population projections",
                  vintage="September 2026", base_year=2025, horizon_year=2055, source_url=CATALOG_URL,
                  retrieved_at=manifest[STATE_FILE]["retrieved_at"],
                  license="No explicit data license identified; cite Alaska Research and Analysis.",
                  citation=f"{AGENCY}. Alaska Population Projections 2025 to 2055. September 2026.",
                  method_notes=("Cohort-component resident population. July 1, 2025 estimates launch the model; later years are projections. "
                                "State Low/Middle/High migration scenarios are annual; mid maps to the publisher's Middle. "
                                "Borough/census-area totals are published at five-year intervals in a single series aligned by the publisher to "
                                "the statewide Middle scenario; no county Low/High series is manufactured. All 30 current county equivalents "
                                "match the saved Census 2020 list, including Chugach 02063 and Copper River 02066; Valdez-Cordova is not substituted. "
                                "Every loaded value agrees with independent PDF Tables 1.4 and 3.1. Regional, age/sex, Alaska Native and "
                                "components detail is published but not loaded. No household or eligible SDO jobs publication was established."))
    insert_observations(conn, rows)
    conn.commit()
    print(f"AK load: {len(rows)} observations")


def verify(conn) -> None:
    shape = conn.execute("SELECT COUNT(*),COUNT(DISTINCT geo_id),MIN(year),MAX(year) FROM observations WHERE state='AK'").fetchone()
    assert tuple(shape) == (303, 31, 2025, 2055)
    for gid, year, scenario, expected in (("02", 2025, "mid", 738737), ("02", 2026, "mid", 739652),
                                         ("02", 2055, "low", 566586), ("02", 2055, "mid", 724936),
                                         ("02", 2055, "high", 925038), ("02063", 2055, "", 5623),
                                         ("02066", 2025, "", 2628), ("02170", 2055, "", 159086)):
        row = conn.execute("SELECT value,datatype FROM observations WHERE state='AK' AND geo_id=? AND year=? AND scenario=?",
                           (gid, year, scenario)).fetchone()
        assert tuple(row) == (expected, "estimate" if year == 2025 else "projection")
    print("AK verify: complete scenarios/current county grid and independent PDF totals ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.ak --acquire")
