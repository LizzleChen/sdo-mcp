"""Washington adapter.

The Office of Financial Management (OFM) publishes two distinct population
products that must never be joined into one series:

* State Population Forecast, November 2025: state only, 1970-2050, with
  estimates through 2025 and projections from 2026.
* 2022 Growth Management Act (GMA) projections: all 39 counties, Low/Medium/
  High scenarios, estimates for 2020-2022 and projections for 2025-2050.

Both were verified 2026-08-10 (plan/EXPANSION.md). The GMA workbook includes
a statewide comparison row, but this adapter deliberately does not load it;
the newer state forecast is the statewide product of record.
"""

from __future__ import annotations

import re
import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source

RAW_DIR = REPO_ROOT / "data" / "raw" / "wa"
STATE_FIPS = "53"

STATE_FILE = "stfc_2025.xlsx"
STATE_URL = "https://ofm.wa.gov/wp-content/uploads/stfc_2025.xlsx"
STATE_SOURCE_ID = "wa-pop-2025-stfc"
STATE_BASE_YEAR = 2025

GMA_FILE = "gma_2022_1yr.xlsx"
GMA_URL = (
    "https://ofm.wa.gov/wp-content/uploads/sites/default/files/public/"
    "dataresearch/pop/GMA/projections2022/gma_2022_1yr.xlsx"
)
GMA_SOURCE_ID = "wa-pop-2022-gma"
GMA_BASE_YEAR = 2022
SCENARIO_SHEETS = {"Low": "low", "Middle": "medium", "High": "high"}

AGENCY = "Washington State Office of Financial Management, Forecasting and Research Division"
LICENSE = "OFM permits use, sharing, and copying with attribution; liability disclaimer applies."

STATE_COMPONENT_FILE = "stfc_2025_components.xlsx"
COUNTY_COMPONENT_FILE = "gma_2022_components_med.xlsx"
COMPONENT_FILES = (
    (STATE_COMPONENT_FILE, "https://ofm.wa.gov/wp-content/uploads/" + STATE_COMPONENT_FILE,
     {"Components of Change"}),
    (COUNTY_COMPONENT_FILE, GMA_URL.rsplit("/", 1)[0] + "/" + COUNTY_COMPONENT_FILE,
     {"Population Change", "Births", "Deaths", "Residual Net Migration"}),
)


def acquire_components() -> None:
    """Archive the two component workbooks without replacing existing snapshots."""
    import io
    import openpyxl

    for name, url, sheets in COMPONENT_FILES:
        body = fetch.get_bytes(url)
        wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        try:
            if set(wb.sheetnames) != sheets:
                raise ValueError(f"WA components: unexpected sheets in {name}")
        finally:
            wb.close()
        fetch.write_snapshot(RAW_DIR / name, body)
        fetch.write_manifest(RAW_DIR, [dict(
            file=name, source=url, retrieved_at=date.today().isoformat(),
            sha256=fetch.sha256(body), records=None,
            note="Original OFM components workbook; annual state / five-year medium county totals.",
        )])


def _whole(value, context: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
        raise ValueError(f"WA components: invalid count at {context}: {value!r}")
    return int(value)


def load_components(conn, counties: dict[str, str]) -> None:
    """Store period totals by ending year; never divide county totals into years."""
    import openpyxl

    fetch.read_manifest(RAW_DIR)
    rows = []
    for name, url, sheets in COMPONENT_FILES:
        state = name == STATE_COMPONENT_FILE
        source_id = "wa-components-2025-stfc" if state else "wa-components-2022-gma"
        wb = openpyxl.load_workbook(RAW_DIR / name, read_only=True, data_only=True)
        try:
            if set(wb.sheetnames) != sheets:
                raise ValueError(f"WA components: unexpected sheets in {name}")
            if state:
                values = list(wb["Components of Change"].values)
                if values[2][0] != "November 2025" or "2026-2050" not in str(values[4][0]):
                    raise ValueError("WA state components: unexpected edition or forecast boundary")
                if tuple(values[8][i] for i in (0, 4, 6, 9)) != ("Period", "Births", "Deaths", "Net Migration"):
                    raise ValueError("WA state components: unexpected columns")
                years = []
                for rec in values[9:]:
                    match = re.fullmatch(r"(\d{4})-(\d{4})", str(rec[0]))
                    if not match:
                        continue
                    start, end = map(int, match.groups())
                    if end - start != 1:
                        continue  # Published multi-year subtotals overlap the annual rows.
                    years.append(end)
                    births, deaths, migration = [_whole(rec[i], str(rec[0])) for i in (4, 6, 9)]
                    if min(births, deaths) < 0:
                        raise ValueError("WA state components: negative births/deaths")
                    if births - deaths + migration != rec[2]:
                        raise ValueError(f"WA state component balance differs in {end}")
                    pop = conn.execute(
                        "SELECT value FROM observations WHERE source_id=? AND variable='population' AND year=?",
                        (STATE_SOURCE_ID, end),
                    ).fetchone()
                    if pop is None or pop[0] != rec[1]:
                        raise ValueError(f"WA state component population differs in {end}")
                    for variable, value in zip(("births", "deaths", "net_migration"), (births, deaths, migration)):
                        rows.append(dict(state="WA", geo_level="state", geo_id=STATE_FIPS,
                                         year=end, variable=variable, value=value, source_id=source_id,
                                         datatype="estimate" if end <= 2025 else "projection"))
                if years != list(range(2011, 2051)):
                    raise ValueError("WA state components: unexpected annual periods")
            else:
                tables = {}
                for sheet in sheets:
                    values = list(wb[sheet].values)
                    if values[1][0] != "2022 GMA Projections - Medium Series" or values[2][14] != "Projection":
                        raise ValueError("WA county components: unexpected edition/scenario/boundary")
                    columns = [(i, y) for i, y in enumerate(range(1965, 2025, 5), 1)]
                    columns += [(i, y) for i, y in enumerate(range(2025, 2055, 5), 14)]
                    for i, year in columns:
                        if values[3][i] != f"{year-5}-{year % 100:02d}":
                            raise ValueError("WA county components: unexpected five-year periods")
                    table = {}
                    for rec in values[4:]:
                        if rec[0] in counties:
                            if rec[0] in table:
                                raise ValueError("WA county components: duplicate county")
                            table[rec[0]] = {year: _whole(rec[i], f"{sheet}/{rec[0]}/{year}") for i, year in columns}
                    if set(table) != set(counties):
                        raise ValueError("WA county components: missing counties")
                    tables[sheet] = table
                for county, gid in counties.items():
                    for year in range(1965, 2055, 5):
                        b, d, m, change = [tables[s][county][year] for s in
                                          ("Births", "Deaths", "Residual Net Migration", "Population Change")]
                        if min(b, d) < 0:
                            raise ValueError("WA county components: negative births/deaths")
                        if b - d + m != change:
                            raise ValueError(f"WA county component balance differs: {county}/{year}")
                        for variable, value in zip(("births", "deaths", "net_migration"), (b, d, m)):
                            rows.append(dict(state="WA", geo_level="county", geo_id=gid,
                                             year=year, variable=variable, value=value, source_id=source_id,
                                             scenario="medium", datatype="estimate" if year <= 2020 else "projection"))
        finally:
            wb.close()
        upsert_source(conn, source_id=source_id, state="WA", agency=AGENCY,
                      product_name="Components of April 1 Population Change" if state else
                      "2022 GMA County Components of Population Change, Medium Series",
                      vintage="November 2025" if state else "2022 GMA projections (March 2023 component tables)",
                      base_year=2025 if state else 2020, horizon_year=2050,
                      source_url=url, retrieved_at=_retrieved(name), license=LICENSE,
                      citation=f"{AGENCY}. {name}.",
                      method_notes=(
                          "Whole-person totals for April 1 to April 1 periods, indexed by ending year. "
                          + ("Annual periods 2010-2011 through 2049-2050; estimates through 2025. "
                             "Overlapping multi-year subtotals are excluded." if state else
                             "Five-year totals, not annual values: 1960-1965 through 2045-2050. "
                             "Medium scenario only. Residual net migration is population change minus "
                             "births plus deaths. OFM cautions against treating the last digit as precise. "
                             "The workbook labels periods ending through 2020 as estimates and 2020-2025 "
                             "onward as projections; base_year=2020 describes this component table, "
                             "not the separate county population product's 2022 base. "
                             "The statewide comparison row is excluded.")))
    insert_observations(conn, rows)


def acquire() -> None:
    """Download and validate both official population workbooks."""
    import io
    import openpyxl

    entries = []
    for name, url, required_sheets in (
        (STATE_FILE, STATE_URL, {"Single Year", "Five Year"}),
        (GMA_FILE, GMA_URL, set(SCENARIO_SHEETS)),
    ):
        body = fetch.get_bytes(url)
        if not body.startswith(b"PK"):
            raise RuntimeError(f"{url}: body is not an xlsx (HTML error page?)")
        wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        try:
            if not required_sheets <= set(wb.sheetnames):
                raise RuntimeError(f"{url}: unexpected workbook sheets {wb.sheetnames!r}")
        finally:
            wb.close()
        RAW_DIR.mkdir(parents=True, exist_ok=True)
        fetch.write_snapshot(RAW_DIR / name, body)
        entries.append({
            "file": name,
            "source": url,
            "retrieved_at": date.today().isoformat(),
            "sha256": fetch.sha256(body),
            "records": None,
            "note": (
                "State forecast; estimates through 2025, projections 2026-2050."
                if name == STATE_FILE else
                "2022 GMA county Low/Medium/High series; estimates 2020-2022, "
                "projections 2025-2050 with no 2023-2024 values."
            ),
        })
        print(f"WA acquire: {name} ({len(body):,} bytes)")
    fetch.write_manifest(RAW_DIR, entries)


def _retrieved(filename: str) -> str:
    return next(
        e["retrieved_at"]
        for e in fetch.read_manifest(RAW_DIR)
        if e["file"] == filename
    )


def load(conn) -> None:
    """Two workbook snapshots -> two strictly separated population products."""
    import openpyxl

    fips_by_name = {
        row["name"].removesuffix(" County"): row["geo_id"]
        for row in conn.execute(
            "SELECT geo_id, name FROM geographies WHERE state='WA' AND geo_level='county'"
        )
    }
    rows: list[dict] = []

    wb = openpyxl.load_workbook(RAW_DIR / STATE_FILE, read_only=True, data_only=True)
    try:
        values = list(wb["Single Year"].iter_rows(values_only=True))
        header = values[8]
        total = next(r for r in values if r[0] == "Total")
        state_points = []
        for index, label in enumerate(header):
            match = re.fullmatch(r"Total (\d{4})", str(label))
            if not match:
                continue
            year = int(match.group(1))
            state_points.append(year)
            rows.append(dict(
                state="WA", geo_level="state", geo_id=STATE_FIPS, year=year,
                variable="population", value=int(total[index]),
                datatype="estimate" if year <= STATE_BASE_YEAR else "projection",
                source_id=STATE_SOURCE_ID,
            ))
        assert state_points == list(range(1970, 2051)), (
            f"WA state years are {state_points[:2]}...{state_points[-2:]}, expected 1970-2050"
        )
    finally:
        wb.close()

    wb = openpyxl.load_workbook(RAW_DIR / GMA_FILE, read_only=True, data_only=True)
    try:
        gma_state: dict[tuple[str, int], int] = {}
        expected_years = [2020, 2021, 2022, *range(2025, 2051)]
        for sheet_name, scenario in SCENARIO_SHEETS.items():
            values = list(wb[sheet_name].iter_rows(values_only=True))
            header = values[4]
            year_columns = [(i, int(v)) for i, v in enumerate(header) if isinstance(v, int)]
            assert [y for _i, y in year_columns] == expected_years, (
                f"WA GMA {sheet_name} years differ from the published sparse series"
            )
            names = set()
            for rec in values[5:]:
                name = str(rec[0] or "").strip()
                if name == "State":
                    for index, year in year_columns:
                        gma_state[(scenario, year)] = int(round(float(rec[index])))
                    continue
                if name not in fips_by_name:
                    continue
                names.add(name)
                for index, year in year_columns:
                    rows.append(dict(
                        state="WA", geo_level="county", geo_id=fips_by_name[name],
                        year=year, variable="population",
                        value=int(round(float(rec[index]))),
                        datatype="estimate" if year <= GMA_BASE_YEAR else "projection",
                        scenario=scenario, source_id=GMA_SOURCE_ID,
                    ))
            if names != set(fips_by_name):
                raise RuntimeError(
                    f"WA load: {sheet_name} county mismatch; missing "
                    f"{sorted(set(fips_by_name) - names)}"
                )

        # These comparison values are published in the GMA workbook but are
        # not inserted. The newer state forecast remains the only state series.
        assert gma_state[("low", 2050)] == 8550172
        assert gma_state[("medium", 2050)] == 9937575
        assert gma_state[("high", 2050)] == 11853278
    finally:
        wb.close()

    upsert_source(
        conn,
        source_id=STATE_SOURCE_ID,
        state="WA",
        agency=AGENCY,
        product_name="Forecast of the State Population by Age and Sex",
        vintage="November 2025 (report published March 2026)",
        base_year=STATE_BASE_YEAR,
        horizon_year=2050,
        source_url=STATE_URL,
        retrieved_at=_retrieved(STATE_FILE),
        license=LICENSE,
        citation=(
            "Washington State Office of Financial Management. Forecast of the "
            "State Population by Age and Sex, November 2025."
        ),
        method_notes=(
            "State-only series. Decennial and intercensal values run through "
            "2020, postcensal estimates through 2025, and forecasts from 2026 "
            "through 2050. Totals come from the workbook's published Total row."
        ),
    )
    upsert_source(
        conn,
        source_id=GMA_SOURCE_ID,
        state="WA",
        agency=AGENCY,
        product_name="2022 Growth Management Act County Population Projections",
        vintage="2022 GMA projections (December 2022 workbook)",
        base_year=GMA_BASE_YEAR,
        horizon_year=2050,
        source_url=GMA_URL,
        retrieved_at=_retrieved(GMA_FILE),
        license=LICENSE,
        citation=(
            "Washington State Office of Financial Management. 2022 Growth "
            "Management Act County Population Projections, Low/Middle/High series."
        ),
        method_notes=(
            "County-only Low/Medium/High series. The workbook labels 2020 as "
            "adjusted Census, 2021-2022 as estimates, and 2025-2050 as "
            "projections; it publishes no 2023-2024 values. Its statewide "
            "comparison row is deliberately not loaded because it is a different "
            "vintage and product from OFM's November 2025 state forecast."
        ),
    )
    n = insert_observations(conn, rows)
    load_components(conn, fips_by_name)
    conn.commit()
    print(f"WA load: {n} observations")


FIXTURES = [
    ("state", "53", "", 2025, 8115100, "estimate", STATE_SOURCE_ID),
    ("state", "53", "", 2050, 9719108, "projection", STATE_SOURCE_ID),
    ("county", "53033", "medium", 2020, 2269675, "estimate", GMA_SOURCE_ID),
    ("county", "53033", "medium", 2025, 2377737, "projection", GMA_SOURCE_ID),
    ("county", "53033", "medium", 2050, 2879176, "projection", GMA_SOURCE_ID),
]


def verify(conn) -> None:
    for level, gid, scenario, year, value, datatype, source_id in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype, source_id FROM observations WHERE state='WA' "
            "AND geo_level=? AND geo_id=? AND variable='population' AND scenario=? AND year=?",
            (level, gid, scenario, year),
        ).fetchone()
        assert row is not None, f"WA verify: missing {level} {gid} {scenario} {year}"
        assert (int(row["value"]), row["datatype"], row["source_id"]) == (
            value, datatype, source_id
        )

    state_shape = conn.execute(
        "SELECT MIN(year) lo, MAX(year) hi, COUNT(*) n, "
        "COUNT(DISTINCT source_id) ns FROM observations WHERE state='WA' "
        "AND geo_level='state' AND variable='population'"
    ).fetchone()
    assert tuple(state_shape) == (1970, 2050, 81, 1)

    county_shape = conn.execute(
        "SELECT COUNT(*) n, COUNT(DISTINCT geo_id) ng, COUNT(DISTINCT year) ny, "
        "COUNT(DISTINCT source_id) ns FROM observations WHERE state='WA' "
        "AND geo_level='county' AND variable='population'"
    ).fetchone()
    assert tuple(county_shape) == (3393, 39, 29, 1)

    gap = conn.execute(
        "SELECT COUNT(*) n FROM observations WHERE state='WA' AND geo_level='county' "
        "AND year IN (2023, 2024)"
    ).fetchone()["n"]
    assert gap == 0, f"WA verify: found {gap} fabricated 2023-2024 county points"

    bad = conn.execute(
        "SELECT lo.geo_id, lo.year FROM observations lo JOIN observations md ON "
        "md.state=lo.state AND md.geo_id=lo.geo_id AND md.year=lo.year "
        "AND md.variable=lo.variable AND md.scenario='medium' JOIN observations hi ON "
        "hi.state=lo.state AND hi.geo_id=lo.geo_id AND hi.year=lo.year "
        "AND hi.variable=lo.variable AND hi.scenario='high' WHERE lo.state='WA' "
        "AND lo.geo_level='county' AND lo.variable='population' AND lo.scenario='low' "
        "AND NOT (lo.value<=md.value AND md.value<=hi.value)"
    ).fetchall()
    assert not bad, f"WA verify: scenario ordering violated: {[tuple(r) for r in bad[:5]]}"
    print(f"WA verify: {len(FIXTURES)} fixtures + product separation + sparse years + scenarios ok")


if __name__ == "__main__":
    if "--acquire-components" in sys.argv:
        acquire_components()
    elif "--acquire" in sys.argv:
        acquire()
        acquire_components()
    else:
        print("usage: python -m sdo_mcp.ingest.wa [--acquire | --acquire-components]")
