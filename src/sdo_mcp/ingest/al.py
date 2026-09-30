"""Alabama CBER/State Data Center 2025-based county population projections."""

from __future__ import annotations

import io
import math
import sys
import warnings
import zipfile
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from xml.etree import ElementTree

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "al"
CATALOG_URL = "https://cber.culverhouse.ua.edu/resources/alabama-demographics/"
PUBLIC_URL = "https://alabama.box.com/s/991v89cfqpkukoa2o7c6eh2wr57elo1t"
DOWNLOAD_URL = ("https://alabama.app.box.com/index.php?rm=box_download_shared_file"
                "&shared_name=991v89cfqpkukoa2o7c6eh2wr57elo1t&file_id=f_1667589139124")
CATALOG_FILE = "demographics_catalog_2026-09-30.html"
WORKBOOK_FILE = "county_population_2025-2050_fall2026.xlsx"
SOURCE_ID = "al-pop-2025-fall2026"
AGENCY = "University of Alabama Center for Business and Economic Research, Alabama State Data Center"
YEARS = (2025, 2030, 2035, 2040, 2045, 2050)
TITLE = ("Alabama's County Decennial Census (1990-2020), Population Estimates "
         "(2021-2025), and Projections (2030-2050)")
SOURCE_LINE = ("Source:  U.S. Census Bureau and Center for Business and Economic Research, "
               "The University of Alabama, October 2024")


def _counties() -> dict[str, str]:
    counties = {}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        fields = line.split("|")
        if fields[0] == "AL":
            counties[fields[4].removesuffix(" County")] = fields[1] + fields[2]
    if len(counties) != 67:
        raise ValueError("AL: expected 67 Census counties")
    return counties


def _modified_at(body: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            core = ElementTree.fromstring(archive.read("docProps/core.xml"))
        modified = core.find("{http://purl.org/dc/terms/}modified")
        return modified.text if modified is not None else ""
    except Exception as exc:
        raise ValueError("AL: workbook has no readable edition metadata") from exc


def _displayed_person(value) -> int:
    if (not isinstance(value, (int, float)) or isinstance(value, bool)
            or not math.isfinite(value) or value < 0):
        raise ValueError(f"AL: invalid population value {value!r}")
    return int(Decimal(str(value)).quantize(1, rounding=ROUND_HALF_UP))


def parse_workbook(body: bytes) -> list[dict]:
    import openpyxl

    try:
        with warnings.catch_warnings():
            # The source's decorative sparkline is unsupported by openpyxl;
            # all population cells and their display formats still parse.
            warnings.filterwarnings("ignore", message="Sparkline Group extension is not supported.*")
            book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("AL: response is not a readable XLSX workbook") from exc
    try:
        if book.sheetnames != ["Projections"]:
            raise ValueError("AL: unexpected workbook sheet")
        sheet = book.active
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Sparkline Group extension is not supported.*")
            grid = list(sheet.iter_rows(values_only=True))
            formats = list(sheet.iter_rows(min_row=4, max_row=71, min_col=10, max_col=15))
        if len(grid) != 74 or any(len(row) != 29 for row in grid):
            raise ValueError("AL: unexpected population table layout")
        if (grid[0][0] != TITLE
                or grid[1][5] != "Population Estimates"
                or grid[1][10] != "Population Projections"
                or tuple(grid[2][9:15]) != YEARS
                or grid[71][0] != SOURCE_LINE
                or _modified_at(body) != "2026-09-24T21:40:11Z"):
            raise ValueError("AL: unexpected workbook title, year grid or edition metadata")
        counties = _counties()
        names = [row[0] for row in grid[4:71]]
        if len(set(names)) != 67 or set(names) != set(counties):
            raise ValueError("AL: missing or duplicate county geography")
        if grid[3][0] != "Alabama":
            raise ValueError("AL: missing state total")
        rows = []
        raw_values = {}
        for source, cells in zip(grid[3:71], formats):
            name = source[0]
            gid = "01" if name == "Alabama" else counties[name]
            if source[28] != name:
                raise ValueError(f"AL: duplicate geography label disagrees for {name}")
            values = source[9:15]
            if any(cell.number_format != "#,##0" for cell in cells):
                raise ValueError(f"AL: unexpected displayed population precision for {name}")
            displayed = tuple(_displayed_person(value) for value in values)
            if abs(values[0] - displayed[0]) > 1e-6:
                raise ValueError(f"AL: 2025 base must be an integer estimate for {name}")
            raw_values[gid] = values
            rows.extend({
                "state": "AL", "geo_level": "state" if gid == "01" else "county",
                "geo_id": gid, "year": year, "variable": "population", "value": value,
                "datatype": "estimate" if year == 2025 else "projection",
                "source_id": SOURCE_ID,
            } for year, value in zip(YEARS, displayed))
        for i, year in enumerate(YEARS):
            if abs(raw_values["01"][i] - sum(values[i] for gid, values in raw_values.items() if gid != "01")) > 1e-5:
                raise ValueError(f"AL: raw county values do not reconcile to the state in {year}")
        return rows
    finally:
        book.close()


def _validate_catalog(body: bytes) -> None:
    page = " ".join(body.decode("utf-8", "replace").lower().split())
    for marker in ("<html", "alabama state data center", "2025 to 2050 county population projections",
                   "991v89cfqpkukoa2o7c6eh2wr57elo1t", "cohort-component model"):
        if marker not in page:
            raise ValueError(f"AL: invalid demographics catalog; missing {marker!r}")


def acquire() -> None:
    if (RAW_DIR / "manifest.json").exists():
        entries = {e["file"] for e in fetch.read_manifest(RAW_DIR)}
        if entries != {CATALOG_FILE, WORKBOOK_FILE}:
            raise ValueError("AL: incomplete saved snapshot manifest")
        _validate_catalog((RAW_DIR / CATALOG_FILE).read_bytes())
        parse_workbook((RAW_DIR / WORKBOOK_FILE).read_bytes())
        print("AL acquire: reused validated snapshots")
        return
    catalog = fetch.get_bytes(CATALOG_URL)
    _validate_catalog(catalog)
    workbook = fetch.get_bytes(DOWNLOAD_URL)
    rows = parse_workbook(workbook)
    entries = []
    for filename, url, body, count in ((CATALOG_FILE, CATALOG_URL, catalog, None),
                                       (WORKBOOK_FILE, DOWNLOAD_URL, workbook, len(rows))):
        fetch.write_snapshot(RAW_DIR / filename, body)
        entries.append({"file": filename, "source": url, "retrieved_at": date.today().isoformat(),
                        "sha256": fetch.sha256(body), "records": count,
                        "note": "Original CBER catalog or Box workbook; the workbook footer's October 2024 credit conflicts with its September 2026 modification date and Fall 2026 filename."})
    fetch.write_manifest(RAW_DIR, entries)


def load(conn) -> None:
    entries = {e["file"]: e for e in fetch.read_manifest(RAW_DIR)}
    _validate_catalog((RAW_DIR / CATALOG_FILE).read_bytes())
    rows = parse_workbook((RAW_DIR / WORKBOOK_FILE).read_bytes())
    upsert_source(
        conn, source_id=SOURCE_ID, state="AL", agency=AGENCY,
        product_name="Alabama state and county population projections", vintage="Fall 2026 file (2025 estimate base)",
        base_year=2025, horizon_year=2050, source_url=PUBLIC_URL,
        retrieved_at=entries[WORKBOOK_FILE]["retrieved_at"],
        license="No explicit data license identified; cite the University of Alabama CBER.",
        citation=f"{AGENCY}. County Population Projections, 2025-2050. Fall 2026 file.",
        method_notes=(
            "The CBER catalog identifies its own cohort-component county projection product. The "
            "workbook labels 2021-2025 as estimates and 2030-2050 as projections; only its 2025 "
            "estimate and five-year projections are loaded. The Box filename says ReleasedFall2026 "
            "and workbook metadata records modification on September 24, 2026, while its footer "
            "still credits October 2024. Fall 2026 is used as the file edition, with that stale "
            "footer disclosed. The source stores fractional model values but displays whole-person "
            "counts with #,##0 formatting. Served values follow the displayed rounding. The "
            "published state row is retained, so rounded county sums can differ by up to three. "
            "No health-planning agency forecast is substituted."
        ),
    )
    insert_observations(conn, rows)
    conn.commit()
    print(f"AL load: {len(rows)} observations")


def verify(conn) -> None:
    fixtures = {("01", 2025): (5193088, "estimate"),
                ("01", 2030): (5365127, "projection"),
                ("01", 2050): (5890544, "projection"),
                ("01003", 2025): (267761, "estimate"),
                ("01003", 2050): (436428, "projection"),
                ("01063", 2050): (4115, "projection"),
                ("01131", 2050): (5321, "projection")}
    for (gid, year), expected in fixtures.items():
        row = conn.execute("SELECT value, datatype FROM observations WHERE "
                           "state='AL' AND variable='population' AND geo_id=? AND year=?",
                           (gid, year)).fetchone()
        assert tuple(row) == expected
    shape = conn.execute("SELECT COUNT(*),COUNT(DISTINCT geo_id),COUNT(DISTINCT year),"
                         "MIN(year),MAX(year) FROM observations WHERE state='AL'").fetchone()
    assert tuple(shape) == (408, 68, 6, 2025, 2050)
    print("AL verify: estimate/projection boundary, displayed values and grid ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.al --acquire")
