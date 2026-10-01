"""Maine OSE April 2025 population: the published 2022/2027/.../2042 grid."""

from __future__ import annotations

import io
import math
import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "me"
BASE_URL = "https://www11.maine.gov/dafs/economist"
CATALOG_URL = BASE_URL + "/demographic-projections"
DATA_URL = BASE_URL + "/sites/maine.gov.dafs.economist/files/inline-files/"
WORKBOOK_URL = DATA_URL + "MaineStateCountyPopulationProjections2042.xlsx"
PDF_URL = DATA_URL + "MaineStateCountyPopulationProjections2042.pdf"
REPORT_URL = DATA_URL + "Maine%20Population%20Outlook%20to%202032.pdf"
PLACE_URL = DATA_URL + "MaineCityTownPopulationProjection2042.xlsx"
ECONOMIC_URL = "https://www1.maine.gov/dafs/economist/economic-forecasting"
CATALOG_FILE = "projection_catalog_2026-09-30.html"
WORKBOOK_FILE = "population_apr2025.xlsx"
PDF_FILE = "population_apr2025.pdf"
REPORT_FILE = "population_outlook_apr2025.pdf"
PLACE_FILE = "municipal_population_apr2025.xlsx"
ECONOMIC_FILE = "economic_catalog_2026-09-30.html"
FILES = {CATALOG_FILE: CATALOG_URL, WORKBOOK_FILE: WORKBOOK_URL,
         PDF_FILE: PDF_URL, REPORT_FILE: REPORT_URL, PLACE_FILE: PLACE_URL,
         ECONOMIC_FILE: ECONOMIC_URL}
SOURCE_ID = "me-pop-apr2025"
AGENCY = "Maine Department of Administrative and Financial Services, Office of the State Economist"
YEARS = (2022, 2027, 2032, 2037, 2042)
AGES = tuple(f"{age}-{age + 4}" for age in range(0, 85, 5)) + ("85+", "Total")


def _counties() -> dict[str, str]:
    counties = {}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        fields = line.split("|")
        if fields[0] == "ME":
            counties[fields[4].removesuffix(" County")] = fields[1] + fields[2]
    if len(counties) != 16:
        raise ValueError("ME: expected 16 Census counties")
    return counties


def _number(value, location: str) -> float:
    if type(value) not in (int, float) or not math.isfinite(value) or value < 0:
        raise ValueError(f"ME: invalid population value at {location}")
    return value


def parse_workbook(body: bytes) -> list[dict]:
    """Keep cached published totals and precision; reject incomplete/changed grids."""
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("ME: response is not a readable XLSX workbook") from exc
    try:
        counties = _counties()
        sheets = [gid[-2:] for gid in counties.values()]
        if book.sheetnames != sheets + ["ME", "summary"]:
            raise ValueError("ME: unexpected population workbook sheets")
        summary = list(book["summary"].iter_rows(values_only=True))
        if (len(summary) != 29 or any(len(row) != 12 for row in summary)
                or summary[0][1:3] != ("Observed", "Projections")
                or summary[1][1:6] != YEARS
                or summary[1][6:10] != ("2022-2027", "2027-2032", "2032-2037", "2037-2042")
                or any(v is not None for row in summary[19:] for v in row)):
            raise ValueError("ME: unexpected year grid or estimate/projection header")
        geos = counties | {"Maine": "23"}
        seen = set()
        totals = {}
        observations = []
        for source in summary[2:19]:
            name = source[0]
            if name not in geos or name in seen:
                raise ValueError(f"ME: unknown or duplicate geography {name!r}")
            seen.add(name)
            gid = geos[name]
            sheet = book["ME" if name == "Maine" else gid[-2:]]
            if (sheet["B1"].value != ("Maine" if name == "Maine" else name.upper())
                    or (name != "Maine" and sheet["A1"].value != int(gid[-3:]))
                    or tuple(c.value for c in sheet[2][2:21]) != AGES
                    or tuple(sheet[f"A{r}"].value for r in (3, 19, 35)) != ("Female", "Male", "Total")):
                raise ValueError(f"ME: mismatched geography or age/sex header for {name}")
            totals[name] = tuple(_number(v, name) for v in source[1:6])
            for index, year in enumerate(YEARS):
                published = totals[name][index]
                # The summary and age/sex sheets cache the same model totals.
                # Binary floating point addition can differ by <1e-8 persons.
                sex_totals = []
                for row_number in (13 + index, 29 + index, 45 + index):
                    row = next(sheet.iter_rows(min_row=row_number, max_row=row_number,
                                               max_col=21, values_only=True))
                    if row[0] != year:
                        raise ValueError(f"ME: changed age/sex year grid for {name}")
                    cells = [_number(v, f"{name}/{year}") for v in row[2:21]]
                    if not math.isclose(sum(cells[:-1]), cells[-1], rel_tol=0, abs_tol=1e-8):
                        raise ValueError(f"ME: age cells do not reconcile for {name}/{year}")
                    sex_totals.append(cells[-1])
                if (not math.isclose(sum(sex_totals[:2]), sex_totals[2], rel_tol=0, abs_tol=1e-8)
                        or not math.isclose(published, sex_totals[2], rel_tol=0, abs_tol=1e-8)):
                    raise ValueError(f"ME: summary and age/sex totals disagree for {name}/{year}")
                observations.append(dict(
                    state="ME", geo_level="state" if gid == "23" else "county",
                    geo_id=gid, year=year, variable="population", value=published,
                    datatype="estimate" if year == 2022 else "projection", source_id=SOURCE_ID,
                ))
        if seen != set(geos):
            raise ValueError("ME: incomplete state/county population grid")
        for index, year in enumerate(YEARS):
            if not math.isclose(sum(totals[name][index] for name in counties),
                                totals["Maine"][index], rel_tol=0, abs_tol=1e-8):
                raise ValueError(f"ME: county totals do not reconcile in {year}")
        return observations
    finally:
        book.close()


def parse_pdf(body: bytes) -> dict[tuple[str, int], int]:
    """Independently read the publisher's rounded summary PDF table."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(body))
        if len(reader.pages) != 18:
            raise ValueError("unexpected PDF page count")
        lines = reader.pages[-1].extract_text().splitlines()
        if (not lines[0].startswith("Observed")
                or tuple(map(int, lines[1].split()[:5])) != YEARS
                or "Projections" not in lines[-1]):
            raise ValueError("unexpected summary table header")
        geos = _counties() | {"Maine": "23"}
        values = {}
        for line in lines[2:-1]:
            tokens = line.split()
            name = tokens[0]
            if name not in geos or len(tokens) != 12:
                raise ValueError(f"unexpected PDF geography or row {name!r}")
            numbers = [int(token.replace(",", "")) for token in tokens[1:6]]
            for year, value in zip(YEARS, numbers):
                key = geos[name], year
                if key in values or value < 0:
                    raise ValueError(f"duplicate/invalid PDF cell {key}")
                values[key] = value
        if len(values) != 85:
            raise ValueError("incomplete PDF summary table")
        return values
    except Exception as exc:
        raise ValueError("ME: response is not the expected population PDF") from exc


def _validate_reference(body: bytes, name: str) -> None:
    if name in (CATALOG_FILE, ECONOMIC_FILE):
        text = " ".join(body.decode("utf-8", "replace").lower().split())
        markers = (("<html", "mainestatecountypopulationprojections2042.xlsx",
                    "mainecitytownpopulationprojection2042.xlsx", "beyond 2032")
                   if name == CATALOG_FILE else
                   ("<html", "independent group of economists", "wage and salary employment"))
        if any(marker not in text for marker in markers):
            raise ValueError(f"ME: invalid publisher reference {name}")
    elif name == REPORT_FILE:
        from pypdf import PdfReader

        fetch.validate_pdf(body, "ME outlook report")
        reader = PdfReader(io.BytesIO(body))
        text = " ".join(" ".join(p.extract_text() for p in reader.pages).split())
        if any(marker not in text for marker in ("APRIL 2025", "base year", "for the projections is 2022",
                                                  "Office of the State Economist", "cohort-component")):
            raise ValueError("ME: unexpected outlook edition or methodology")
    elif name == PLACE_FILE:
        import openpyxl

        try:
            book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            sheet = book.active
            if (book.sheetnames != ["Sheet1"] or sheet["A1"].value != "Town"
                    or sheet["B1"].value != "County"
                    or tuple(c.value for c in sheet[2][2:9]) != (2012, 2017, *YEARS)
                    or "fewer than 10 people" not in str(sheet["A518"].value)):
                raise ValueError("unexpected municipal grid")
        except Exception as exc:
            raise ValueError("ME: response is not the expected municipal workbook") from exc
        finally:
            if "book" in locals():
                book.close()


def _validate_all(bodies: dict[str, bytes]) -> list[dict]:
    rows = parse_workbook(bodies[WORKBOOK_FILE])
    pdf = parse_pdf(bodies[PDF_FILE])
    for row in rows:
        if abs(pdf[row["geo_id"], row["year"]] - row["value"]) > 0.5:
            raise ValueError("ME: workbook and independently rounded PDF disagree")
    for name in (CATALOG_FILE, REPORT_FILE, PLACE_FILE, ECONOMIC_FILE):
        _validate_reference(bodies[name], name)
    return rows


def acquire() -> None:
    if (RAW_DIR / "manifest.json").exists():
        if set(e["file"] for e in fetch.read_manifest(RAW_DIR)) != set(FILES):
            raise ValueError("ME: incomplete saved snapshot manifest")
        _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
        print("ME acquire: reused validated April 2025 snapshots")
        return
    bodies = {name: fetch.get_bytes(url) for name, url in FILES.items()}
    rows = _validate_all(bodies)
    entries = []
    for name, url in FILES.items():
        body = bodies[name]
        fetch.write_snapshot(RAW_DIR / name, body)
        entries.append(dict(file=name, source=url, retrieved_at=date.today().isoformat(),
                            sha256=fetch.sha256(body), records=len(rows) if name == WORKBOOK_FILE else None,
                            note="Original Maine OSE April 2025 publication or supporting catalog. Fractional workbook totals are retained; the PDF displays nearest-person values. Municipal rows and demographic detail are archived but not served."))
    fetch.write_manifest(RAW_DIR, entries)


def load(conn) -> None:
    entries = {e["file"]: e for e in fetch.read_manifest(RAW_DIR)}
    rows = _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
    upsert_source(
        conn, source_id=SOURCE_ID, state="ME", agency=AGENCY,
        product_name="Maine state and county population projections to 2042",
        vintage="April 2025", base_year=2022, horizon_year=2042,
        source_url=WORKBOOK_URL, retrieved_at=entries[WORKBOOK_FILE]["retrieved_at"],
        license="No explicit data license identified; cite the Maine Office of the State Economist.",
        citation=f"{AGENCY}. Maine Population Outlook 2022 to 2032 and Population Projections to 2042. April 2025.",
        method_notes=(
            "Cohort-component projections with a 2022 historical estimate base; the published grid is "
            "2022, 2027, 2032, 2037 and 2042. No intermediate years are generated. The summary's "
            "cached fractional model totals are retained at workbook precision; they match the "
            "published age/sex totals and the independent PDF's rounded figures. State and 16 "
            "county totals reconcile within binary floating point precision. The publisher replaces "
            "the June 2023 edition and encourages limiting use beyond 2032. Pandemic-adjusted survival "
            "and county-specific migration windows are documented in the outlook report. Published "
            "municipal forecasts and age/sex detail are archived but not served; municipalities with "
            "fewer than 10 projected residents were suppressed by the publisher. Historical component "
            "charts are not future component forecasts. The linked DOL outlook and independently "
            "authored CEFC employment forecast are not substituted for an OSE jobs product."
        ),
    )
    insert_observations(conn, rows)
    conn.commit()
    print(f"ME load: {len(rows)} observations")


def verify(conn) -> None:
    fixtures = {("23", 2022): (1389338, "estimate"),
                ("23", 2027): (1419888.2879359922, "projection"),
                ("23", 2042): (1424856.838353645, "projection"),
                ("23005", 2032): (322972.8624317943, "projection"),
                ("23021", 2042): (16667.90336057443, "projection"),
                ("23031", 2032): (232135.78662340573, "projection")}
    for (gid, year), expected in fixtures.items():
        row = conn.execute("SELECT value,datatype FROM observations WHERE state='ME' "
                           "AND variable='population' AND geo_id=? AND year=?", (gid, year)).fetchone()
        assert tuple(row) == expected, f"ME fixture {gid}/{year}"
    shape = conn.execute("SELECT COUNT(*),COUNT(DISTINCT geo_id),COUNT(DISTINCT year),"
                         "MIN(year),MAX(year),SUM(datatype='estimate') FROM observations WHERE state='ME'").fetchone()
    assert tuple(shape) == (85, 17, 5, 2022, 2042, 17)
    print("ME verify: exact five-year grid, precision and checked PDF anchors ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.me --acquire")
