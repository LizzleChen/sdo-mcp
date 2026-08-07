"""Georgia adapter.

Source: Governor's Office of Planning and Budget (publisher of record);
produced since the 2024 vintage by the Applied Demography Program at UGA's
Carl Vinson Institute of Government. 2026 vintage (released 2026-07-21):
2026-2065 annual, state + 159 counties, Excel only. Verified 2026-08-07
(plan/RESEARCH.md).

Georgia quirks handled here:
  - download URLs are vintage-specific slugs — acquire() scrapes the landing
    page rather than hard-coding them
  - the xlsx are slightly malformed (missing drawing part, wrong declared
    dimensions): openpyxl needs read_only=True + ws.reset_dimensions()
  - county files carry names only, no FIPS — joined against the seeded Census
    county list; any name that doesn't match is a loud ingest error
  - everything OPB publishes is projection: the base comes from Census
    Vintage 2025 estimates, there is no OPB estimates series
"""

from __future__ import annotations

import re
import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, upsert_source, insert_observations

LANDING = "https://opb.georgia.gov/census-data/population-projections"
RAW_DIR = REPO_ROOT / "data" / "raw" / "ga"
STATE_FIPS = "13"

VINTAGE = "2026"
BASE_YEAR = 2025   # Census Vintage 2025 estimates; every served year is projection
YEARS = (2026, 2065)
SOURCE_ID = "ga-pop-2026"

# landing-page slug fragment -> snapshot filename
FILES = {
    "state-residential-population": "state_total.xlsx",
    "county-residential-population": "county_total.xlsx",
}

AGENCY = (
    "Governor's Office of Planning and Budget (OPB); produced by the Applied "
    "Demography Program, Carl Vinson Institute of Government, University of Georgia"
)
CITATION = (
    "Georgia Governor's Office of Planning and Budget / UGA Carl Vinson "
    "Institute of Government. Long-Term Population Projections, 2026 vintage "
    "(2026-2065). Retrieved from opb.georgia.gov."
)
LICENSE = (
    "No stated terms; public data, non-commercial use per project decision "
    "2026-08-07. Cite both OPB and UGA."
)


def acquire() -> None:
    """Scrape the landing page for the current vintage's totals files and
    snapshot them into data/raw/ga/."""
    html = fetch.get_bytes(LANDING).decode("utf-8", "replace")
    hrefs = set(re.findall(r'href="([^"]*?/document/census-data/[^"]*?/download)"', html))
    entries = []
    for slug, name in FILES.items():
        matches = [h for h in hrefs if slug in h]
        if len(matches) != 1:
            raise RuntimeError(f"GA acquire: {len(matches)} landing-page links match {slug!r}")
        url = matches[0]
        if url.startswith("/"):
            url = "https://opb.georgia.gov" + url
        body = fetch.get_bytes(url)
        if not body.startswith(b"PK"):
            raise RuntimeError(f"{url}: body is not an xlsx (soft-200 trap)")
        RAW_DIR.mkdir(parents=True, exist_ok=True)
        (RAW_DIR / name).write_bytes(body)
        entries.append({
            "file": name, "source": url,
            "retrieved_at": date.today().isoformat(),
            "sha256": fetch.sha256(body), "records": None,
            "note": "vintage-specific URL scraped from the landing page",
        })
        print(f"GA acquire: {name} ({len(body):,} bytes)")
    fetch.write_manifest(RAW_DIR, entries)


def _sheet(name: str) -> list[tuple]:
    import openpyxl

    # read_only + reset_dimensions: the files ship with a missing drawing part
    # and wrong declared dimensions (1x1), which break a default load.
    wb = openpyxl.load_workbook(RAW_DIR / name, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        ws.reset_dimensions()
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _retrieved(name: str) -> str:
    entry = next(e for e in fetch.read_manifest(RAW_DIR) if e["file"] == name)
    return entry["retrieved_at"]


def load(conn) -> None:
    """Snapshots -> observations + sources. No network."""
    fips_by_name = {
        row["name"].removesuffix(" County").strip().lower(): row["geo_id"]
        for row in conn.execute(
            "SELECT geo_id, name FROM geographies WHERE state='GA' AND geo_level='county'"
        )
    }

    rows: list[dict] = []

    state = _sheet("state_total.xlsx")
    years = [int(y) for y in state[0]]
    assert (years[0], years[-1]) == YEARS, f"GA state year span {years[0]}-{years[-1]}"
    for year, val in zip(years, state[1]):
        rows.append(dict(state="GA", geo_level="state", geo_id=STATE_FIPS, year=year,
                         variable="population", value=int(val),
                         datatype="projection", source_id=SOURCE_ID))

    county = _sheet("county_total.xlsx")
    years = [int(y) for y in county[0][1:]]
    assert (years[0], years[-1]) == YEARS, f"GA county year span {years[0]}-{years[-1]}"
    unmatched = []
    for rec in county[1:]:
        name = str(rec[0]).strip()
        gid = fips_by_name.get(name.lower())
        if gid is None:
            unmatched.append(name)
            continue
        for year, val in zip(years, rec[1:]):
            rows.append(dict(state="GA", geo_level="county", geo_id=gid, year=year,
                             variable="population", value=int(val),
                             datatype="projection", source_id=SOURCE_ID))
    if unmatched:
        raise RuntimeError(f"GA load: county names with no Census FIPS match: {unmatched}")

    upsert_source(
        conn, source_id=SOURCE_ID, state="GA", agency=AGENCY,
        product_name="Long-Term Population Projections, 2026 vintage",
        vintage=VINTAGE, base_year=BASE_YEAR, horizon_year=YEARS[1],
        source_url=LANDING, retrieved_at=_retrieved("state_total.xlsx"),
        license=LICENSE, citation=CITATION,
        method_notes=(
            "Cohort-component with Gaussian Process Regression for rates; "
            "COVID-era (2020-21) mortality excluded from fitting. Base = Census "
            "Vintage 2025 estimates; OPB publishes no estimates series, so "
            "every served year is a projection."
        ),
    )
    n = insert_observations(conn, rows)
    conn.commit()
    print(f"GA load: {n} observations")


# Known published values recorded in plan/RESEARCH.md at verification time.
FIXTURES = [
    ("state",  "13",    2026, 11422803),
    ("county", "13001", 2026, 19287),     # Appling County
]


def verify(conn) -> None:
    for level, gid, year, value in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype FROM observations WHERE state='GA' AND"
            " geo_level=? AND geo_id=? AND variable='population' AND year=?"
            " AND scenario='' AND age_group='all' AND sex='all' AND race='all'",
            (level, gid, year),
        ).fetchone()
        assert row is not None, f"GA verify: missing {level} {gid} {year}"
        assert int(row["value"]) == value and row["datatype"] == "projection", (
            f"GA verify: {level} {gid} {year} -> {row['value']}/{row['datatype']},"
            f" expected {value}/projection"
        )

    span = conn.execute(
        "SELECT MIN(year) lo, MAX(year) hi, COUNT(DISTINCT geo_id) n"
        " FROM observations WHERE state='GA' AND geo_level='county'"
    ).fetchone()
    assert (span["lo"], span["hi"]) == YEARS, f"GA span {span['lo']}-{span['hi']}"
    assert span["n"] == 159, f"GA has {span['n']} counties, expected 159"
    print(f"GA verify: {len(FIXTURES)} round-trip values + spans ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.ga --acquire")
