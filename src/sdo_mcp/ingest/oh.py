"""Ohio Research Office's original 2023-edition population workbook.

Serve only its published Total_TOT cells; age/sex columns and two publisher
PDFs provide independent checks. The current catalog needs a browser capture.
"""
from __future__ import annotations

import argparse
import io
import json
import math
import re
from datetime import date, datetime
from pathlib import Path

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "oh"
CATALOG_URL = "https://development.ohio.gov/about-us/research/population"
URL = "https://development.ohio.gov/static/research/pop1/Projections-by-County.xlsx"
DOWNLOAD_URL = "https://dam.assets.ohio.gov/raw/upload/development.ohio.gov/research/pop1/Projections-by-County.xlsx"
REPORT_URL = "https://dam.assets.ohio.gov/image/upload/development.ohio.gov/research/pop1/State-of-Ohio-Population-Projections-Report-2023.pdf"
OVERVIEW_URL = "https://dam.assets.ohio.gov/image/upload/development.ohio.gov/research/pop1/State-of-Ohio-Population-Projections-Overview-2020-to-2050.pdf"
REPORT_LINK = "https://development.ohio.gov/static/research/pop1/State-of-Ohio-Population-Projections-Report-2023.pdf"
FILE = "population_2020-2050_report2023.xlsx"
REPORT_FILE = "population_report_2023.pdf"
OVERVIEW_FILE = "population_overview_2020-2050.pdf"
CATALOG_FILE = "projection_catalog_browser_excerpt_2026-09-30.json"
FILES = {FILE: DOWNLOAD_URL, REPORT_FILE: REPORT_URL, OVERVIEW_FILE: OVERVIEW_URL, CATALOG_FILE: CATALOG_URL}
SOURCE_ID = "oh-pop-2023"
AGENCY = "Ohio Department of Development, Office of Research"
VINTAGE = "2023 edition"
YEARS = tuple(range(2020, 2051, 5))
AGES = tuple(f"{age}to{age+4}" for age in range(0, 85, 5)) + ("85+",)
HEADERS = ("Area", "Year") + tuple(label for sex in ("M", "F", "TOT")
    for label in (*(f"{age}_{sex}" for age in AGES), f"Total_{sex}")) + (None,)
# The state's historical sex-age estimates differ from its published both-sex
# age estimates. Every other row reconciles by age; all 623 total cells reconcile.
STATE_2020_AGE_DIFFERENCES = (-106, -71, 2, 154, 746, 501, 349, 340, 292,
                            105, -112, -200, -809, -383, -551, -185, -65, -7)
# Independently published in report Table 1 and overview page 2.
AGE_ANCHORS = (("0-to-14", 2170245, 1987475), ("15-to-24", 1515496, 1384137),
               ("25-to-64", 6065987, 5792462), ("65-to-74", 1228376, 1030175),
               ("75-to-84", 586511, 641020), ("85+", 232833, 288627))
AGE_SLICES = ((0, 3), (3, 5), (5, 13), (13, 15), (15, 17), (17, 18))


def _geographies() -> dict[str, str]:
    geos = {"Ohio": "39"}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        fields = line.split("|")
        if fields[0] == "OH":
            geos[fields[4].removesuffix(" County")] = fields[1] + fields[2]
    if len(geos) != 89:
        raise ValueError("OH: expected state and all 88 counties")
    return geos


def _integer(value, label: str) -> int:
    if (not isinstance(value, (int, float)) or isinstance(value, bool)
            or not math.isfinite(value) or value < 0 or int(value) != value):
        raise ValueError(f"OH: invalid integer {label} {value!r}")
    return int(value)


def _parse_grid(grid: list[tuple]) -> list[dict]:
    if (len(grid) != 624 or any(len(row) != 60 for row in grid)
            or tuple(grid[0]) != HEADERS):
        raise ValueError("OH: unexpected workbook header or shape")
    geos = _geographies()
    totals = {}
    rows = []
    for source in grid[1:]:
        name, year = source[:2]
        if not isinstance(name, str) or name not in geos:
            raise ValueError(f"OH: unknown geography {name!r}")
        year = _integer(year, "year")
        if year not in YEARS or source[59] is not None:
            raise ValueError("OH: unexpected year or extra data")
        gid = geos[name]
        key = gid, year
        if key in totals:
            raise ValueError(f"OH: duplicate geography/year {name} {year}")
        cells = [_integer(value, "population") for value in source[2:59]]
        male, female, both = cells[:19], cells[19:38], cells[38:57]
        if (sum(male[:-1]) != male[-1] or sum(female[:-1]) != female[-1]
                or sum(both[:-1]) != both[-1] or male[-1] + female[-1] != both[-1]):
            raise ValueError(f"OH: age/sex totals do not reconcile for {name} {year}")
        differences = tuple(m + f - t for m, f, t in zip(male[:-1], female[:-1], both[:-1], strict=True))
        expected = STATE_2020_AGE_DIFFERENCES if key == ("39", 2020) else (0,) * 18
        if differences != expected:
            raise ValueError(f"OH: unexpected age/sex cell difference for {name} {year}")
        if gid == "39" and year in (2020, 2050):
            anchor_values = [anchor[1 if year == 2020 else 2] for anchor in AGE_ANCHORS]
            if [sum(both[start:end]) for start, end in AGE_SLICES] != anchor_values:
                raise ValueError("OH: state ages disagree with independent report anchors")
        totals[key] = both[-1]
        rows.append(dict(state="OH", geo_level="state" if gid == "39" else "county",
            geo_id=gid, year=year, variable="population", value=both[-1],
            datatype="estimate" if year == 2020 else "projection", source_id=SOURCE_ID))
    expected = {(gid, year) for gid in geos.values() for year in YEARS}
    if set(totals) != expected:
        raise ValueError("OH: incomplete geography/year grid")
    for year in YEARS:
        if sum(value for (gid, y), value in totals.items() if gid != "39" and y == year) != totals["39", year]:
            raise ValueError(f"OH: county/state totals do not reconcile in {year}")
    return rows


def parse(body: bytes) -> list[dict]:
    import openpyxl
    try:
        book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("OH: response is not a readable XLSX workbook") from exc
    try:
        if (book.sheetnames != ["data"] or book.properties.creator != "Howison, Jeffrey"
                or book.properties.modified != datetime(2023, 1, 6, 19, 7, 39)):
            raise ValueError("OH: unexpected workbook sheets or edition")
        return _parse_grid(list(book["data"].iter_rows(values_only=True)))
    finally:
        book.close()


def _pdf_pages(body: bytes, count: int, label: str) -> list[str]:
    from pypdf import PdfReader
    fetch.validate_pdf(body, label)
    reader = PdfReader(io.BytesIO(body))
    if len(reader.pages) != count:
        raise ValueError(f"OH: unexpected {label} page count")
    return [" ".join(page.extract_text().split()) for page in reader.pages]


def _validate_references(report: bytes, overview: bytes) -> None:
    pages = _pdf_pages(report, 16, "report")
    summary = _pdf_pages(overview, 2, "overview")
    if ("Population Projections Report 2023" not in pages[0]
            or "for Ohio and its 88 counties" not in pages[1]
            or "Ohio Department of Development Research Office" not in pages[1]
            or "2016 to 2020" not in pages[2] or "cohort component method" not in pages[2]
            or "county total populations are controlled" not in pages[3]
            or "population estimates for 2020 were used as the base population instead of the census count" not in pages[14]
            or "State of Ohio Population Projections Overview: 2020 to 2050" not in summary[0]
            or "Ohio Department of Development Research Office (2022)" not in summary[0]):
        raise ValueError("OH: unexpected reference edition, producer or base year")
    for text in (pages[1], summary[1]):
        if "Total Population 11,799,448 11,123,896" not in text:
            raise ValueError("OH: reference state totals disagree")
        for label, base, final in AGE_ANCHORS:
            if f"{label} {base:,} {final:,}" not in text:
                raise ValueError("OH: reference age anchors disagree")
    if "Ohio 11,799,448 11,123,896" not in summary[0]:
        raise ValueError("OH: overview state totals disagree")


def _validate_catalog(body: bytes) -> dict:
    try:
        catalog = json.loads(body)
        retrieved = date.fromisoformat(catalog["retrieved_at"])
        links = {item["href"]: " ".join(item["text"].split()) for item in catalog["links"]}
        if (catalog["format"] != "browser-rendered-catalog-excerpt"
                or catalog["source_url"] != CATALOG_URL or catalog["page_title"] != "Population | Development"
                or "not the original HTTP response HTML" not in catalog["extraction"]
                or "Office of Research" not in catalog["overview"]
                or "fertility, mortality, and migration" not in catalog["report_note"]
                or "Projections by County, 2020 to 2050" not in links[URL]
                or "Population Projections Report 2023" not in links[REPORT_LINK]
                or retrieved > date.today()):
            raise ValueError("unexpected catalog content")
    except (ValueError, TypeError, KeyError) as exc:
        raise ValueError("OH: invalid browser catalog excerpt or selected product links") from exc
    return catalog


def _validate_all(bodies: dict[str, bytes]) -> list[dict]:
    _validate_catalog(bodies[CATALOG_FILE])
    _validate_references(bodies[REPORT_FILE], bodies[OVERVIEW_FILE])
    return parse(bodies[FILE])


def acquire(*, catalog_excerpt: bytes | None = None) -> None:
    if (RAW_DIR / "manifest.json").exists():
        entries = fetch.read_manifest(RAW_DIR)
        if {entry["file"] for entry in entries} != set(FILES):
            raise ValueError("OH: incomplete saved manifest")
        _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
        print("OH acquire: reused validated 2023-edition snapshots")
        return
    if catalog_excerpt is None:
        raise ValueError("OH: new acquisition requires a read-only browser catalog excerpt; pass --catalog-excerpt PATH")
    catalog = _validate_catalog(catalog_excerpt)
    bodies = {name: fetch.get_bytes(url) for name, url in FILES.items() if name != CATALOG_FILE}
    bodies[CATALOG_FILE] = catalog_excerpt
    rows = _validate_all(bodies)
    entries = []
    for name, url in FILES.items():
        fetch.write_snapshot(RAW_DIR / name, bodies[name])
        entries.append(dict(file=name, source=url,
            retrieved_at=catalog["retrieved_at"] if name == CATALOG_FILE else date.today().isoformat(),
            sha256=fetch.sha256(bodies[name]), records=len(rows) if name == FILE else None,
            note=("Read-only browser-rendered catalog EXCERPT, not original HTTP HTML. Records observed workbook and Report 2023 links; catalog HTTP requests returned 404."
                  if name == CATALOG_FILE else "Unmodified publisher original. Workbook fetched from publisher DAM raw asset corresponding to the catalog's static link; report and overview independently corroborate state endpoints and ages. 2020 estimates base; exact five-year totals. Overview cites 2022 while current report/catalog label 2023. Overview map Ottawa 2050 differs by three from the selected workbook.")))
    fetch.write_manifest(RAW_DIR, entries)


def load(conn) -> None:
    entries = {entry["file"]: entry for entry in fetch.read_manifest(RAW_DIR)}
    if set(entries) != set(FILES):
        raise ValueError("OH: incomplete saved manifest")
    rows = _validate_all({name: (RAW_DIR / name).read_bytes() for name in FILES})
    geos = {(row["geo_level"], row["geo_id"]) for row in rows}
    seeded = {(row["geo_level"], row["geo_id"]) for row in conn.execute("SELECT * FROM geographies WHERE state='OH'")}
    if seeded != geos:
        raise ValueError("OH: seeded geographies differ from publication")
    upsert_source(conn, source_id=SOURCE_ID, state="OH", agency=AGENCY,
        product_name="Projections by County, 2020 to 2050; state and county Total_TOT",
        vintage=VINTAGE, base_year=2020, horizon_year=2050, source_url=URL,
        retrieved_at=entries[FILE]["retrieved_at"],
        license="No explicit data license identified; cite the Ohio Department of Development Research Office.",
        citation="Ohio Department of Development Research Office. Population Projections Report 2023 and accompanying Projections by County, 2020 to 2050 workbook; data sheet, Total_TOT column.",
        method_notes=("The original publisher workbook independently publishes Ohio and all 88 county total populations at seven five-year points. "
            "Total_TOT cells are served unchanged, with no interpolation. Report endnote 2 says the 2020 estimates are the base; "
            "county totals match Census 2020 while age/sex cohorts may differ. 2020 is estimate; 2025-2050 remain projection. "
            "The cohort-component model uses 2016-2020 demographic rates, then controls county projections to the separately projected state. "
            "All county sums equal the published state totals and every row's age and sex totals reconcile. The state 2020 sex-age sums differ "
            "from its separate both-sex age cells by published offsets summing zero; those both-sex ages match report Table 1. "
            "Report Table 1 and the overview independently corroborate state 2020/2050 totals and broad age groups. "
            "The overview map prints Ottawa 2050 as 31,374; the workbook's 31,371 is selected consistently. "
            "Vintage uses the current catalog's Report 2023 label and workbook modified January 6, 2023; the overview cites a 2022 series. "
            "An exact release date or later rebasing was not established. The report discusses conditional revisions after 2025 estimates, "
            "but the browser-rendered catalog still links this workbook and Report 2023 on September 30, 2026. "
            "The workbook is fetched from the publisher's DAM raw asset; an archived browser catalog excerpt records the canonical static link. "
            "Statewide projected deaths, ten-year net migration and demographic detail are archived but not served; complete components coverage remains unverified. "
            "Household population counts persons, not household units. Household-count, subcounty and eligible Research Office jobs forecast coverage remains unverified."))
    insert_observations(conn, rows)
    conn.commit()
    print(f"OH load: {len(rows)} observations")


# State endpoints independently read from report/overview; county 2050 values
# independently read from the overview map (Ottawa excluded because it differs).
FIXTURES = (("39", 2020, 11799448), ("39", 2050, 11123896),
            ("39035", 2050, 1030507), ("39049", 2050, 1487885),
            ("39051", 2050, 36801), ("39107", 2050, 41766),
            ("39163", 2050, 9117), ("39175", 2050, 16233))


def verify(conn) -> None:
    rows = conn.execute("SELECT * FROM observations WHERE state='OH'").fetchall()
    expected = {(gid, year) for gid in _geographies().values() for year in YEARS}
    assert len(rows) == 623 and {(row["geo_id"], row["year"]) for row in rows} == expected
    assert all(row["source_id"] == SOURCE_ID and row["variable"] == "population" and row["scenario"] == ""
        and row["age_group"] == row["sex"] == row["race"] == "all"
        and row["datatype"] == ("estimate" if row["year"] == 2020 else "projection")
        and row["geo_level"] == ("state" if row["geo_id"] == "39" else "county")
        and math.isfinite(row["value"]) and row["value"] >= 0 and float(row["value"]).is_integer() for row in rows)
    totals = {(row["geo_id"], row["year"]): row["value"] for row in rows}
    for gid, year, value in FIXTURES:
        assert totals[gid, year] == value
    for year in YEARS:
        assert sum(row["value"] for row in rows if row["geo_level"] == "county" and row["year"] == year) == totals["39", year]
    print("OH verify: 623 complete totals and independent publisher anchors ok")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--acquire", action="store_true")
    parser.add_argument("--catalog-excerpt", type=Path, help="Read-only browser catalog capture for first acquisition")
    args = parser.parse_args()
    if args.acquire:
        acquire(catalog_excerpt=args.catalog_excerpt.read_bytes() if args.catalog_excerpt else None)
    else:
        parser.print_help()
