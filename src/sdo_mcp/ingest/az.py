"""Arizona adapter.

Source: Arizona Office of Economic Opportunity (OEO), 2025 State and County
Population Projections. The archive contains Low/Medium/High workbooks for
Arizona and all 15 counties, annual from July 1, 2025 through 2060. Verified
2026-08-10 (plan/EXPANSION.md).

The 2025 population in every scenario exactly matches OEO's official 2025
estimate. It is labeled estimate; 2026 onward is projection. The workbooks
also publish births, deaths, and total net migration for every scenario.
"""

from __future__ import annotations

import io
import sys
from datetime import date
from zipfile import ZipFile

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source

RAW_DIR = REPO_ROOT / "data" / "raw" / "az"
FILE = "All_Series_2025-2060.zip"
URL = (
    "https://oeo.az.gov/sites/default/files/data/pop/Projections_2025_SC/"
    "All_Series_2025-2060.zip"
)

STATE_FIPS = "04"
SOURCE_ID = "az-pop-2025"
BASE_YEAR = 2025
HORIZON_YEAR = 2060
SCENARIOS = ("low", "medium", "high")

AGENCY = "Arizona Office of Economic Opportunity (OEO), State Demographer"
CITATION = (
    "Arizona Office of Economic Opportunity. Arizona State and County "
    "Population Projections, 2025-2060, 2025 edition. Retrieved from oeo.az.gov."
)
LICENSE = "No separate license statement; official public statistics under az.gov terms."


def _members(archive: ZipFile) -> list[str]:
    return [
        name for name in archive.namelist()
        if name.endswith(".xlsx") and not name.startswith("__MACOSX/")
    ]


def acquire() -> None:
    """Download the Cloudflare-guarded archive and validate its contents."""
    body = fetch.get_with_js(URL)
    try:
        with ZipFile(io.BytesIO(body)) as archive:
            members = _members(archive)
            if len(members) != 48:
                raise RuntimeError(f"AZ acquire: archive has {len(members)} workbooks, expected 48")
            for name in members:
                with archive.open(name) as member:
                    if not member.read(4).startswith(b"PK"):
                        raise RuntimeError(
                            f"AZ acquire: {name!r} is not an xlsx workbook"
                        )
    except Exception as exc:
        if isinstance(exc, RuntimeError):
            raise
        raise RuntimeError(f"{URL}: body is not the expected zip archive: {exc}") from exc

    RAW_DIR.mkdir(parents=True, exist_ok=True)
    (RAW_DIR / FILE).write_bytes(body)
    fetch.write_manifest(RAW_DIR, [{
        "file": FILE,
        "source": URL,
        "retrieved_at": date.today().isoformat(),
        "sha256": fetch.sha256(body),
        "records": 48,
        "note": (
            "48 workbooks: Arizona + 15 counties x Low/Medium/High. "
            "Downloaded through a challenge-capable client and validated as "
            "xlsx files before snapshotting."
        ),
    }])
    print(f"AZ acquire: {FILE} ({len(body):,} bytes, 48 workbooks)")


def _retrieved() -> str:
    return next(
        e["retrieved_at"] for e in fetch.read_manifest(RAW_DIR) if e["file"] == FILE
    )


def load(conn) -> None:
    """Archive snapshot -> population and components by scenario."""
    import openpyxl

    fips_by_name = {
        row["name"].removesuffix(" County"): row["geo_id"]
        for row in conn.execute(
            "SELECT geo_id, name FROM geographies WHERE state='AZ' AND geo_level='county'"
        )
    }

    rows: list[dict] = []
    seen: dict[str, set[str]] = {s: set() for s in SCENARIOS}
    with ZipFile(RAW_DIR / FILE) as archive:
        members = _members(archive)
        assert len(members) == 48, f"AZ load: {len(members)} workbooks, expected 48"
        for member in members:
            scenario = member.split("/")[-2].lower()
            assert scenario in SCENARIOS, f"AZ load: unknown scenario path {member!r}"
            geography = member.rsplit("/", 1)[-1].split("_", 1)[0]
            seen[scenario].add(geography)
            if geography == "Arizona":
                level, gid = "state", STATE_FIPS
            else:
                level, gid = "county", fips_by_name.get(geography)
                if gid is None:
                    raise RuntimeError(f"AZ load: county {geography!r} has no Census FIPS match")

            wb = openpyxl.load_workbook(
                io.BytesIO(archive.read(member)), read_only=True, data_only=True
            )
            try:
                ws = wb["1. Total Pop & Components"]
                header = [str(v or "").strip() for v in next(
                    ws.iter_rows(min_row=4, max_row=4, values_only=True)
                )]
                expected = {
                    "Year": 0,
                    "Population": 1,
                    "Births": 4,
                    "Deaths": 5,
                    "Total Net Migration2": 9,
                }
                for label, index in expected.items():
                    assert header[index] == label, (
                        f"AZ load: {member} column {index + 1} is {header[index]!r}, "
                        f"expected {label!r}"
                    )
                points = 0
                for rec in ws.iter_rows(min_row=5, values_only=True):
                    if not isinstance(rec[0], int) or not BASE_YEAR <= rec[0] <= HORIZON_YEAR:
                        continue
                    year = int(rec[0])
                    datatype = "estimate" if year == BASE_YEAR else "projection"
                    for variable, index in (
                        ("population", 1),
                        ("births", 4),
                        ("deaths", 5),
                        ("net_migration", 9),
                    ):
                        rows.append(dict(
                            state="AZ", geo_level=level, geo_id=gid, year=year,
                            variable=variable, value=int(round(float(rec[index]))),
                            datatype=datatype, scenario=scenario, source_id=SOURCE_ID,
                        ))
                    points += 1
                assert points == 36, f"AZ load: {member} has {points} years, expected 36"
            finally:
                wb.close()

    expected_geographies = set(fips_by_name) | {"Arizona"}
    for scenario, names in seen.items():
        if names != expected_geographies:
            raise RuntimeError(
                f"AZ load: {scenario} geography mismatch; missing "
                f"{sorted(expected_geographies - names)}, extra {sorted(names - expected_geographies)}"
            )

    upsert_source(
        conn,
        source_id=SOURCE_ID,
        state="AZ",
        agency=AGENCY,
        product_name="Arizona State and County Population Projections, 2025-2060",
        vintage="2025 edition (files restamped April 2026)",
        base_year=BASE_YEAR,
        horizon_year=HORIZON_YEAR,
        source_url=URL,
        retrieved_at=_retrieved(),
        license=LICENSE,
        citation=CITATION,
        method_notes=(
            "Annual July 1 Low/Medium/High scenarios. The 2025 launch population "
            "equals OEO's official July 1, 2025 estimate in all three scenarios; "
            "it is labeled estimate and 2026 onward is labeled projection. "
            "Population and component values are rounded to whole persons from "
            "the detailed workbook cells; the separate SummaryTable is rounded "
            "to the nearest 100 and is not used. Place projections remain on an "
            "older county vintage and are deliberately not mixed into this product. "
            "Published scenario ordering has one small crossover: Yuma County's "
            "2026 Low value is 61 people above Medium; the series order normally "
            "holds from 2027 onward."
        ),
    )
    n = insert_observations(conn, rows)
    conn.commit()
    print(f"AZ load: {n} observations")


FIXTURES = [
    ("state", "04", "medium", 2025, 7718747, "estimate"),
    ("state", "04", "medium", 2060, 9750566, "projection"),
    ("state", "04", "low", 2060, 8854828, "projection"),
    ("state", "04", "high", 2060, 12294500, "projection"),
    ("county", "04019", "medium", 2025, 1093761, "estimate"),  # Pima
    ("county", "04019", "medium", 2060, 1143575, "projection"),
]


def verify(conn) -> None:
    for level, gid, scenario, year, value, datatype in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype FROM observations WHERE state='AZ' AND "
            "geo_level=? AND geo_id=? AND variable='population' AND scenario=? AND year=?",
            (level, gid, scenario, year),
        ).fetchone()
        assert row is not None, f"AZ verify: missing {level} {gid} {scenario} {year}"
        assert int(row["value"]) == value and row["datatype"] == datatype, (
            f"AZ verify: {level} {gid} {scenario} {year} -> "
            f"{row['value']}/{row['datatype']}, expected {value}/{datatype}"
        )

    shape = conn.execute(
        "SELECT COUNT(*) n, COUNT(DISTINCT geo_id) ng, COUNT(DISTINCT year) ny "
        "FROM observations WHERE state='AZ' AND variable='population'"
    ).fetchone()
    assert (shape["n"], shape["ng"], shape["ny"]) == (1728, 16, 36)

    shared_base = conn.execute(
        "SELECT geo_id, COUNT(DISTINCT value) n FROM observations WHERE state='AZ' "
        "AND variable='population' AND year=2025 GROUP BY geo_id HAVING n!=1"
    ).fetchall()
    assert not shared_base, f"AZ verify: scenario launch values differ: {[tuple(r) for r in shared_base]}"

    bad = conn.execute(
        "SELECT lo.geo_id, lo.year FROM observations lo "
        "JOIN observations md ON md.state=lo.state AND md.geo_level=lo.geo_level "
        "AND md.geo_id=lo.geo_id AND md.year=lo.year AND md.variable=lo.variable "
        "AND md.scenario='medium' JOIN observations hi ON hi.state=lo.state "
        "AND hi.geo_level=lo.geo_level AND hi.geo_id=lo.geo_id AND hi.year=lo.year "
        "AND hi.variable=lo.variable AND hi.scenario='high' WHERE lo.state='AZ' "
        "AND lo.variable='population' AND lo.scenario='low' "
        "AND NOT (lo.value<=md.value AND md.value<=hi.value)"
    ).fetchall()
    # The source itself has one small crossover: Yuma 2026 Low rounds to
    # 225,869 and Medium to 225,808. Preserve it and assert it is the only
    # exception, so neither a source quirk nor a future parsing error is hidden.
    anomalies = [(r["geo_id"], r["year"]) for r in bad]
    assert anomalies == [("04027", 2026)], (
        f"AZ verify: unexpected scenario-order exceptions: {anomalies[:10]}"
    )

    components = conn.execute(
        "SELECT variable, COUNT(*) n FROM observations WHERE state='AZ' AND "
        "variable IN ('births','deaths','net_migration') GROUP BY variable"
    ).fetchall()
    assert {r["variable"]: r["n"] for r in components} == {
        "births": 1728,
        "deaths": 1728,
        "net_migration": 1728,
    }
    print(
        f"AZ verify: {len(FIXTURES)} fixtures + shape + shared base + scenario "
        "ordering (one published Yuma-2026 crossover) ok"
    )


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.az --acquire")
