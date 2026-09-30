"""Kentucky State Data Center Vintage 2022 population and household projections.

The official workbooks contain statewide, county, and Area Development
District data. The current server geography contract supports state and county
queries, so this adapter serves only the state and 120 county totals.
"""

from __future__ import annotations

import io
import math
import sys
from datetime import date

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "ky"
POPULATION_FILE = "population_projections_2022.xlsx"
HOUSEHOLDS_FILE = "household_projections_2022.xlsx"
CATALOG_FILE = "data_catalog_2026-09-29.html"
CODEBOOK_FILE = "population_codebook_2022.pdf"
POPULATION_URL = "https://louisville.widen.net/s/rzn2jnhjph/population-projections-by-sex-and-five-year-age-group"
HOUSEHOLDS_URL = "https://louisville.widen.net/s/hbwbqlghqh/household-projections"
CATALOG_URL = "https://centers.louisville.edu/kentucky-state-data-center/data-downloads"
CODEBOOK_URL = "https://louisville.widen.net/content/mwcdfskvwf/original/Population-Projections-by-Sex-and-Five-Year-Age-Group.pdf?download=true&u=kwgjzs"
STATE_FIPS = "21"
POPULATION_SOURCE_ID = "ky-pop-2022"
HOUSEHOLDS_SOURCE_ID = "ky-households-2022"
AGENCY = "Kentucky State Data Center, University of Louisville"
LICENSE = "No explicit data license identified; cite the Kentucky State Data Center."
VINTAGE = "2022 (released August 2022)"
BASE_YEAR = 2021  # Last estimate input named by the catalog; Census anchor is 2020.
HORIZON_YEAR = 2050
YEARS = tuple(range(2020, 2051, 5))
HOUSEHOLD_YEARS = (2010, *YEARS)
SEXES = (0, 1)
AGE_GROUPS = tuple(range(1, 19))
POPULATION_HEADER = (
    "GEO", "FIPS", "NAME", "MALE", "AGEGRP", "N2010", "HHPOP2010", "GQ2010",
    "N2020", "HHPOP2020", "GQ2020", "N2025", "HHPOP2025", "GQ2025",
    "N2030", "HHPOP2030", "GQ2030", "N2035", "HHPOP2035", "GQ2035",
    "N2040", "HHPOP2040", "GQ2040", "N2045", "HHPOP2045", "GQ2045",
    "N2050", "HHPOP2050", "GQ2050",
)
HOUSEHOLDS_HEADER = (
    "GEO", "FIPS", "NAME", "HH2010", "HH2020", "HH2025", "HH2030",
    "HH2035", "HH2040", "HH2045", "HH2050",
)


def _county_geographies() -> dict[str, str]:
    """Return Census county FIPS -> canonical county name for Kentucky."""
    geographies = {}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        parts = line.split("|")
        if len(parts) >= 5 and parts[0] == "KY":
            geographies[parts[1] + parts[2]] = parts[4]
    if len(geographies) != 120:
        raise ValueError(f"KY: expected 120 Census counties, found {len(geographies)}")
    return geographies


def _integer(value, *, label: str) -> int:
    if (not isinstance(value, (int, float)) or isinstance(value, bool)
            or not math.isfinite(float(value)) or int(value) != value
            or value < 0):
        raise ValueError(f"KY: invalid {label}: {value!r}")
    return int(value)


def parse_population(body: bytes) -> list[dict]:
    """Aggregate the published age-by-sex rows into total population."""
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("KY: response is not a readable population workbook") from exc
    try:
        if workbook.sheetnames != ["Population-Projections-by-Sex-a"]:
            raise ValueError(f"KY: unexpected population sheets: {workbook.sheetnames!r}")
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        if tuple(next(rows, ())) != POPULATION_HEADER:
            raise ValueError("KY: unexpected population workbook header")

        counties = _county_geographies()
        seen = set()
        totals: dict[tuple[str, str, int], int] = {}
        geography_counts = {"State": 0, "County": 0, "ADD": 0}
        add_names = set()
        for row in rows:
            if not row or all(value is None for value in row):
                continue
            if len(row) != len(POPULATION_HEADER):
                raise ValueError("KY: malformed population row")
            geography, fips, name, male, age_group, *values = row
            if geography not in geography_counts:
                raise ValueError(f"KY: unexpected geography type {geography!r}")
            if (type(male) is not int or type(age_group) is not int
                    or male not in SEXES or age_group not in AGE_GROUPS):
                raise ValueError(f"KY: unexpected sex/age group {male!r}/{age_group!r}")
            geography_counts[geography] += 1

            if geography == "State":
                if fips is not None or name != "Kentucky":
                    raise ValueError(f"KY: invalid state geography {fips!r}/{name!r}")
                level, geo_id = "state", STATE_FIPS
            elif geography == "County":
                if type(fips) is not int:
                    raise ValueError(f"KY: invalid county FIPS {fips!r}")
                geo_id = str(fips).zfill(5)
                if counties.get(geo_id, "").removesuffix(" County") != name:
                    raise ValueError(f"KY: county FIPS/name mismatch {geo_id}/{name!r}")
                level = "county"
            else:
                if fips is not None or not isinstance(name, str) or not name:
                    raise ValueError(f"KY: invalid Area Development District {fips!r}/{name!r}")
                add_names.add(name)
                level, geo_id = "add", name

            key = (geography, str(geo_id), male, age_group)
            if key in seen:
                raise ValueError(f"KY: duplicate population row {key}")
            seen.add(key)

            # Validate every published total column, including the historical
            # columns not loaded into the MCP.
            population = [
                _integer(value, label=f"population at {name}/{year}")
                for year, value in zip(HOUSEHOLD_YEARS, values[::3])
            ]
            if geography == "ADD":
                continue
            for year, value in zip(YEARS, population[1:]):
                total_key = (level, str(geo_id), year)
                totals[total_key] = totals.get(total_key, 0) + value

        if geography_counts != {"State": 36, "County": 120 * 36, "ADD": 15 * 36}:
            raise ValueError(f"KY: unexpected population geography row counts {geography_counts}")
        if len(add_names) != 15:
            raise ValueError(f"KY: expected 15 distinct Area Development Districts, found {len(add_names)}")
        if len(seen) != 136 * 36:
            raise ValueError(f"KY: incomplete population grid; found {len(seen)} rows")

        expected_geos = {("state", STATE_FIPS)} | {
            ("county", fips) for fips in counties
        }
        if {(level, gid) for level, gid, _year in totals} != expected_geos:
            raise ValueError("KY: population totals do not match state and county geography coverage")
        observations = [
            {
                "state": "KY", "geo_level": level, "geo_id": geo_id,
                "year": year, "variable": "population", "value": value,
                "datatype": "estimate" if year == 2020 else "projection",
                "source_id": POPULATION_SOURCE_ID,
            }
            for (level, geo_id, year), value in sorted(totals.items())
        ]
        return observations
    finally:
        workbook.close()


def parse_households(body: bytes) -> list[dict]:
    """Read published state and county household totals."""
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("KY: response is not a readable household workbook") from exc
    try:
        if workbook.sheetnames != ["Household-Projections"]:
            raise ValueError(f"KY: unexpected household sheets: {workbook.sheetnames!r}")
        rows = workbook.active.iter_rows(values_only=True)
        if tuple(next(rows, ())) != HOUSEHOLDS_HEADER:
            raise ValueError("KY: unexpected household workbook header")

        counties = _county_geographies()
        seen = set()
        observations = []
        add_names = set()
        for row in rows:
            if not row or all(value is None for value in row):
                continue
            if len(row) != len(HOUSEHOLDS_HEADER):
                raise ValueError("KY: malformed household row")
            geography, fips, name, *values = row
            if geography == "State":
                if fips is not None or name != "Kentucky":
                    raise ValueError(f"KY: invalid household state geography {fips!r}/{name!r}")
                level, geo_id = "state", STATE_FIPS
            elif geography == "County":
                if type(fips) is not int:
                    raise ValueError(f"KY: invalid household county code {fips!r}")
                geo_id = f"21{fips:03d}"
                if counties.get(geo_id, "").removesuffix(" County") != name:
                    raise ValueError(f"KY: household county code/name mismatch {geo_id}/{name!r}")
                level = "county"
            elif geography == "ADD":
                if fips is not None or not isinstance(name, str) or not name:
                    raise ValueError(f"KY: invalid household Area Development District {fips!r}/{name!r}")
                add_names.add(name)
                level, geo_id = "add", name
            else:
                raise ValueError(f"KY: unexpected household geography type {geography!r}")
            key = (level, geo_id)
            if key in seen:
                raise ValueError(f"KY: duplicate household geography {key}")
            seen.add(key)

            if len(values) != len(HOUSEHOLD_YEARS):
                raise ValueError(f"KY: incomplete household years for {name}")
            validated = [
                _integer(value, label=f"households at {name}/{year}")
                for value, year in zip(values, HOUSEHOLD_YEARS)
            ]
            if level == "add":
                continue
            for year, value in zip(HOUSEHOLD_YEARS[1:], validated[1:]):
                observations.append({
                    "state": "KY", "geo_level": level, "geo_id": geo_id,
                    "year": year, "variable": "households", "value": value,
                    "datatype": "estimate" if year == 2020 else "projection",
                    "source_id": HOUSEHOLDS_SOURCE_ID,
                })

        expected = {("state", STATE_FIPS)} | {("county", fips) for fips in counties}
        if not expected <= seen or len(seen) != 136 or len(add_names) != 15:
            raise ValueError(
                f"KY: household geographies differ; state/counties={len(expected & seen)}, "
                f"all={len(seen)}, ADDs={len(add_names)}"
            )
        if len(observations) != 121 * len(YEARS):
            raise ValueError(f"KY: incomplete household totals, found {len(observations)}")
        return observations
    finally:
        workbook.close()


def _validate_catalog(body: bytes) -> None:
    text = body.decode("utf-8", "replace")
    normalized = " ".join(text.split()).lower()
    for marker in ("<html", "kentucky state data center", "vintage 2022",
                   "august 2022", "area development districts", "2020 decennial census",
                   "2021 population estimates", POPULATION_URL, HOUSEHOLDS_URL):
        if marker not in normalized:
            raise ValueError(f"KY: unexpected data catalog page; missing {marker!r}")


def _validate_codebook(body: bytes) -> None:
    from pypdf import PdfReader

    fetch.validate_pdf(body, "KY codebook")
    text = " ".join(" ".join(page.extract_text() for page in
                            PdfReader(io.BytesIO(body)).pages).split())
    for marker in ("0 = Female", "1 = Male", "18 = 85+", "Total population in [YEAR]"):
        if marker not in text:
            raise ValueError(f"KY: unexpected population codebook; missing {marker!r}")


def acquire() -> None:
    """Download and validate immutable source snapshots."""
    required = {POPULATION_FILE, HOUSEHOLDS_FILE, CATALOG_FILE, CODEBOOK_FILE}
    if (RAW_DIR / "manifest.json").exists():
        entries = fetch.read_manifest(RAW_DIR)
        if not required <= {entry["file"] for entry in entries}:
            raise ValueError("KY: incomplete saved manifest")
        parse_population((RAW_DIR / POPULATION_FILE).read_bytes())
        parse_households((RAW_DIR / HOUSEHOLDS_FILE).read_bytes())
        _validate_catalog((RAW_DIR / CATALOG_FILE).read_bytes())
        _validate_codebook((RAW_DIR / CODEBOOK_FILE).read_bytes())
        print("KY acquire: reused validated Vintage 2022 snapshots")
        return

    population = fetch.get_bytes(POPULATION_URL)
    population_rows = parse_population(population)
    households = fetch.get_bytes(HOUSEHOLDS_URL)
    household_rows = parse_households(households)
    catalog = fetch.get_bytes(CATALOG_URL)
    _validate_catalog(catalog)
    codebook = fetch.get_bytes(CODEBOOK_URL)
    _validate_codebook(codebook)

    today = date.today().isoformat()
    snapshots = [
        (POPULATION_FILE, POPULATION_URL, population, len(population_rows),
         "Original Kentucky State Data Center Vintage 2022 population workbook. State and county totals are aggregated from its published five-year age-by-sex rows; Area Development District rows remain archived but are not served."),
        (HOUSEHOLDS_FILE, HOUSEHOLDS_URL, households, len(household_rows),
         "Original Kentucky State Data Center Vintage 2022 household workbook. State and county totals are served; Area Development District rows remain archived but are not served."),
        (CATALOG_FILE, CATALOG_URL, catalog, None,
         "Kentucky State Data Center data catalog saved as source-page evidence for the Vintage 2022 files."),
        (CODEBOOK_FILE, CODEBOOK_URL, codebook, None,
         "Original KSDC population codebook defining geography, sex, age-group and total-population columns."),
    ]
    entries = []
    for filename, url, body, count, note in snapshots:
        fetch.write_snapshot(RAW_DIR / filename, body)
        entries.append({
            "file": filename, "source": url, "retrieved_at": today,
            "sha256": fetch.sha256(body), "records": count, "note": note,
        })
    fetch.write_manifest(RAW_DIR, entries)
    print(f"KY acquire: {len(population_rows)} population and {len(household_rows)} household totals")


def load(conn) -> None:
    entries = fetch.read_manifest(RAW_DIR)
    _validate_catalog((RAW_DIR / CATALOG_FILE).read_bytes())
    _validate_codebook((RAW_DIR / CODEBOOK_FILE).read_bytes())
    population_retrieved = next(
        entry["retrieved_at"] for entry in entries if entry["file"] == POPULATION_FILE
    )
    household_retrieved = next(
        entry["retrieved_at"] for entry in entries if entry["file"] == HOUSEHOLDS_FILE
    )
    population = parse_population((RAW_DIR / POPULATION_FILE).read_bytes())
    households = parse_households((RAW_DIR / HOUSEHOLDS_FILE).read_bytes())

    actual_geographies = {
        (row["geo_level"], row["geo_id"])
        for row in conn.execute("SELECT geo_level, geo_id FROM geographies WHERE state='KY'")
    }
    expected_geographies = {("state", STATE_FIPS)} | {
        ("county", fips) for fips in _county_geographies()
    }
    if actual_geographies != expected_geographies:
        raise ValueError("KY: seeded state and county geographies do not match the source")

    upsert_source(
        conn, source_id=POPULATION_SOURCE_ID, state="KY", agency=AGENCY,
        product_name="Kentucky State, County, and Area Development District Population Projections",
        vintage=VINTAGE, base_year=BASE_YEAR, horizon_year=HORIZON_YEAR,
        source_url=POPULATION_URL, retrieved_at=population_retrieved, license=LICENSE,
        citation=("Kentucky State Data Center, University of Louisville. Population and Household "
                  "Projections: Kentucky, Kentucky Counties, and Area Development Districts, 2020-2050. August 2022."),
        method_notes=(
            "Vintage 2022 uses 2020 Decennial Census and 2021 Population Estimates as its population base. "
            "The report's model anchor is the 2020 Census; base_year=2021 records the last underlying "
            "estimate year identified by the catalog, not the first table year. No 2021 row is published. "
            "The workbook publishes 2020 historical totals and five-year projections from 2025 through 2050. "
            "Historical Census-based totals use the MCP estimate label; all 2025-2050 values are projections. "
            "The source supplies population by sex and five-year age group; totals are exact sums of those "
            "published categories. State totals equal the sum of the published county totals. No annual "
            "values or Area Development District totals are generated."
        ),
    )
    upsert_source(
        conn, source_id=HOUSEHOLDS_SOURCE_ID, state="KY", agency=AGENCY,
        product_name="Kentucky State, County, and Area Development District Household Projections",
        vintage=VINTAGE, base_year=BASE_YEAR, horizon_year=HORIZON_YEAR,
        source_url=HOUSEHOLDS_URL, retrieved_at=household_retrieved, license=LICENSE,
        citation=("Kentucky State Data Center, University of Louisville. Population and Household "
                  "Projections: Kentucky, Kentucky Counties, and Area Development Districts, 2020-2050. August 2022."),
        method_notes=(
            "The Vintage 2022 workbook publishes 2020 historical household totals and five-year household "
            "projections from 2025 through 2050 for the state, all 120 counties, and 15 Area Development "
            "Districts. Headship rates from the 2020 ACS are calibrated to 2020 Census household counts, "
            "then applied to the projected household population. base_year=2021 records the last underlying "
            "population estimate input identified by the catalog; the household-count anchor is 2020. "
            "No 2021 household value is published or generated. Historical Census-based totals use the MCP "
            "estimate label. This server loads only state and county totals and preserves publisher values. "
            "No annual values or Area Development District totals are generated."
        ),
    )
    count = insert_observations(conn, population + households)
    conn.commit()
    print(f"KY load: {count} observations")


FIXTURES = [
    ("population", "state", STATE_FIPS, 2020, 4505836, "estimate"),
    ("population", "state", STATE_FIPS, 2025, 4580143, "projection"),
    ("population", "state", STATE_FIPS, 2050, 4785233, "projection"),
    ("population", "county", "21001", 2020, 18903, "estimate"),
    ("population", "county", "21001", 2050, 16291, "projection"),
    ("population", "county", "21111", 2050, 838210, "projection"),
    ("population", "county", "21067", 2050, 398219, "projection"),
    ("households", "state", STATE_FIPS, 2020, 1797937, "estimate"),
    ("households", "state", STATE_FIPS, 2050, 1944624, "projection"),
    ("households", "county", "21001", 2020, 7364, "estimate"),
    ("households", "county", "21001", 2050, 6775, "projection"),
]


def verify(conn) -> None:
    for variable, level, geo_id, year, value, datatype in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype FROM observations WHERE state='KY' AND variable=? "
            "AND geo_level=? AND geo_id=? AND year=?",
            (variable, level, geo_id, year),
        ).fetchone()
        assert row is not None and tuple(row) == (value, datatype), (
            f"KY fixture: {variable}/{geo_id}/{year}"
        )

    for variable in ("population", "households"):
        shape = conn.execute(
            "SELECT COUNT(*), COUNT(DISTINCT geo_id), MIN(year), MAX(year), "
            "SUM(datatype='estimate') FROM observations WHERE state='KY' AND variable=?",
            (variable,),
        ).fetchone()
        assert tuple(shape) == (847, 121, 2020, 2050, 121), f"KY {variable} shape: {tuple(shape)}"

        differences = conn.execute(
            "SELECT s.year, SUM(c.value)-s.value FROM observations s "
            "JOIN observations c ON c.state=s.state AND c.variable=s.variable "
            "AND c.year=s.year AND c.geo_level='county' "
            "WHERE s.state='KY' AND s.variable=? AND s.geo_level='state' "
            "GROUP BY s.year ORDER BY s.year",
            (variable,),
        ).fetchall()
        assert [tuple(row) for row in differences] == [(year, 0) for year in YEARS], (
            f"KY {variable} county totals do not sum to state totals"
        )
    print(f"KY verify: {len(FIXTURES)} fixtures and complete state/county grids ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.ky --acquire")
