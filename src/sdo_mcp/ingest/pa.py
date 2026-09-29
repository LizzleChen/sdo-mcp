"""Pennsylvania State Data Center's 2023 state/county population edition.

Archive the original workbook and supporting documents. Serve its published
totals only: July 2020 estimates and five-year projections from 2025 to 2050.
"""

from __future__ import annotations

import io
import sys
from datetime import date, datetime

from ..common import fetch
from ..store import REPO_ROOT, insert_observations, upsert_source
from .seed import COUNTY_FILE

RAW_DIR = REPO_ROOT / "data" / "raw" / "pa"
FILE = "population_2023.xlsx"
URL = "https://www.rural.pa.gov/getfile.cfm?file=Resources%2Fexcel%2FPopulation+Projections+by+Gender+and+Age+2020+to+2050.xlsx"
DATA_URL = "https://www.rural.pa.gov/data/population-projections"
METHOD_URL = "https://www.rural.pa.gov/getfile.cfm?file=Resources%2Freports%2Fassets%2F260%2FPennsylvania+Population+Projections+for+the+Next+30+Years+Web.pdf"
RELEASE_URL = "https://www.rural.pa.gov/getfile.cfm?file=Resources%2FPDFs%2FPennsylvania+Population+Projections+Press+Release.pdf"
REFERENCES = (
    ("methodology_2023.pdf", METHOD_URL, "An Overview of the Methodology, 2020 to 2050"),
    ("release_2023-10-04.pdf", RELEASE_URL, "October 4, 2023"),
    ("projection_catalog_2026-09-29.html", DATA_URL, "all 67 counties"),
)
SOURCE_ID = "pa-pop-2023"
AGENCY = "Pennsylvania State Data Center, Penn State Harrisburg, for the Center for Rural Pennsylvania"
LICENSE = "No explicit data license identified in the reviewed products; cite the Pennsylvania State Data Center and Center for Rural Pennsylvania."
YEARS = tuple(range(2020, 2051, 5))
AGES = ("Total Population", *(f"Ages {a} to {a + 4}" for a in range(0, 85, 5)), "Ages 85 and Over")
HEADER = ("FIPS Code", "Geography", "Age Group") + tuple(
    f"{year} {sex} ({'Est.' if year == 2020 else 'Proj.'})"
    for year in YEARS for sex in ("Total", "Male", "Female")
)
BASE_NOTE = ("The base population is from the U.S. Census Bureau's population estimate "
             "for the resident population on July 1, 2020, from the Vintage 2022 Population Estimates")


def _geographies() -> dict[int, str]:
    geos = {42: "Pennsylvania"}
    for line in COUNTY_FILE.read_text().splitlines()[1:]:
        parts = line.split("|")
        if parts[0] == "PA":
            geos[int(parts[1] + parts[2])] = parts[4]
    if len(geos) != 68:
        raise ValueError("PA: expected Pennsylvania and 67 Census counties")
    return geos


def parse(body: bytes) -> list[dict]:
    """Require the edition, estimate labels, geography grid and valid numbers."""
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("PA: response is not a readable population workbook") from exc
    try:
        if workbook.sheetnames != ["Source", "Population Projections"]:
            raise ValueError("PA: unexpected workbook sheets")
        source = workbook["Source"]
        if (str(source["A2"].value).strip() != BASE_NOTE or
                source["A9"].value != "The Pennsylvania State Data Center" or
                source["A12"].value != datetime(2023, 9, 5)):
            raise ValueError("PA: unexpected source, base year or edition date")
        rows = iter(workbook["Population Projections"].iter_rows(values_only=True))
        if (next(rows, (None,))[0] != "Pennsylvania Population Projections, 2020 to 2050" or
                next(rows, (None,))[0] != "Source: Pennsylvania State Data Center for the Center for Rural Pennsylvania"):
            raise ValueError("PA: unexpected population title or attribution")
        header = tuple(" ".join(str(v).split()) for v in next(rows, ()))
        if header != HEADER:
            raise ValueError("PA: unexpected year/estimate/projection header")
        geos = _geographies()
        expected = {(code, age) for code in geos for age in AGES}
        seen = set()
        observations = []
        for row in rows:
            if all(v is None for v in row):
                continue
            if len(row) != len(HEADER):
                raise ValueError("PA: malformed population row")
            code, name, age, *values = row
            if type(code) is not int or code not in geos or geos[code] != name:
                raise ValueError(f"PA: unknown geography or mismatched name: {code!r}/{name!r}")
            key = code, age
            if key in seen or key not in expected:
                raise ValueError(f"PA: duplicate or unexpected geography/age: {key}")
            seen.add(key)
            if any(type(v) is not int or v < 0 for v in values):
                raise ValueError(f"PA: invalid integer population value in {key}")
            if age != "Total Population":
                continue
            for year, value in zip(YEARS, values[::3]):
                observations.append(dict(
                    state="PA", geo_level="state" if code == 42 else "county",
                    geo_id=str(code), year=year, variable="population", value=value,
                    datatype="estimate" if year == 2020 else "projection", source_id=SOURCE_ID,
                ))
        if seen != expected:
            raise ValueError(f"PA: incomplete geography/age grid; missing {len(expected - seen)} rows")
        return observations
    finally:
        workbook.close()


def _validate_reference(body: bytes, name: str, marker: str) -> None:
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        fetch.validate_pdf(body, name)
        text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(body)).pages)
    else:
        text = body.decode("utf-8-sig")
        if "<html" not in text.lower():
            raise ValueError(f"PA: unexpected reference page {name}")
    if marker not in " ".join(text.split()):
        raise ValueError(f"PA: unexpected reference document {name}; missing {marker!r}")


def acquire() -> None:
    """Validate before saving; never replace originals or their retrieval dates."""
    if (RAW_DIR / "manifest.json").exists():
        entries = fetch.read_manifest(RAW_DIR)
        required = {FILE} | {name for name, _, _ in REFERENCES}
        if not required <= {entry["file"] for entry in entries}:
            raise ValueError("PA: incomplete saved manifest")
        parse((RAW_DIR / FILE).read_bytes())
        for name, _, marker in REFERENCES:
            _validate_reference((RAW_DIR / name).read_bytes(), name, marker)
        print("PA acquire: reused validated 2023 snapshots")
        return
    body = fetch.get_bytes(URL)
    rows = parse(body)
    downloads = [(FILE, URL, body, len(rows))]
    for name, url, marker in REFERENCES:
        reference = fetch.get_bytes(url)
        _validate_reference(reference, name, marker)
        downloads.append((name, url, reference, None))
    entries = []
    for name, url, body, count in downloads:
        fetch.write_snapshot(RAW_DIR / name, body)
        entries.append(dict(
            file=name, source=url, retrieved_at=date.today().isoformat(),
            sha256=fetch.sha256(body), records=count,
            note="Original Pennsylvania 2023 population edition or supporting evidence. Workbook prepared September 5, 2023; release October 4, 2023; methodology December 2023. 476 state/county totals served; age/sex detail retained in original only.",
        ))
    fetch.write_manifest(RAW_DIR, entries)
    print(f"PA acquire: {len(rows)} totals and three reference documents")


def load(conn) -> None:
    entries = fetch.read_manifest(RAW_DIR)
    retrieved = next(e["retrieved_at"] for e in entries if e["file"] == FILE)
    rows = parse((RAW_DIR / FILE).read_bytes())
    actual = {(r["geo_level"], r["geo_id"]) for r in conn.execute(
        "SELECT geo_level, geo_id FROM geographies WHERE state='PA'"
    )}
    if actual != {(r["geo_level"], r["geo_id"]) for r in rows}:
        raise ValueError("PA: seeded geographies do not match the source geography set")
    upsert_source(
        conn, source_id=SOURCE_ID, state="PA", agency=AGENCY,
        product_name="Pennsylvania State and County Population Projections, 2020-2050",
        vintage="2023", base_year=2020, horizon_year=2050,
        source_url=URL, retrieved_at=retrieved, license=LICENSE,
        citation="Pennsylvania State Data Center for the Center for Rural Pennsylvania. Population Projections, 2020-2050. Released October 4, 2023.",
        method_notes=(
            "Workbook prepared September 5, 2023, released October 4, 2023; methodology report December 2023. "
            "July 1, 2020 estimates from Census Vintage 2022 form the population base: base_year=2020, "
            "not the Census estimate vintage 2022. Published 2025-2050 points are projections, including "
            "2025 when retrieved after that year. Cohort-component methods apply fertility, survival and "
            "migration assumptions to five-year age/sex groups, with separate group-quarters treatment. "
            "Only published five-year population totals are served; no annual interpolation. "
            "County totals sum exactly to the published state totals in this workbook. "
            "Age/sex detail is archived but not served. The report's statewide components table is not loaded. "
            f"Methodology: {METHOD_URL}"
        ),
    )
    count = insert_observations(conn, rows)
    conn.commit()
    print(f"PA load: {count} observations")


# Cross-checked against the methodology report's Table 4, pages 20-21.
FIXTURES = [
    ("42", 2020, 12994440), ("42", 2030, 13185540), ("42", 2050, 13195897),
    ("42003", 2050, 1226933), ("42027", 2050, 177459), ("42053", 2050, 5585),
    ("42083", 2050, 37123), ("42101", 2020, 1600600), ("42101", 2050, 1836216),
]


def verify(conn) -> None:
    for gid, year, value in FIXTURES:
        row = conn.execute(
            "SELECT value, datatype FROM observations WHERE state='PA' AND geo_id=? "
            "AND year=? AND variable='population'", (gid, year),
        ).fetchone()
        expected = (value, "estimate" if year == 2020 else "projection")
        assert row is not None and tuple(row) == expected, f"PA fixture: {gid}/{year}"
    shape = conn.execute(
        "SELECT COUNT(*), COUNT(DISTINCT geo_id), MIN(year), MAX(year), "
        "SUM(datatype='estimate') FROM observations WHERE state='PA'"
    ).fetchone()
    assert tuple(shape) == (476, 68, 2020, 2050, 68)
    print(f"PA verify: {len(FIXTURES)} fixtures and complete state/county grid ok")


if __name__ == "__main__":
    if "--acquire" in sys.argv:
        acquire()
    else:
        print("usage: python -m sdo_mcp.ingest.pa --acquire")
