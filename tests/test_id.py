"""Idaho original workbook grids, unknown estimate boundary and API semantics."""
import io
import zipfile
from functools import lru_cache
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook

from sdo_mcp import server, store
from sdo_mcp.ingest import id as idaho


@lru_cache
def _bodies():
    return {name: (idaho.RAW_DIR / name).read_bytes() for name in idaho.FILES}


def _changed(edit, file=idaho.STATE_FILE):
    book = load_workbook(io.BytesIO(_bodies()[file]), data_only=True)
    edit(book)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _parse_state(body):
    return idaho.parse_workbook(body, ("Total", *(v[0] for v in idaho.REGIONS.values())))


def test_state_exact_grid_and_unknown_estimate_boundary():
    result = server.get_forecast("ID", "state", "Idaho", "population")
    assert result["status"] == "ok"
    assert result["series"] == [dict(year=y, value=v, datatype="projection")
                                for y, v in zip(idaho.YEARS, idaho.STATE_VALUES, strict=True)]
    source = result["source"]
    assert source["vintage"] == "August 7, 2025 (revised)"
    assert source["base_year"] is None and source["horizon_year"] == 2034
    assert source["url"] == idaho.FILES[idaho.STATE_FILE]
    assert source["retrieved_at"] == "2026-10-02"
    assert "last underlying estimate year" in source["method_notes"]
    assert "every loaded point remains projection" in source["method_notes"]


@pytest.mark.parametrize("name,gid,year,value", [
    ("Ada", "16001", 2024, 533654), ("Ada", "16001", 2025, 542752), ("Ada", "16001", 2034, 618118),
    ("Canyon", "16027", 2034, 320984), ("Clark", "16033", 2024, 825), ("Clark", "16033", 2025, 817),
    ("Clark", "16033", 2034, 725), ("Camas", "16025", 2034, 1425),
])
def test_county_name_id_and_original_anchors(name, gid, year, value):
    result = server.get_forecast("ID", "county", name, "population", year, year)
    assert result == server.get_forecast("ID", "county", gid, "population", year, year)
    assert result["series"] == [dict(year=year, value=value, datatype="projection")]
    assert result["source"]["url"] == idaho.CATALOG_URL


def test_independent_original_xml_cell_fixtures():
    # Read numeric XML directly, independently of the openpyxl adapter.
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    for file, sheet, cells in [
        (idaho.STATE_FILE, 1, {"B21":2047050, "C21":2080283, "L21":2364544}),
        ("Population_Southwestern_revised2025-08-07.xlsx", 2, {"B21":533654,"C21":542752,"L21":618118}),
        ("Population_Eastern_revised2025-08-07.xlsx", 4, {"B21":825,"C21":817,"L21":725}),
        ("Population_South-central_revised2025-08-07.xlsx", 3, {"B21":1289,"L21":1425}),
    ]:
        with zipfile.ZipFile(io.BytesIO(_bodies()[file])) as archive:
            root = ET.fromstring(archive.read(f"xl/worksheets/sheet{sheet}.xml"))
        for cell, expected in cells.items():
            value = root.find(f".//s:c[@r='{cell}']/s:v", ns)
            assert value is not None and int(value.text) == expected


def test_complete_counties_deferred_regions_and_other_products():
    conn = store.connect()
    try:
        idaho.verify(conn)
        assert conn.execute("SELECT COUNT(DISTINCT geo_id) FROM observations WHERE state='ID' AND geo_level='county'").fetchone()[0] == 44
    finally:
        conn.close()
    assert server.get_forecast("ID", "economic_region", "Northern", "population")["status"] == "not_loaded"
    assert server.get_forecast("ID", "place", "Boise", "population")["status"] == "not_published"
    for variable in ("households", "employment"):
        assert server.get_forecast("ID", "county", "Ada", variable)["status"] == "coverage_unknown"
    assert server.get_components("ID", "Ada")["status"] == "coverage_unknown"
    assert server.describe_source("ID")["review"]["status"] == "never_reviewed"


@pytest.mark.parametrize("start,end", [(2020,2023),(2035,2050)])
def test_no_generated_years(start, end):
    result = server.get_forecast("ID", "state", "ID", "population", start, end)
    assert result["status"] == "empty_range"
    assert result["available_start_year"] == 2024 and result["available_end_year"] == 2034


def test_invalid_scenario_and_unknown_county():
    assert server.get_forecast("ID", "county", "Ada", "population", scenario="High")["status"] == "error"
    result = server.get_forecast("ID", "county", "No such county", "population")
    assert result["status"] == "error" and result["candidates"] is None
    ambiguous = server.get_forecast("ID", "county", "Bon", "population")
    assert ambiguous["status"] == "error" and len(ambiguous["candidates"]) == 2


@pytest.mark.parametrize("cell,value,error", [
    ("A1","Population","header"),("L1",2035,"year grid"),
    ("A21","Population","row labels"),("A2","Total","row labels"),
    ("B21",None,"integer"),("B21","NA","integer"),("B21",-1,"integer"),
    ("B21",float("inf"),"integer"),("B21",2047050.5,"integer"),
    ("B21",True,"integer"),("B21","=SUM(B2:B19)","integer"),
    ("B2",123,"disagree"),("B26",123,"disagree"),
])
def test_changed_headers_labels_and_bad_population_rejected(cell,value,error):
    with pytest.raises(ValueError, match=error):
        _parse_state(_changed(lambda w:w["Total"].__setitem__(cell,value)))


def test_missing_duplicate_and_unknown_geographies_rejected():
    for edit in (lambda w:w.remove(w["6. East"]), lambda w:w.create_sheet("Total copy"),
                 lambda w:setattr(w["6. East"],"title","Unknown")):
        with pytest.raises(ValueError, match="geography sheets"):
            _parse_state(_changed(edit))
    with pytest.raises(ValueError, match="trailing"):
        _parse_state(_changed(lambda w:w["Total"].__setitem__("A31","unexpected")))


@pytest.mark.parametrize("body", [b"<html>temporarily unavailable</html>",b"[]",b"PK broken"])
def test_soft_200_rejected(body):
    with pytest.raises(ValueError, match="readable XLSX"):
        _parse_state(body)


def test_wrong_readable_workbook_rejected():
    output=io.BytesIO()
    Workbook().save(output)
    with pytest.raises(ValueError, match="geography sheets"):
        _parse_state(output.getvalue())


def test_snapshot_reuse_has_no_network_or_retrieval_change(monkeypatch):
    original=(idaho.RAW_DIR/"manifest.json").read_bytes()
    monkeypatch.setattr(idaho.fetch,"get_bytes",lambda *_:pytest.fail("unexpected network"))
    idaho.acquire()
    assert (idaho.RAW_DIR/"manifest.json").read_bytes()==original


@pytest.mark.parametrize("file", list(idaho.FILES))
def test_failed_acquisition_saves_nothing(file,tmp_path,monkeypatch):
    bodies=dict(_bodies())
    bodies[file]=b"<html>temporary error</html>"
    monkeypatch.setattr(idaho,"RAW_DIR",tmp_path/"id")
    monkeypatch.setattr(idaho.fetch,"get_bytes",lambda url:bodies[next(k for k,v in idaho.FILES.items() if v==url)])
    with pytest.raises(ValueError,match="readable XLSX|catalog|readable PDF|product evidence"):
        idaho.acquire()
    assert not idaho.RAW_DIR.exists()


def test_missing_seeded_geography_prevents_load(tmp_path):
    conn=store.connect(tmp_path/"db.sqlite",create=True)
    try:
        store.init_db(conn)
        with pytest.raises(ValueError,match="seeded geographies"):
            idaho.load(conn)
        assert conn.execute("SELECT COUNT(*) FROM sources").fetchone()[0]==0
    finally:
        conn.close()
