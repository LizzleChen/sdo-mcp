"""Ohio original workbook validation, source controls and public semantics."""

import io
import json
from datetime import datetime
from functools import lru_cache

import openpyxl
import pytest
from pypdf import PdfWriter

from sdo_mcp import server, store
from sdo_mcp.ingest import oh


@lru_cache
def _bodies():
    return {name: (oh.RAW_DIR / name).read_bytes() for name in oh.FILES}


@lru_cache
def _grid():
    book = openpyxl.load_workbook(io.BytesIO(_bodies()[oh.FILE]), read_only=True, data_only=True)
    try:
        return tuple(tuple(row) for row in book["data"].iter_rows(values_only=True))
    finally:
        book.close()


def _changed_grid(name="Ohio", year=2020):
    grid = [list(row) for row in _grid()]
    row = next(row for row in grid[1:] if row[:2] == [name, year])
    return grid, row


@lru_cache
def _reference_pages():
    return {
        "report": oh._pdf_pages(_bodies()[oh.REPORT_FILE], 16, "report"),
        "overview": oh._pdf_pages(_bodies()[oh.OVERVIEW_FILE], 2, "overview"),
    }


def test_exact_state_points_and_source_metadata():
    result = server.get_forecast("OH", "state", "Ohio", "population")
    assert result["status"] == "ok"
    assert result["geo"] == dict(level="state", id="39", name="Ohio")
    assert result["scenario"] is None
    # Independently checked against the rendered Development county-totals
    # table; report Table 1 and the overview also corroborate both endpoints.
    assert result["series"] == [
        dict(year=year, value=value, datatype="estimate" if year == 2020 else "projection")
        for year, value in zip(range(2020, 2051, 5),
                               (11799448, 11765739, 11694767, 11574119, 11425531, 11267043, 11123896))
    ]
    source = result["source"]
    assert source["vintage"] == "2023 edition"
    assert source["base_year"] == 2020
    assert source["horizon_year"] == 2050
    assert source["url"] == oh.URL
    assert source["retrieved_at"] == next(
        entry["retrieved_at"] for entry in oh.fetch.read_manifest(oh.RAW_DIR) if entry["file"] == oh.FILE
    )
    assert "Total_TOT" in source["citation"]
    assert "2020 estimates are the base" in source["method_notes"]
    assert "overview cites a 2022 series" in source["method_notes"]
    assert "An exact release date or later rebasing was not established" in source["method_notes"]


@pytest.mark.parametrize("name,gid,year,value", [
    # Franklin base and first projection: independently rendered county table.
    ("Franklin", "39049", 2020, 1323807), ("Franklin", "39049", 2025, 1371832),
    # Horizon values independently read from the publisher's overview map,
    # including Fulton, Mercer and Vinton omitted by the housing reprint.
    ("Franklin", "39049", 2050, 1487885), ("Cuyahoga", "39035", 2050, 1030507),
    ("Fulton", "39051", 2050, 36801), ("Mercer", "39107", 2050, 41766),
    ("Vinton", "39163", 2050, 9117), ("Wyandot", "39175", 2050, 16233),
])
def test_county_names_fips_and_independent_anchors(name, gid, year, value):
    result = server.get_forecast("OH", "county", name, "population", year, year)
    assert result == server.get_forecast("OH", "county", gid, "population", year, year)
    assert result["geo"] == dict(level="county", id=gid, name=f"{name} County")
    assert result["series"] == [dict(year=year, value=value,
                                     datatype="estimate" if year == 2020 else "projection")]


def test_workbook_ottawa_total_retains_disclosed_map_difference():
    result = server.get_forecast("OH", "county", "Ottawa", "population", 2050, 2050)
    assert result["series"] == [dict(year=2050, value=31371, datatype="projection")]
    assert "overview map prints Ottawa 2050 as 31,374" in result["source"]["method_notes"]


def test_complete_county_grid_and_published_state_controls():
    conn = store.connect()
    try:
        oh.verify(conn)
        rows = conn.execute("SELECT * FROM observations WHERE state='OH'").fetchall()
        assert len(rows) == 623
        assert {row["geo_level"] for row in rows} == {"state", "county"}
        assert {row["variable"] for row in rows} == {"population"}
        counties = conn.execute(
            "SELECT geo_id, COUNT(*) FROM observations WHERE state='OH' AND geo_level='county' GROUP BY geo_id"
        ).fetchall()
        assert len(counties) == 88 and all(row[1] == 7 for row in counties)
        for year in range(2020, 2051, 5):
            county_sum = sum(row["value"] for row in rows if row["year"] == year and row["geo_level"] == "county")
            state = next(row["value"] for row in rows if row["year"] == year and row["geo_level"] == "state")
            assert county_sum == state
    finally:
        conn.close()


@pytest.mark.parametrize("start,end", [(2010, 2019), (2021, 2024), (2026, 2029), (2051, 2060)])
def test_no_generated_years_or_extended_horizon(start, end):
    for level, geo in (("state", "OH"), ("county", "Vinton")):
        result = server.get_forecast("OH", level, geo, "population", start, end)
        assert result["status"] == "empty_range"
        assert result["available_start_year"] == 2020
        assert result["available_end_year"] == 2050
        assert "series" not in result


def test_unverified_products_remain_unknown_and_inventory_matches():
    for level, geo in (("state", "OH"), ("county", "Franklin")):
        for variable in ("households", "employment"):
            assert server.get_forecast("OH", level, geo, variable)["status"] == "coverage_unknown"
        assert server.get_components("OH", geo)["status"] == "coverage_unknown"
        assert server.get_forecast("OH", level, geo, "population", scenario="high")["status"] == "error"
    assert server.get_forecast("OH", "place", "Columbus", "population")["status"] == "coverage_unknown"
    summary = server.list_states()["states"]["OH"]
    for level in ("state", "county"):
        assert summary["variables"]["population"][level]["load_status"] == "loaded"
    assert summary["scenarios"] is None
    description = server.describe_source("OH")
    assert description["variables"] == summary["variables"]
    assert description["review"]["status"] == "never_reviewed"
    assert len(description["sources_loaded"]) == 1


@pytest.mark.parametrize("body", [b"<html>Service unavailable</html>", b"PKbroken", b"[]"])
def test_soft_200_workbook_rejected(body):
    with pytest.raises(ValueError, match="readable XLSX workbook"):
        oh.parse(body)


@pytest.mark.parametrize("change", ["sheet", "creator", "edition"])
def test_readable_workbook_with_wrong_edition_rejected(change):
    book = openpyxl.Workbook()
    book.active.title = "data" if change != "sheet" else "households"
    book.properties.creator = "Howison, Jeffrey" if change != "creator" else "Other producer"
    book.properties.modified = datetime(2024, 1, 6) if change == "edition" else datetime(2023, 1, 6, 19, 7, 39)
    output = io.BytesIO()
    book.save(output)
    book.close()
    with pytest.raises(ValueError, match="sheets or edition"):
        oh.parse(output.getvalue())


@pytest.mark.parametrize("column,value,error", [
    (0, "Unknown County", "unknown geography"), (0, "Adams", "duplicate geography/year"),
    (1, 2026, "unexpected year"), (1, 2020, "duplicate geography/year"),
    (2, None, "invalid integer population"), (2, True, "invalid integer population"),
    (2, -1, "invalid integer population"), (2, float("nan"), "invalid integer population"),
    (2, float("inf"), "invalid integer population"), (2, 0.5, "invalid integer population"),
    (2, "100", "invalid integer population"), (59, 1, "extra data"),
])
def test_changed_geography_year_and_invalid_values_rejected(column, value, error):
    grid, row = _changed_grid("Franklin", 2025)
    row[column] = value
    with pytest.raises(ValueError, match=error):
        oh._parse_grid(grid)


@pytest.mark.parametrize("change", ["header", "missing_row", "missing_column", "extra_column"])
def test_changed_header_or_grid_shape_rejected(change):
    grid = [list(row) for row in _grid()]
    if change == "header":
        grid[0][58] = "Household_TOT"
    elif change == "missing_row":
        grid.pop()
    elif change == "missing_column":
        grid[1].pop()
    else:
        grid[1].append(0)
    with pytest.raises(ValueError, match="header or shape"):
        oh._parse_grid(grid)


def test_changed_age_sum_rejected():
    grid, row = _changed_grid("Franklin", 2025)
    row[2] += 1
    with pytest.raises(ValueError, match="age/sex totals"):
        oh._parse_grid(grid)


def test_published_state_base_age_offsets_are_checked_exactly():
    grid, row = _changed_grid()
    # Preserve every row total while changing the published sex-age offsets.
    row[2] += 1
    row[3] -= 1
    with pytest.raises(ValueError, match="age/sex cell difference"):
        oh._parse_grid(grid)


def test_reconciled_state_ages_must_match_independent_report():
    grid, row = _changed_grid("Ohio", 2050)
    # Preserve male/female/both totals and age reconciliation, but transfer
    # one person across report Table 1's broad age-group boundary.
    for column, delta in ((2, 1), (5, -1), (40, 1), (43, -1)):
        row[column] += delta
    with pytest.raises(ValueError, match="independent report anchors"):
        oh._parse_grid(grid)


def test_reconciled_county_change_must_match_published_state():
    grid, row = _changed_grid("Franklin", 2025)
    for column in (2, 20, 40, 58):
        row[column] += 1
    with pytest.raises(ValueError, match="county/state totals"):
        oh._parse_grid(grid)


@pytest.mark.parametrize("label,page,old,new,error", [
    ("report", 0, "Population Projections Report 2023", "Population Projections Report 2024", "edition"),
    ("report", 14, "population estimates for 2020", "population estimates for 2021", "base year"),
    ("overview", 1, "11,123,896", "11,123,897", "state totals"),
    ("report", 1, "2,170,245", "2,170,246", "age anchors"),
])
def test_changed_reference_edition_base_or_independent_anchors_rejected(monkeypatch, label, page, old, new, error):
    pages = {key: list(value) for key, value in _reference_pages().items()}
    assert old in pages[label][page]
    pages[label][page] = pages[label][page].replace(old, new, 1)
    monkeypatch.setattr(oh, "_pdf_pages", lambda _body, _count, key: pages[key])
    with pytest.raises(ValueError, match=error):
        oh._validate_references(_bodies()[oh.REPORT_FILE], _bodies()[oh.OVERVIEW_FILE])


def test_wrong_readable_reference_pdf_rejected():
    writer = PdfWriter()
    writer.add_blank_page(612, 792)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="report page count"):
        oh._validate_references(output.getvalue(), _bodies()[oh.OVERVIEW_FILE])


@pytest.mark.parametrize("key,value", [
    ("format", "http-response"), ("source_url", "https://example.com/population"),
    ("retrieved_at", "9999-12-31"), ("page_title", "Error"),
    ("extraction", "Downloaded original HTML"), ("links", []),
])
def test_invalid_catalog_capture_rejected(key, value):
    catalog = json.loads(_bodies()[oh.CATALOG_FILE])
    catalog[key] = value
    with pytest.raises(ValueError, match="invalid browser catalog"):
        oh._validate_catalog(json.dumps(catalog).encode())


def test_snapshot_reuse_is_offline_and_preserves_every_file(monkeypatch):
    before = {name: (oh.RAW_DIR / name).read_bytes() for name in (*oh.FILES, "manifest.json")}
    monkeypatch.setattr(oh.fetch, "get_bytes", lambda *_: pytest.fail("unexpected network call"))
    oh.acquire()
    assert {name: (oh.RAW_DIR / name).read_bytes() for name in before} == before


def test_new_acquisition_requires_valid_browser_capture_before_network(tmp_path, monkeypatch):
    raw = tmp_path / "oh"
    monkeypatch.setattr(oh, "RAW_DIR", raw)
    monkeypatch.setattr(oh.fetch, "get_bytes", lambda *_: pytest.fail("unexpected network call"))
    with pytest.raises(ValueError, match="browser catalog excerpt"):
        oh.acquire()
    with pytest.raises(ValueError, match="invalid browser catalog"):
        oh.acquire(catalog_excerpt=b"<html>Not a browser capture</html>")
    assert not raw.exists()


@pytest.mark.parametrize("failed_file", [oh.FILE, oh.REPORT_FILE, oh.OVERVIEW_FILE])
def test_invalid_download_archives_nothing(tmp_path, monkeypatch, failed_file):
    bodies = dict(_bodies())
    catalog = bodies[oh.CATALOG_FILE]
    bodies[failed_file] = b"<html>temporary error</html>"
    by_url = {url: bodies[name] for name, url in oh.FILES.items()}
    raw = tmp_path / "oh"
    monkeypatch.setattr(oh, "RAW_DIR", raw)
    monkeypatch.setattr(oh.fetch, "get_bytes", lambda url: by_url[url])
    with pytest.raises(ValueError, match="readable PDF|readable XLSX workbook"):
        oh.acquire(catalog_excerpt=catalog)
    assert not raw.exists()


def test_successful_acquisition_preserves_originals_and_reuses_manifest(tmp_path, monkeypatch):
    bodies = dict(_bodies())
    catalog = bodies[oh.CATALOG_FILE]
    by_url = {url: bodies[name] for name, url in oh.FILES.items()}
    raw = tmp_path / "oh"
    monkeypatch.setattr(oh, "RAW_DIR", raw)
    monkeypatch.setattr(oh.fetch, "get_bytes", lambda url: by_url[url])
    oh.acquire(catalog_excerpt=catalog)
    assert {name: (raw / name).read_bytes() for name in oh.FILES} == bodies
    entries = {entry["file"]: entry for entry in oh.fetch.read_manifest(raw)}
    assert set(entries) == set(oh.FILES)
    assert entries[oh.FILE]["records"] == 623
    assert entries[oh.FILE]["source"] == oh.DOWNLOAD_URL
    assert entries[oh.CATALOG_FILE]["source"] == oh.CATALOG_URL
    assert "EXCERPT, not original HTTP HTML" in entries[oh.CATALOG_FILE]["note"]
    before = (raw / "manifest.json").read_bytes()
    monkeypatch.setattr(oh.fetch, "get_bytes", lambda *_: pytest.fail("unexpected network call"))
    # Supplying a new excerpt never replaces a saved, validated edition.
    oh.acquire(catalog_excerpt=b"new catalog response")
    assert (raw / "manifest.json").read_bytes() == before
    assert {name: (raw / name).read_bytes() for name in oh.FILES} == bodies


def test_tampered_saved_snapshot_is_rejected_without_retrieval(tmp_path, monkeypatch):
    bodies = dict(_bodies())
    manifest = (oh.RAW_DIR / "manifest.json").read_bytes()
    raw = tmp_path / "oh"
    raw.mkdir()
    for name, body in bodies.items():
        (raw / name).write_bytes(body)
    (raw / "manifest.json").write_bytes(manifest)
    (raw / oh.FILE).write_bytes(b"changed original")
    monkeypatch.setattr(oh, "RAW_DIR", raw)
    monkeypatch.setattr(oh.fetch, "get_bytes", lambda *_: pytest.fail("unexpected network call"))
    with pytest.raises(ValueError, match="fingerprint mismatch"):
        oh.acquire()
    assert (raw / oh.FILE).read_bytes() == b"changed original"
    assert (raw / "manifest.json").read_bytes() == manifest


def test_incomplete_saved_manifest_is_rejected_without_retrieval(tmp_path, monkeypatch):
    raw = tmp_path / "oh"
    raw.mkdir()
    (raw / "manifest.json").write_text("[]\n")
    monkeypatch.setattr(oh, "RAW_DIR", raw)
    monkeypatch.setattr(oh.fetch, "get_bytes", lambda *_: pytest.fail("unexpected network call"))
    with pytest.raises(ValueError, match="incomplete saved manifest"):
        oh.acquire()
    assert {path.name for path in raw.iterdir()} == {"manifest.json"}
